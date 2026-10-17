"""Journaux de caisse et de banque joints aux rapports planifiés."""

from __future__ import annotations

from datetime import datetime, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest

from app.services.journaux_annexes import construire_annexes

from tests.test_justificatifs_sorties import _payer
from tests.test_treasury_flows import _admin, _caisse, _depense_poste, _org, _service


@pytest.mark.asyncio
async def test_le_journal_de_caisse_de_la_periode_est_joint_en_pdf(db_session, monkeypatch):
    db = db_session
    org = await _org(db)
    poste = await _depense_poste(db, org)
    service = await _service(db, org)
    await _caisse(db, org, usd=Decimal("1000"))
    await db.commit()
    user = await _admin(db, org)
    org_id = org.id

    await _payer(db, monkeypatch, org, user, poste, service)

    tz = ZoneInfo("UTC")
    debut = datetime.now(tz).replace(hour=0, minute=0, second=0, microsecond=0) - timedelta(days=3)
    fin = debut + timedelta(days=7)
    annexes = await construire_annexes(
        db, tenant_id=org_id, organisation="Organisation test", start=debut, end=fin, tz=tz
    )

    noms = [nom for nom, _ in annexes]
    # Même nom que le PDF tiré de l'écran Rapports.
    assert any(nom.startswith("Journal_") and "_Caisse_USD_" in nom for nom in noms), noms
    # La caisse CDF n'a ni mouvement ni solde : elle n'alourdit pas le courriel.
    assert not any("_Caisse_CDF_" in nom for nom in noms), noms
    for nom, contenu in annexes:
        assert nom.endswith(".pdf")
        assert contenu.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_le_rapport_mensuel_joint_le_journal_en_excel_au_modele_de_l_ecran(db_session, monkeypatch):
    import io

    from openpyxl import load_workbook

    db = db_session
    org = await _org(db)
    poste = await _depense_poste(db, org)
    service = await _service(db, org)
    await _caisse(db, org, usd=Decimal("1000"))
    await db.commit()
    user = await _admin(db, org)
    org_id = org.id
    await _payer(db, monkeypatch, org, user, poste, service)

    tz = ZoneInfo("UTC")
    debut = datetime.now(tz).replace(hour=0, minute=0, second=0, microsecond=0) - timedelta(days=3)
    annexes = await construire_annexes(
        db, tenant_id=org_id, organisation="Organisation test", start=debut, end=debut + timedelta(days=7),
        tz=tz, format="xlsx",
    )
    nom, contenu = next((n, c) for n, c in annexes if "_Caisse_USD_" in n)
    assert nom.endswith(".xlsx")

    ws = load_workbook(io.BytesIO(contenu)).active
    assert ws["A1"].value.startswith("JOURNAL DE TRÉSORERIE — Caisse USD")
    assert [ws.cell(row=4, column=c).value for c in (1, 2, 3, 4, 5)] == [
        "Solde initial", "Total entrées", "Total sorties", "Solde final", "Mouvements",
    ]
    assert [ws.cell(row=7, column=c).value for c in range(1, 7)] == ["N°", "Date", "Libellé", "Entrée", "Sortie", "Solde"]
    assert ws.cell(row=8, column=1).value == 1
    assert ws.cell(row=8, column=5).value == 100
    total = ws.max_row
    assert ws.cell(row=total, column=1).value == "TOTAL MOUVEMENTS"
    assert ws.freeze_panes == "A8"
