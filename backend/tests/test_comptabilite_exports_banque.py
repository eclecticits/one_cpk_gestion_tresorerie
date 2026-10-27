"""Contrôles du contrat d'export et du rapprochement bancaire."""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

import app.db.session  # noqa: F401 — active le cloisonnement tenant
from app.core.tenant_context import set_current_tenant_id
from app.models.compte_bancaire import CompteBancaire
from app.models.organisation import Organisation
from app.modules.comptabilite.models import (
    ComptaCompte,
    ComptaEcriture,
    ComptaExercice,
    ComptaJournal,
    ComptaLigneEcriture,
    ComptaReferentiel,
    ComptaSociete,
)
from app.modules.comptabilite.schemas.exports import ExportComptableFiltres
from app.modules.comptabilite.services.export_service import (
    ENTETE_LOGICIEL,
    creer_xlsx_logiciel,
    historiser_export,
    preparer_export_comptable,
)
from app.modules.comptabilite.services.rapprochement_bancaire import (
    Candidat,
    importer_releve,
    parser_releve,
    scorer,
)


def _org() -> Organisation:
    suffix = uuid.uuid4().hex[:10]
    now = datetime.now(timezone.utc)
    return Organisation(
        nom=f"Export {suffix}",
        slug=f"export-{suffix}",
        plan_type="ACTIVE",
        status_abonnement="ACTIVE",
        limite_utilisateurs=10,
        is_active=True,
        created_at=now,
        updated_at=now,
    )


async def _setup(db, organisation_id: int):
    societe = ComptaSociete(
        organisation_id=organisation_id,
        code="SOC",
        raison_sociale="Société export",
    )
    referentiel = ComptaReferentiel(
        organisation_id=organisation_id,
        code="REF",
        libelle="Plan export",
        type_referentiel="PERSONNALISE",
    )
    db.add_all([societe, referentiel])
    await db.flush()
    banque = ComptaCompte(
        organisation_id=organisation_id,
        referentiel_id=referentiel.id,
        numero="5210000001",
        libelle="Banque",
        nature="ACTIF",
        sens_normal="DEBIT",
    )
    produit = ComptaCompte(
        organisation_id=organisation_id,
        referentiel_id=referentiel.id,
        numero="7060000001",
        libelle="Cotisations",
        nature="PRODUIT",
        sens_normal="CREDIT",
    )
    db.add_all([banque, produit])
    await db.flush()
    journal = ComptaJournal(
        organisation_id=organisation_id,
        societe_id=societe.id,
        code="BQ",
        libelle="Banque",
        type_journal="BQ",
    )
    exercice = ComptaExercice(
        organisation_id=organisation_id,
        societe_id=societe.id,
        code="2026",
        date_debut=date(2026, 1, 1),
        date_fin=date(2026, 12, 31),
        referentiel_id=referentiel.id,
        statut="OUVERT",
    )
    compte_bancaire = CompteBancaire(
        organisation_id=organisation_id,
        intitule="Banque non mappée",
        numero_compte="001-TEST",
        devise="USD",
        account_type="BANK",
    )
    db.add_all([journal, exercice, compte_bancaire])
    await db.flush()
    return societe, banque, produit, journal, exercice, compte_bancaire


async def _ecriture(
    db,
    *,
    organisation_id: int,
    societe: ComptaSociete,
    banque: ComptaCompte,
    produit: ComptaCompte,
    journal: ComptaJournal,
    exercice: ComptaExercice,
    numero: str,
    devise: str = "USD",
    statut: str = "VALIDEE",
):
    ecriture = ComptaEcriture(
        organisation_id=organisation_id,
        societe_id=societe.id,
        exercice_id=exercice.id,
        journal_id=journal.id,
        numero=numero,
        date_ecriture=date(2026, 10, 15),
        reference_piece=f"REC-{numero}",
        libelle="Cotisation octobre",
        statut=statut,
        devise=devise,
        module_origine="encaissements",
        type_origine="encaissement",
        objet_origine_id=str(uuid.uuid4()),
    )
    db.add(ecriture)
    await db.flush()
    db.add_all([
        ComptaLigneEcriture(
            organisation_id=organisation_id,
            societe_id=societe.id,
            ecriture_id=ecriture.id,
            compte_id=banque.id,
            ordre=1,
            libelle="Encaissement cotisation",
            debit=Decimal("125.40"),
            credit=Decimal("0"),
            debit_tenue=Decimal("125.40"),
            credit_tenue=Decimal("0"),
            devise=devise,
        ),
        ComptaLigneEcriture(
            organisation_id=organisation_id,
            societe_id=societe.id,
            ecriture_id=ecriture.id,
            compte_id=produit.id,
            ordre=2,
            libelle="Cotisation membre",
            debit=Decimal("0"),
            credit=Decimal("125.40"),
            debit_tenue=Decimal("0"),
            credit_tenue=Decimal("125.40"),
            devise=devise,
        ),
    ])
    await db.flush()
    return ecriture


async def test_export_excel_strict_filtre_statuts_devises_et_reexport(db_session):
    set_current_tenant_id(None)
    org = _org()
    db_session.add(org)
    await db_session.flush()
    set_current_tenant_id(org.id)
    societe, banque, produit, journal, exercice, _ = await _setup(db_session, org.id)
    active = await _ecriture(
        db_session,
        organisation_id=org.id,
        societe=societe,
        banque=banque,
        produit=produit,
        journal=journal,
        exercice=exercice,
        numero="BQ-2026-00001",
    )
    await _ecriture(
        db_session,
        organisation_id=org.id,
        societe=societe,
        banque=banque,
        produit=produit,
        journal=journal,
        exercice=exercice,
        numero="BQ-2026-00002",
        devise="CDF",
    )
    await _ecriture(
        db_session,
        organisation_id=org.id,
        societe=societe,
        banque=banque,
        produit=produit,
        journal=journal,
        exercice=exercice,
        numero="BQ-2026-00003",
        statut="ANNULEE",
    )
    filtres = ExportComptableFiltres(
        date_debut=date(2026, 10, 1),
        date_fin=date(2026, 10, 31),
        journal="BANQUE",
        devise="USD",
        format="XLSX",
    )
    selection = await preparer_export_comptable(db_session, organisation_id=org.id, filtres=filtres)
    assert selection.apercu.exportable is True
    assert selection.apercu.nombre_ecritures == 1
    assert selection.apercu.nombre_lignes == 2
    assert selection.apercu.total_debit == selection.apercu.total_credit == Decimal("125.40")

    contenu = creer_xlsx_logiciel(selection.lignes_export)
    feuille = load_workbook(io.BytesIO(contenu), data_only=True).active
    assert tuple(cell.value for cell in feuille[1]) == ENTETE_LOGICIEL
    assert feuille.max_column == 6
    assert feuille.max_row == 3
    assert isinstance(feuille["A2"].value, datetime)
    assert feuille["A2"].number_format == "mm-dd-yy"
    assert isinstance(feuille["C2"].value, int)
    assert feuille["C2"].number_format == "0"
    assert isinstance(feuille["E2"].value, (int, float))
    assert feuille["F2"].value is None

    await historiser_export(
        db_session,
        organisation_id=org.id,
        user_id=None,
        type_export="LOGICIEL_COMPTABLE",
        filtres=filtres,
        nom_fichier="import.xlsx",
        contenu=contenu,
        empreinte_selection=selection.empreinte,
        nombre_ecritures=1,
        nombre_lignes=2,
        total_debit=Decimal("125.40"),
        total_credit=Decimal("125.40"),
        ecriture_ids=[active.id],
    )
    await db_session.flush()
    seconde_selection = await preparer_export_comptable(
        db_session, organisation_id=org.id, filtres=filtres
    )
    assert seconde_selection.apercu.exportable is False
    assert seconde_selection.apercu.nombre_deja_exportees == 1
    assert "DEJA_EXPORTEE" in {item.code for item in seconde_selection.apercu.anomalies}

    filtres_reexport = filtres.model_copy(update={"autoriser_reexport": True})
    reexport = await preparer_export_comptable(
        db_session, organisation_id=org.id, filtres=filtres_reexport
    )
    assert reexport.apercu.exportable is True
    assert reexport.premier_export_id is not None
    set_current_tenant_id(None)


async def test_filtre_compte_bancaire_refuse_un_mapping_absent(db_session):
    set_current_tenant_id(None)
    org = _org()
    db_session.add(org)
    await db_session.flush()
    set_current_tenant_id(org.id)
    *_, compte_bancaire = await _setup(db_session, org.id)
    filtres = ExportComptableFiltres(
        date_debut=date(2026, 10, 1),
        date_fin=date(2026, 10, 31),
        journal="BANQUE",
        compte_bancaire_id=compte_bancaire.id,
        devise="USD",
    )
    with pytest.raises(HTTPException) as exc_info:
        await preparer_export_comptable(db_session, organisation_id=org.id, filtres=filtres)
    assert exc_info.value.status_code == 400
    assert "compte comptable associé" in exc_info.value.detail
    set_current_tenant_id(None)


def test_parser_releve_preserve_solde_negatif_et_exige_debit_credit_distincts():
    contenu = (
        "Date;Date de valeur;Référence;Bénéficiaire;Libellé;Débit;Crédit;Solde\n"
        "15/10/2026;16/10/2026;FR-01;Fournisseur;Frais bancaires;15,50;;(120,75)\n"
        "17/10/2026;;REC-10;Membre;Cotisation;;125,40;4,65\n"
    ).encode("utf-8")
    lignes = parser_releve(contenu, "releve.csv")
    assert len(lignes) == 2
    assert lignes[0].debit == Decimal("15.50")
    assert lignes[0].credit == Decimal("0")
    assert lignes[0].solde == Decimal("-120.75")
    assert lignes[1].credit == Decimal("125.40")

    candidat = Candidat(
        operation_type="encaissement",
        operation_id=str(uuid.uuid4()),
        date_operation=date(2026, 10, 17),
        reference="REC-10",
        libelle="Cotisation",
        tiers="Membre",
        debit=Decimal("0"),
        credit=Decimal("125.40"),
    )
    assert scorer(lignes[1], candidat) >= Decimal("95")
    assert scorer(lignes[0], candidat) == Decimal("0")


async def test_import_releve_ne_valide_jamais_seul_et_bloque_fichier_duplique(db_session):
    set_current_tenant_id(None)
    org = _org()
    db_session.add(org)
    await db_session.flush()
    set_current_tenant_id(org.id)
    *_, compte_bancaire = await _setup(db_session, org.id)
    contenu = (
        "Date;Référence;Libellé;Débit;Crédit;Solde\n"
        "20/10/2026;FR-20;Frais de tenue de compte;10,00;;90,00\n"
    ).encode("utf-8")
    releve = await importer_releve(
        db_session,
        organisation_id=org.id,
        user_id=None,
        compte_bancaire_id=compte_bancaire.id,
        filename="releve_octobre.csv",
        content=contenu,
    )
    assert releve.nombre_lignes == 1
    assert releve.nombre_rapprochees == 0
    assert releve.nombre_propositions == 0

    with pytest.raises(HTTPException) as exc_info:
        await importer_releve(
            db_session,
            organisation_id=org.id,
            user_id=None,
            compte_bancaire_id=compte_bancaire.id,
            filename="copie.csv",
            content=contenu,
        )
    assert exc_info.value.status_code == 409
    set_current_tenant_id(None)
