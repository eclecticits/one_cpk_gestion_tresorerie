"""Les avances « à valoir » sautent aux yeux dans le journal et ses exports.

Une avance à valoir est une créance sur le bénéficiaire : tant qu'elle n'est ni
justifiée ni rendue, les utilisateurs et les autorités doivent la voir. Le
journal la signale (bordeaux dans les exports) et dit où elle en est.
"""

import io
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest

from app.api.v1.endpoints.reports import journal_tresorerie
from app.models.organisation import Organisation
from app.models.requisition import Requisition
from app.models.retour_caisse import RetourCaisse
from app.models.sortie_fonds import SortieFonds
from app.models.user import User
from app.services.journaux_annexes import Identite, JournalAnnexe, _pdf, _xlsx

AVANCE_LE = datetime(2026, 9, 15, 10, tzinfo=timezone.utc)


async def _avance_a_valoir(db, *, justificatif="ATTENDU", echeance=AVANCE_LE + timedelta(hours=48)):
    org = Organisation(nom="À valoir", slug=f"av-{uuid.uuid4().hex[:8]}", is_active=True)
    db.add(org)
    await db.flush()
    user = User(id=uuid.uuid4(), email=f"a{uuid.uuid4().hex[:6]}@ex.com", role="admin", organisation_id=org.id)
    db.add(user)
    req = Requisition(
        organisation_id=org.id,
        numero_requisition=f"REQ-{uuid.uuid4().hex[:8]}",
        reference_numero=f"REF-{uuid.uuid4().hex[:8]}",
        objet="Mission Matadi", mode_paiement="cash", type_requisition="classique",
        status="PAYEE", montant_total=Decimal("500"), devise="USD",
        a_valoir=True, instance_beneficiaire="Commission Formation",
        notes_a_valoir="Carburant et per diem",
    )
    db.add(req)
    await db.flush()
    avance = SortieFonds(
        organisation_id=org.id, requisition_id=req.id, type_sortie="autre",
        reference_numero=f"SF-{uuid.uuid4().hex[:6]}", montant_paye=Decimal("500"),
        date_paiement=AVANCE_LE, mode_paiement="cash", devise="USD", canal="CAISSE",
        statut="VALIDE", motif="Avance mission Matadi", beneficiaire="Jean MUKENDI",
        justificatif_statut=justificatif, justificatif_echeance=echeance,
    )
    ordinaire = SortieFonds(
        organisation_id=org.id, type_sortie="autre",
        reference_numero=f"SF-{uuid.uuid4().hex[:6]}", montant_paye=Decimal("40"),
        date_paiement=AVANCE_LE, mode_paiement="cash", devise="USD", canal="CAISSE",
        statut="VALIDE", motif="Fournitures", beneficiaire="Papeterie",
    )
    db.add_all([avance, ordinaire])
    await db.flush()
    db.add(RetourCaisse(
        organisation_id=org.id, sortie_fonds_id=avance.id, type_retour="reliquat_avance",
        reference_numero=f"RT-{uuid.uuid4().hex[:6]}", motif="Reliquat", montant=Decimal("120"),
        devise="USD", canal="CAISSE", mode="cash", date_retour=AVANCE_LE + timedelta(days=3),
        statut="VALIDE",
    ))
    await db.commit()
    return org, user


async def _releve(db, org, user):
    return await journal_tresorerie(
        canal="CAISSE", devise="USD", date_debut="2026-09-01", date_fin="2026-09-30",
        user=user, db=db, tenant_id=org.id,
    )


@pytest.mark.asyncio
async def test_le_journal_signale_l_avance_a_valoir_et_dit_ou_elle_en_est(db_session):
    org, user = await _avance_a_valoir(db_session)
    releve = await _releve(db_session, org, user)

    avance = next(l for l in releve.lignes if l.libelle_base == "Avance mission Matadi")
    ordinaire = next(l for l in releve.lignes if l.libelle_base == "Fournitures")

    assert avance.a_valoir is True
    # Échéance dépassée depuis longtemps : le justificatif est en retard.
    assert avance.a_valoir_statut == "EN_RETARD"
    # 500 remis, 120 rendus : il en reste 380.
    assert avance.a_valoir_reste == Decimal("380")
    assert "Commission Formation" in avance.a_valoir_alerte
    assert "EN RETARD" in avance.a_valoir_alerte
    assert "Reste : 380,00 USD" in avance.a_valoir_alerte
    assert "Carburant et per diem" in avance.a_valoir_alerte

    assert ordinaire.a_valoir is False
    assert ordinaire.a_valoir_alerte is None


@pytest.mark.asyncio
async def test_un_justificatif_recu_ne_laisse_rien_a_justifier(db_session):
    org, user = await _avance_a_valoir(db_session, justificatif="RECU", echeance=None)
    releve = await _releve(db_session, org, user)

    avance = next(l for l in releve.lignes if l.a_valoir)
    assert avance.a_valoir_statut == "RECU"
    assert avance.a_valoir_reste == Decimal("0")


def _journal(lignes):
    return JournalAnnexe(
        nom_compte="Caisse USD", devise="USD", lignes=lignes,
        solde_initial=Decimal("1000"), total_entrees=Decimal("0"),
        total_sorties=Decimal("540"), solde_final=Decimal("460"),
    )


def _ligne(libelle, montant, **a_valoir):
    return SimpleNamespace(
        date=AVANCE_LE, libelle=libelle, libelle_base=libelle, precision=None, reference=None,
        entree=Decimal("0"), sortie=Decimal(montant), solde=Decimal("0"),
        a_valoir=bool(a_valoir), a_valoir_statut=a_valoir.get("statut"),
        a_valoir_reste=a_valoir.get("reste"), a_valoir_alerte=a_valoir.get("alerte"),
    )


def test_l_excel_envoye_colore_l_avance_en_bordeaux_et_porte_une_note():
    from openpyxl import load_workbook

    journal = _journal([
        _ligne("Avance mission", "500", statut="EN_RETARD", reste=Decimal("380"), alerte="AVANCE À VALOIR\nReste : 380"),
        _ligne("Fournitures", "40"),
    ])
    tz = ZoneInfo("UTC")
    contenu = _xlsx(journal, Identite(nom="Org", contact="", logo=None), AVANCE_LE, AVANCE_LE, tz)
    ws = load_workbook(io.BytesIO(contenu)).active

    assert ws.cell(row=4, column=5).value == "À valoir (1)"
    assert ws.cell(row=5, column=5).value == 380
    assert "dont 1 au justificatif en retard" in ws["A6"].value

    avance, ordinaire = ws.cell(row=8, column=3), ws.cell(row=9, column=3)
    assert avance.value == "À VALOIR — Avance mission"
    assert avance.font.color.rgb.endswith("800020")
    assert avance.comment is not None and "Reste : 380" in avance.comment.text
    assert ordinaire.value == "Fournitures"
    assert ordinaire.comment is None


def test_le_pdf_envoye_se_construit_avec_une_avance_a_valoir():
    journal = _journal([_ligne("Avance mission", "500", statut="ATTENDU", reste=Decimal("500"), alerte="x")])
    tz = ZoneInfo("UTC")
    contenu = _pdf(journal, Identite(nom="Org", contact="", logo=None), AVANCE_LE, AVANCE_LE, tz)
    assert contenu.startswith(b"%PDF")
