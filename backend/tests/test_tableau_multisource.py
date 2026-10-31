from __future__ import annotations

import io
import inspect
from datetime import date, datetime

import openpyxl
import pytest
from sqlalchemy import select

from app.modules.tableau.models import (
    TableauDossier,
    TableauCaDeclaration,
    TableauInsuranceDeclaration,
    TableauImport,
    TableauMemberIdentity,
    TableauPersonneMoraleSnapshot,
    TableauPersonnePhysiqueSnapshot,
    TableauSourceRow,
)
from app.modules.tableau.service import build_preparatory_tableau_at_date, ca_present_at_date, get_ca_declarations_at_date, get_insurance_declarations_at_date, get_personne_physique_at_date, import_source_snapshot, insurance_status_at_date
from app.modules.tableau.source_adapters import file_sha256, parse_source_excel, row_sha256
import app.modules.tableau.service as tableau_service


def _xlsx(headers: list[str], rows: list[list[object]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("source_type", "headers", "row"),
    [
        ("personnes_physiques", ["N° d'ordre", "Nom", "Statut"], ["EC/18.00062", "DUPONT Jean", "indépendant"]),
        ("personnes_morales", ["N° d'ordre", "Société", "Ville"], ["SEC/18.00001", "CABINET X", "Kinshasa"]),
        ("chiffres_affaires", ["N° d'ordre", "Année", "Devise", "CA facturé"], ["EC/18.00062", 2025, "USD", 1000]),
        ("assurances", ["Membre", "Déclaré", "Souscrit", "Assureur"], ["SEC/18.00001", "Oui", "Oui", "Assureur X"]),
    ],
)
def test_source_adapter_accepts_each_source_type(source_type, headers, row):
    result = parse_source_excel(_xlsx(headers, [row]), source_type)
    assert len(result.rows) == 1
    assert result.rows[0].raw_data[headers[0]] == row[0]
    assert result.rows[0].business_key_candidate == str(row[0])


def test_source_snapshot_preserves_raw_row_and_hash():
    content = _xlsx(
        ["N° Ordre", "Nom", "Assuré", "CA facturé", "Date naissance", "Vide"],
        [["  EC/18.00062 ", "DUPONT Jean", "OUI", "1 250 000,50", datetime(1980, 1, 2), None]],
    )
    result = parse_source_excel(content, "personnes_physiques")
    raw = result.rows[0].raw_data
    assert raw["N° Ordre"] == "  EC/18.00062 "
    assert raw["Assuré"] == "OUI"
    assert raw["CA facturé"] == "1 250 000,50"
    assert raw["Date naissance"] == "1980-01-02T00:00:00"
    assert raw["Vide"] is None
    assert result.rows[0].normalized_data["numero_ordre"] == "  EC/18.00062 "
    assert len(file_sha256(content)) == 64


def test_file_and_row_fingerprints_are_distinct_concepts():
    content_a = _xlsx(["N° d'ordre", "Nom"], [["EC/18.00062", "DUPONT"]])
    raw = {"N° d'ordre": "EC/18.00062", "Nom": "DUPONT"}
    assert row_sha256(raw) != file_sha256(content_a)
    assert row_sha256(raw) == row_sha256(dict(raw))


@pytest.mark.parametrize("order_header", ["N° Ordre", "N° d'ordre", "N° ordre", "Numéro d'ordre", "N°Ordrre"])
def test_real_order_header_variants_are_recognized_without_overmatching(order_header):
    accepted = parse_source_excel(_xlsx([order_header, "Nom"], [["EC/18.00062", "DUPONT"]]), "personnes_physiques")
    assert len(accepted.rows) == 1
    rejected = parse_source_excel(_xlsx(["Référence libre", "Nom"], [["EC/18.00062", "DUPONT"]]), "personnes_physiques")
    assert rejected.rows == []
    assert any(error["champ"] == "headers" for error in rejected.errors)


def test_modele_membres_pp_cp_maps_all_business_columns():
    headers = [
        "N°Ordrre", "Nom", "Sexe", "Actif", "Date naissance", "Téléphone", "Email", "Province",
        "Statut", "N° impôt", "Assuré", "AMLCO", "% 120 FOR", "%80 FOR",
    ]
    values = [
        "EC/17.00002", "EXEMPLE BOSILA Marie", "F", "oui", "13/12/1974", "(+243)810000001",
        "marie@exemple.cd", "KINSHASA", "indépendant", "A0000001X", "oui", "non", 82, 123,
    ]

    result = parse_source_excel(_xlsx(headers, [values]), "personnes_physiques")

    assert result.errors == []
    assert len(result.rows) == 1
    assert result.rows[0].business_key_candidate == "EC/17.00002"
    assert result.rows[0].normalized_data == {
        "numero_ordre": "EC/17.00002",
        "membre": "EXEMPLE BOSILA Marie",
        "sexe": "F",
        "actif": "oui",
        "date_naissance": "13/12/1974",
        "telephone": "(+243)810000001",
        "email": "marie@exemple.cd",
        "province_attache": "KINSHASA",
        "statut": "indépendant",
        "numero_impot": "A0000001X",
        "assure": "oui",
        "amlco": "non",
        "pourcentage_120_for": 82,
        "pourcentage_80_for": 123,
    }


def test_real_export_column_sets_are_accepted():
    fixtures = {
        "personnes_physiques": [
            "N° Ordre", "Nom", "Sexe", "Actif", "Date naissance", "Téléphone", "Email", "Ville",
            "Statut", "N° impôt", "Assuré", "AMLCO", "% 120 FOR", "% 80 FOR",
        ],
        "personnes_morales": [
            "N° Ordre", "Société", "Téléphone", "Email", "Ville", "Numéro impôt", "Cotisation",
            "Solde", "Assuré", "Nb employés", "Nb EC", "Nb StA",
        ],
        "chiffres_affaires": [
            "N° d'ordre", "Membre", "Actif", "Année", "Devise", "CA facturé", "CA collecté", "Date mise à jour",
        ],
        "assurances": [
            "Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur", "Période",
        ],
    }
    values = {
        "personnes_physiques": ["EC/18.00062", "DUPONT", "M", "Oui", date(1980, 1, 2), "0800000000", "a@b.cd", "Kinshasa", "indépendant", "NIF1", "Oui", "Non", 80, 20],
        "personnes_morales": ["SEC/18.00001", "SOCIETE X", "0800000000", "a@b.cd", "Kinshasa", "NIF2", "Oui", 0, "Oui", 4, 2, 1],
        "chiffres_affaires": ["EC/18.00062", "DUPONT", "Oui", 2025, "USD", "1 250 000,50", "1 000 000", date(2026, 1, 2)],
        "assurances": ["EC/18.00062", "Oui", "Oui", date(2026, 12, 31), 2026, "Assureur X", "annuelle"],
    }
    for source_type, headers in fixtures.items():
        result = parse_source_excel(_xlsx(headers, [values[source_type]]), source_type)
        assert len(result.rows) == 1, (source_type, result.errors)
        assert result.errors == [], (source_type, result.errors)


def test_modele_membres_pm_cp_maps_all_business_columns():
    headers = [
        "N° Ordre", "Société", "Associé gérant", "Téléphone", "Email", "Province", "Sigle",
        "Raison sociale", "Lieu création", "Numéro impôt", "Solde", "Nb empl", "Nb EC", "Nb StA",
    ]
    values = [
        "SEC/18.00001", "EXEMPLE & CO SAS", "EXEMPLE Gerant", "(+243)810000002",
        "contact@exemple.cd", "KINSHASA", "ABN", "Exemple & Co", "Kinshasa", "A0000002X",
        0, 10, 1, 0,
    ]

    result = parse_source_excel(_xlsx(headers, [values]), "personnes_morales")

    assert result.errors == []
    assert len(result.rows) == 1
    assert result.rows[0].normalized_data == {
        "numero_ordre": "SEC/18.00001",
        "membre": "EXEMPLE & CO SAS",
        "associe_gerant": "EXEMPLE Gerant",
        "telephone": "(+243)810000002",
        "email": "contact@exemple.cd",
        "province_attache": "KINSHASA",
        "sigle": "ABN",
        "raison_sociale": "Exemple & Co",
        "lieu_creation": "Kinshasa",
        "numero_impot": "A0000002X",
        "solde": 0,
        "nb_employes": 10,
        "nb_ec": 1,
        "nb_sta": 0,
    }


def test_modele_ca_membres_cp_maps_all_business_columns():
    headers = ["N° d'ordre", "Membre", "Actif", "AN.", "Dev", "CA facturé", "CA collecté", "Date maj"]
    values = ["SEC/18.00001", "EXEMPLE & CO SAS", "oui", 2024, "CDF", 834140885, 765311125, "12/12/2025"]

    result = parse_source_excel(_xlsx(headers, [values]), "chiffres_affaires")

    assert result.errors == []
    assert len(result.rows) == 1
    assert result.rows[0].business_key_candidate == "SEC/18.00001"
    assert result.rows[0].normalized_data == {
        "numero_ordre": "SEC/18.00001",
        "membre": "EXEMPLE & CO SAS",
        "actif": "oui",
        "annee": 2024,
        "devise": "CDF",
        "ca_facture": 834140885,
        "ca_collecte": 765311125,
        "date_mise_a_jour": "12/12/2025",
    }


def test_modele_declarations_assurance_gestionnaire_cp_maps_all_business_columns():
    headers = ["Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur", "Période"]
    values = [
        "SEC/20.00005 EXEMPLE CONSULTANTS SARL", "09/08/2024", "29/05/2024", "21/05/2025",
        2024, "MAYFAIR Insurance Congo SA", "2024 - 2025",
    ]

    result = parse_source_excel(_xlsx(headers, [values]), "assurances")

    assert result.errors == []
    assert len(result.rows) == 1
    assert result.rows[0].normalized_data == {
        "membre": "SEC/20.00005 EXEMPLE CONSULTANTS SARL",
        "declare": "09/08/2024",
        "souscrit": "29/05/2024",
        "fin_couverture": "21/05/2025",
        "annee_souscription": 2024,
        "assureur": "MAYFAIR Insurance Congo SA",
        "periode": "2024 - 2025",
    }


@pytest.mark.asyncio
async def test_snapshot_import_is_idempotent_for_same_file_and_situation(db_session, test_admin_user):
    content = _xlsx(["N° d'ordre", "Nom"], [["EC/18.00062", "DUPONT Jean"]])
    kwargs = dict(
        db=db_session,
        user=test_admin_user,
        organisation_id=test_admin_user.organisation_id,
        file_name="physiques.xlsx",
        content=content,
        source_type="personnes_physiques",
        exercice="2026",
        date_situation=date(2026, 6, 1),
    )
    first = await import_source_snapshot(**kwargs)
    second = await import_source_snapshot(**kwargs)
    assert first.duplicate_detected is False
    assert second.duplicate_detected is True
    rows = (await db_session.execute(select(TableauSourceRow).where(TableauSourceRow.import_id == first.imp.id))).scalars().all()
    assert len(rows) == 1
    assert rows[0].raw_data["Nom"] == "DUPONT Jean"
    dossiers = (await db_session.execute(select(TableauDossier).where(TableauDossier.import_id == first.imp.id))).scalars().all()
    assert dossiers == []


@pytest.mark.asyncio
async def test_same_file_new_situation_creates_new_snapshot(db_session, test_admin_user):
    content = _xlsx(["N° d'ordre", "Nom"], [["EC/18.00063", "MUKENDI Paul"]])
    common = dict(db=db_session, user=test_admin_user, organisation_id=test_admin_user.organisation_id,
                  file_name="physiques-2.xlsx", content=content, source_type="personnes_physiques", exercice="2026")
    first = await import_source_snapshot(**common, date_situation=date(2026, 6, 2))
    second = await import_source_snapshot(**common, date_situation=date(2026, 7, 2))
    assert first.imp.id != second.imp.id
    assert second.duplicate_detected is False


@pytest.mark.asyncio
async def test_same_bytes_different_source_type_are_distinct(db_session, test_admin_user):
    content = _xlsx(["N° d'ordre", "Nom", "Année"], [["EC/18.00064", "KABAMBA", 2025]])
    common = dict(db=db_session, user=test_admin_user, organisation_id=test_admin_user.organisation_id,
                  file_name="same-bytes.xlsx", content=content, exercice="2026", date_situation=date(2026, 6, 3))
    first = await import_source_snapshot(**common, source_type="personnes_physiques")
    second = await import_source_snapshot(**common, source_type="personnes_morales")
    assert first.duplicate_detected is False
    assert second.duplicate_detected is False
    assert first.imp.id != second.imp.id


@pytest.mark.asyncio
async def test_same_file_different_organisation_is_distinct(db_session, test_admin_user, test_organisation):
    from app.models.organisation import Organisation

    other = Organisation(nom="Autre organisation", slug="tableau-multisource-other", is_active=True)
    db_session.add(other)
    await db_session.flush()
    content = _xlsx(["N° d'ordre", "Nom"], [["EC/18.00065", "KABAMBA"]])
    common = dict(db=db_session, user=test_admin_user, file_name="org.xlsx", content=content,
                  source_type="personnes_physiques", exercice="2026", date_situation=date(2026, 6, 4))
    first = await import_source_snapshot(organisation_id=test_admin_user.organisation_id, **common)
    second = await import_source_snapshot(organisation_id=other.id, **common)
    assert first.duplicate_detected is False
    assert second.duplicate_detected is False


@pytest.mark.asyncio
async def test_pp_creates_internal_ec_identity_and_snapshot_without_tableau_dossier(db_session, test_admin_user):
    content = _xlsx(["N° Ordre", "Nom", "Statut", "Actif"], [["EC/18.00070", "DUPONT Jean", "indépendant", "Oui"]])
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "pp.xlsx", content,
        "personnes_physiques", "2026", date(2026, 9, 1),
    )
    identity = (await db_session.execute(select(TableauMemberIdentity).where(TableauMemberIdentity.numero_ordre == "EC/18.00070"))).scalar_one()
    snapshot = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.source_import_id == outcome.imp.id))).scalar_one()
    assert identity.member_kind == "EC"
    assert identity.person_kind == "PHYSIQUE"
    assert snapshot.identity_id == identity.id
    assert snapshot.statut_normalise == "INDEPENDANT"
    assert (await db_session.execute(select(TableauDossier).where(TableauDossier.import_id == outcome.imp.id))).scalars().all() == []


@pytest.mark.asyncio
async def test_modele_membres_pp_cp_materializes_the_expected_snapshot(db_session, test_admin_user):
    headers = [
        "N°Ordrre", "Nom", "Sexe", "Actif", "Date naissance", "Téléphone", "Email", "Province",
        "Statut", "N° impôt", "Assuré", "AMLCO", "% 120 FOR", "%80 FOR",
    ]
    values = [
        "EC/25.00604", "EXEMPLE MALU Jean", "M", "oui", "14/09/1992", "0810000003",
        "jean@exemple.cd", "KINSHASA", "indépendant", "E0000003R", "oui", "non", 153, 229,
    ]
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "Membres PP CP.xlsx",
        _xlsx(headers, [values]), "personnes_physiques", "2026", date(2026, 10, 31),
    )

    snapshot = (await db_session.execute(
        select(TableauPersonnePhysiqueSnapshot).where(
            TableauPersonnePhysiqueSnapshot.source_import_id == outcome.imp.id
        )
    )).scalar_one()
    assert snapshot.numero_ordre == "EC/25.00604"
    assert snapshot.nom == "EXEMPLE MALU Jean"
    assert snapshot.actif is True
    assert snapshot.date_naissance == date(1992, 9, 14)
    assert snapshot.province_attache == "KINSHASA"
    assert snapshot.numero_impot == "E0000003R"
    assert snapshot.nif == "E0000003R"
    assert snapshot.assure_synthetique is True
    assert snapshot.amlco is False
    assert str(snapshot.pourcentage_120_for) == "153.00"
    assert str(snapshot.pourcentage_80_for) == "229.00"


@pytest.mark.asyncio
async def test_modele_membres_pp_cp_ne_signale_que_les_vraies_anomalies(db_session, test_admin_user):
    headers = ["N°Ordrre", "Nom", "Date naissance", "Statut", "N° impôt", "Assuré", "AMLCO"]
    rows = [
        # Salarié : N/A et date nulle du portail, sans colonne Employeur dans le modèle.
        ["EC/16.99101", "EXEMPLE Salarie", "30/11/-0001", "salarié", None, "N/A", "N/A"],
        ["EC/16.99102", "EXEMPLE Cabinet", "15/06/1980", "en cabinet", None, "oui", "non"],
        ["EC/16.99103", "EXEMPLE Sans Nif", "31/02/1980", "indépendant", None, "oui", "non"],
    ]
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "Membres PP CP.xlsx",
        _xlsx(headers, rows), "personnes_physiques", "2026", date(2026, 10, 31),
    )

    source_rows = (await db_session.execute(
        select(TableauSourceRow).where(TableauSourceRow.import_id == outcome.imp.id).order_by(TableauSourceRow.line_number)
    )).scalars().all()
    assert [row.normalization_status for row in source_rows] == ["accepted", "accepted", "warning"]
    salarie = (await db_session.execute(
        select(TableauPersonnePhysiqueSnapshot).where(
            TableauPersonnePhysiqueSnapshot.source_import_id == outcome.imp.id,
            TableauPersonnePhysiqueSnapshot.numero_ordre == "EC/16.99101",
        )
    )).scalar_one()
    assert salarie.date_naissance is None
    assert salarie.assure_synthetique is None
    assert salarie.amlco is None
    # Le rapport d'import annonce ce que portent les lignes.
    assert {(error["ligne"], error["champ"]) for error in outcome.errors} == {(4, "date_naissance"), (4, "nif")}


@pytest.mark.asyncio
async def test_employeur_vide_reste_signale_quand_la_colonne_existe(db_session, test_admin_user):
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "pp-employeur.xlsx",
        _xlsx(["N° Ordre", "Nom", "Statut", "Employeur"], [["EC/16.99104", "EXEMPLE Salarie", "salarié", None]]),
        "personnes_physiques", "2026", date(2026, 10, 31),
    )
    assert [error["champ"] for error in outcome.errors] == ["nom_employeur"]


@pytest.mark.asyncio
async def test_pm_creates_internal_sec_identity(db_session, test_admin_user):
    content = _xlsx(["N° Ordre", "Société", "Nb employés", "Nb EC", "Nb StA"], [["SEC/18.00071", "SOCIETE X", 5, 2, 1]])
    await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "pm.xlsx", content,
        "personnes_morales", "2026", date(2026, 9, 1),
    )
    identity = (await db_session.execute(select(TableauMemberIdentity).where(TableauMemberIdentity.numero_ordre == "SEC/18.00071"))).scalar_one()
    snapshot = (await db_session.execute(select(TableauPersonneMoraleSnapshot).where(TableauPersonneMoraleSnapshot.numero_ordre == "SEC/18.00071"))).scalar_one()
    assert identity.member_kind == "SEC"
    assert identity.person_kind == "MORALE"
    assert snapshot.nb_ec == 2


@pytest.mark.asyncio
async def test_modele_membres_pm_cp_materializes_the_expected_snapshot(db_session, test_admin_user):
    headers = [
        "N° Ordre", "Société", "Associé gérant", "Téléphone", "Email", "Province", "Sigle",
        "Raison sociale", "Lieu création", "Numéro impôt", "Solde", "Nb empl", "Nb EC", "Nb StA",
    ]
    values = [
        "SEC/26.00999", "CABINET TEST SARL", "GERANT Test", "+243810000000", "test@cabinet.cd",
        "KINSHASA", "CTS", "CABINET TEST", "Kinshasa", "A0000001X", 1500, 12, 2, 3,
    ]
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "Membres PM CP.xlsx",
        _xlsx(headers, [values]), "personnes_morales", "2026", date(2026, 10, 31),
    )

    snapshot = (await db_session.execute(
        select(TableauPersonneMoraleSnapshot).where(
            TableauPersonneMoraleSnapshot.source_import_id == outcome.imp.id
        )
    )).scalar_one()
    assert snapshot.numero_ordre == "SEC/26.00999"
    assert snapshot.societe == "CABINET TEST SARL"
    assert snapshot.associe_gerant == "GERANT Test"
    assert snapshot.sigle == "CTS"
    assert snapshot.raison_sociale == "CABINET TEST"
    assert snapshot.lieu_creation == "Kinshasa"
    assert snapshot.province_attache == "KINSHASA"
    assert snapshot.numero_impot == "A0000001X"
    assert str(snapshot.solde) == "1500.00"
    assert (snapshot.nb_employes, snapshot.nb_ec, snapshot.nb_sta) == (12, 2, 3)


@pytest.mark.asyncio
async def test_identity_is_reused_and_snapshots_are_append_only(db_session, test_admin_user):
    first_content = _xlsx(["N° Ordre", "Nom", "Téléphone"], [["EC/18.00072", "DUPONT", "A"]])
    second_content = _xlsx(["N° Ordre", "Nom", "Téléphone"], [["EC/18.00072", "DUPONT", "B"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "a.xlsx", first_content, "personnes_physiques", "2026", date(2026, 9, 1))
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "b.xlsx", second_content, "personnes_physiques", "2026", date(2026, 9, 12))
    identities = (await db_session.execute(select(TableauMemberIdentity).where(TableauMemberIdentity.numero_ordre == "EC/18.00072"))).scalars().all()
    snapshots = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.numero_ordre == "EC/18.00072").order_by(TableauPersonnePhysiqueSnapshot.date_situation))).scalars().all()
    assert len(identities) == 1
    assert [snapshot.telephone for snapshot in snapshots] == ["A", "B"]


@pytest.mark.asyncio
async def test_snapshot_can_be_reconstructed_at_a_situation_date(db_session, test_admin_user):
    for situation, phone in ((date(2026, 9, 1), "A"), (date(2026, 9, 12), "B")):
        content = _xlsx(["N° Ordre", "Nom", "Téléphone"], [["EC/18.00073", "DUPONT", phone]])
        await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, f"{phone}.xlsx", content, "personnes_physiques", "2026", situation)
    snapshot = await get_personne_physique_at_date(db_session, test_admin_user.organisation_id, " ec/18.00073 ", date(2026, 9, 5))
    assert snapshot is not None and snapshot.telephone == "A"


@pytest.mark.asyncio
async def test_incompatible_prefix_is_anomaly_without_conversion(db_session, test_admin_user):
    content = _xlsx(["N° Ordre", "Nom"], [["SEC/18.00074", "SOCIETE X"]])
    outcome = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "bad-pp.xlsx", content, "personnes_physiques", "2026", date(2026, 9, 1))
    row = (await db_session.execute(select(TableauSourceRow).where(TableauSourceRow.import_id == outcome.imp.id))).scalar_one()
    assert row.business_key_candidate == "SEC/18.00074"
    assert row.normalization_status == "warning"
    assert "SEC/18.00074" in {i.numero_ordre for i in (await db_session.execute(select(TableauMemberIdentity))).scalars().all()}
    assert (await db_session.execute(select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.source_row_id == row.id))).scalars().all()


@pytest.mark.asyncio
async def test_absence_in_later_partial_import_does_not_deactivate_member(db_session, test_admin_user):
    content = _xlsx(["N° Ordre", "Nom", "Actif"], [["EC/18.00075", "DUPONT", "Oui"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "initial.xlsx", content, "personnes_physiques", "2026", date(2026, 9, 1))
    empty = _xlsx(["N° Ordre", "Nom", "Actif"], [])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "partial.xlsx", empty, "personnes_physiques", "2026", date(2026, 9, 12))
    identity = (await db_session.execute(select(TableauMemberIdentity).where(TableauMemberIdentity.numero_ordre == "EC/18.00075"))).scalar_one()
    assert identity.classification_status == "classified"


@pytest.mark.asyncio
async def test_unknown_values_remain_null_and_are_not_defaulted(db_session, test_admin_user):
    content = _xlsx(["N° Ordre", "Nom", "Actif", "AMLCO", "% 120 FOR"], [["EC/18.00076", "DUPONT", "PEUT-ETRE", "?", "inconnu"]])
    outcome = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "unknown.xlsx", content, "personnes_physiques", "2026", date(2026, 9, 1))
    snapshot = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.source_import_id == outcome.imp.id))).scalar_one()
    assert snapshot.actif is None and snapshot.amlco is None and snapshot.pourcentage_120_for is None
    row = (await db_session.execute(select(TableauSourceRow).where(TableauSourceRow.import_id == outcome.imp.id))).scalar_one()
    assert row.normalization_status == "warning"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("statut", "header", "value", "field", "expected"),
    [
        ("en cabinet", "Cabinet d'attache", "SEC/22.00088", "cabinet_attache", "SEC/22.00088"),
        ("indépendant", "NIF", "A1502382C", "nif", "A1502382C"),
        ("salarié", "Nom de l'employeur", "Ministère des Finances", "nom_employeur", "Ministère des Finances"),
    ],
)
async def test_pp_contextual_status_fields_are_snapshot_values(db_session, test_admin_user, statut, header, value, field, expected):
    content = _xlsx(["N° Ordre", "Nom", "Statut", header], [["EC/18.00077", "DUPONT", statut, value]])
    outcome = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, f"contextual-{field}.xlsx", content, "personnes_physiques", "2026", date(2026, 9, 1))
    snapshot = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.source_import_id == outcome.imp.id))).scalar_one()
    assert getattr(snapshot, field) == expected
    assert snapshot.nif is None if field != "nif" else True


@pytest.mark.asyncio
async def test_pp_status_change_keeps_identity_and_old_snapshot(db_session, test_admin_user):
    for name, situation, statut, header, value in (
        ("ind.xlsx", date(2026, 9, 1), "indépendant", "NIF", "A1502382C"),
        ("cab.xlsx", date(2026, 9, 12), "en cabinet", "Cabinet d'attache", "SEC/22.00088"),
    ):
        await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, name, _xlsx(["N° Ordre", "Nom", "Statut", header], [["EC/18.00078", "DUPONT", statut, value]]), "personnes_physiques", "2026", situation)
    identity = (await db_session.execute(select(TableauMemberIdentity).where(TableauMemberIdentity.numero_ordre == "EC/18.00078"))).scalar_one()
    snapshots = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.numero_ordre == "EC/18.00078").order_by(TableauPersonnePhysiqueSnapshot.date_situation))).scalars().all()
    assert len(snapshots) == 2 and snapshots[0].nif == "A1502382C" and snapshots[1].cabinet_attache == "SEC/22.00088"
    assert identity.member_kind == "EC" and identity.person_kind == "PHYSIQUE"


@pytest.mark.asyncio
async def test_pp_contextual_mismatch_is_warning_and_raw_is_preserved(db_session, test_admin_user):
    content = _xlsx(["N° Ordre", "Nom", "Statut", "NIF"], [["EC/18.00079", "DUPONT", "indépendant", "SEC/22.00088"]])
    outcome = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "bad-context.xlsx", content, "personnes_physiques", "2026", date(2026, 9, 1))
    snapshot = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.source_import_id == outcome.imp.id))).scalar_one()
    row = (await db_session.execute(select(TableauSourceRow).where(TableauSourceRow.import_id == outcome.imp.id))).scalar_one()
    assert snapshot.nif is None
    assert row.raw_data["NIF"] == "SEC/22.00088"
    assert "SEC" in str(row.normalization_errors)


@pytest.mark.asyncio
async def test_insurance_ec_and_sec_are_rattachees_without_dossier(db_session, test_admin_user):
    pp = _xlsx(["N° Ordre", "Nom"], [["EC/18.00090", "DUPONT"]])
    pm = _xlsx(["N° Ordre", "Société"], [["SEC/18.00091", "SOCIETE"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ass-pp.xlsx", pp, "personnes_physiques", "2026", date(2026, 9, 1))
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ass-pm.xlsx", pm, "personnes_morales", "2026", date(2026, 9, 1))
    content = _xlsx(["Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur", "Période"], [["EC/18.00090", "Oui", "Oui", date(2027, 12, 31), 2026, "  Assureur X  ", "annuelle"], ["SEC/18.00091", "Oui", "Non", date(2025, 12, 31), 2026, "Assureur Y", "annuelle"]])
    outcome = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "assurances.xlsx", content, "assurances", "2026", date(2026, 9, 2))
    rows = (await db_session.execute(select(TableauInsuranceDeclaration).where(TableauInsuranceDeclaration.source_import_id == outcome.imp.id).order_by(TableauInsuranceDeclaration.id))).scalars().all()
    assert [row.member_identity_id is not None for row in rows] == [True, True]
    assert rows[0].declare_normalise is True and rows[0].souscrit_normalise is True
    assert rows[0].assureur_normalise == "assureur x" and rows[0].periode_normalisee == "annuelle"
    assert (await db_session.execute(select(TableauDossier).where(TableauDossier.import_id == outcome.imp.id))).scalars().all() == []


@pytest.mark.asyncio
async def test_modele_declarations_assurance_gestionnaire_cp_materializes_dates(db_session, test_admin_user):
    numero = "SEC/20.00005"
    await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "pm-assurance-modele.xlsx",
        _xlsx(["N° Ordre", "Société"], [[numero, "EXEMPLE CONSULTANTS SARL"]]),
        "personnes_morales", "2026", date(2026, 10, 30),
    )
    headers = ["Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur", "Période"]
    values = [
        f"{numero} EXEMPLE CONSULTANTS SARL", "09/08/2024", "29/05/2024", "21/05/2025",
        2024, "MAYFAIR Insurance Congo SA", "2024 - 2025",
    ]
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id,
        "Déclarations assurance gestionnaire CP.xlsx", _xlsx(headers, [values]),
        "assurances", "2026", date(2026, 10, 31),
    )

    declaration = (await db_session.execute(
        select(TableauInsuranceDeclaration).where(TableauInsuranceDeclaration.source_import_id == outcome.imp.id)
    )).scalar_one()
    source_row = (await db_session.execute(
        select(TableauSourceRow).where(TableauSourceRow.import_id == outcome.imp.id)
    )).scalar_one()
    assert declaration.member_identity_id is not None
    assert declaration.numero_ordre_normalise == numero
    assert declaration.declare_normalise is True
    assert declaration.date_declaration == date(2024, 8, 9)
    assert declaration.souscrit_normalise is True
    assert declaration.date_souscription == date(2024, 5, 29)
    assert declaration.fin_couverture == date(2025, 5, 21)
    assert declaration.annee_souscription == 2024
    assert declaration.assureur_normalise == "mayfair insurance congo sa"
    assert declaration.periode_normalisee == "2024 - 2025"
    assert source_row.normalization_status == "accepted"


@pytest.mark.asyncio
async def test_insurance_orphan_and_multiple_history_are_preserved(db_session, test_admin_user):
    pp = _xlsx(["N° Ordre", "Nom"], [["EC/18.00092", "DUPONT"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ass-pp2.xlsx", pp, "personnes_physiques", "2026", date(2026, 9, 1))
    headers = ["Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur"]
    first = _xlsx(headers, [["EC/18.00092", "Oui", "Oui", date(2026, 12, 31), 2026, "A"]])
    second = _xlsx(headers, [["EC/18.00092", "Oui", "Oui", date(2027, 12, 31), 2026, "B"], ["INCONNU", "?", "?", "bad", "bad", ""]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ass-1.xlsx", first, "assurances", "2026", date(2026, 9, 2))
    out = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ass-2.xlsx", second, "assurances", "2026", date(2026, 9, 10))
    rows = (await db_session.execute(select(TableauInsuranceDeclaration).where(TableauInsuranceDeclaration.numero_ordre_normalise == "EC/18.00092"))).scalars().all()
    orphan = (await db_session.execute(select(TableauInsuranceDeclaration).where(TableauInsuranceDeclaration.source_import_id == out.imp.id, TableauInsuranceDeclaration.numero_ordre_normalise.is_(None)))).scalar_one()
    assert len(rows) == 2 and orphan.member_identity_id is None
    assert len(await get_insurance_declarations_at_date(db_session, test_admin_user.organisation_id, "EC/18.00092", date(2026, 9, 10))) == 2


@pytest.mark.asyncio
async def test_insurance_tri_state_temporal_status(db_session, test_admin_user):
    pp = _xlsx(["N° Ordre", "Nom"], [["EC/18.00093", "DUPONT"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ass-pp3.xlsx", pp, "personnes_physiques", "2026", date(2026, 9, 1))
    headers = ["Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur"]
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "future.xlsx", _xlsx(headers, [["EC/18.00093", "Oui", "Oui", date(2027, 12, 31), 2026, "A"]]), "assurances", "2026", date(2026, 9, 2))
    assert await insurance_status_at_date(db_session, test_admin_user.organisation_id, "EC/18.00093", date(2026, 10, 1)) == "TRUE"
    assert await insurance_status_at_date(db_session, test_admin_user.organisation_id, "EC/18.00093", date(2028, 1, 1)) == "FALSE"


@pytest.mark.asyncio
async def test_preparatory_projection_is_independent_and_exposes_source_indicators(db_session, test_admin_user):
    pp = _xlsx(["N° Ordre", "Nom", "Statut", "Actif", "Assuré"], [["EC/18.00094", "DUPONT", "en cabinet", "Non", "Oui"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "proj-pp.xlsx", pp, "personnes_physiques", "2026", date(2026, 9, 1))
    ca = _xlsx(["N° d'ordre", "Membre", "Année", "Devise", "CA facturé"], [["EC/18.00094", "DUPONT", 2025, "USD", "100"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "proj-ca.xlsx", ca, "chiffres_affaires", "2026", date(2026, 9, 2))
    projection = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 3))
    row = next(item for item in projection.rows if item.numero_ordre == "EC/18.00094")
    assert row.tableau_category == "EC Cabinet" and row.actif is False
    assert row.ca_present is True and row.ca_declaration_count == 1
    assert row.insurance_status == "UNKNOWN" and row.assure_source is True


@pytest.mark.asyncio
async def test_projection_temporal_selection_and_future_exclusion(db_session, test_admin_user):
    for name, situation, phone in (("temporal-a.xlsx", date(2026, 9, 1), "A"), ("temporal-b.xlsx", date(2026, 9, 12), "B")):
        await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, name, _xlsx(["N° Ordre", "Nom", "Téléphone", "Statut"], [["EC/18.00100", "DUPONT", phone, "indépendant"]]), "personnes_physiques", "2026", situation)
    before = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 10))
    after = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 15))
    assert next(r for r in before.rows if r.numero_ordre == "EC/18.00100").source_dates["personne"] == date(2026, 9, 1)
    assert next(r for r in after.rows if r.numero_ordre == "EC/18.00100").source_dates["personne"] == date(2026, 9, 12)


@pytest.mark.asyncio
async def test_projection_same_date_is_deterministic_and_reports_ambiguity(db_session, test_admin_user):
    for name, phone in (("same-a.xlsx", "A"), ("same-b.xlsx", "B")):
        await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, name, _xlsx(["N° Ordre", "Nom", "Téléphone", "Statut"], [["EC/18.00101", "DUPONT", phone, "indépendant"]]), "personnes_physiques", "2026", date(2026, 9, 5))
    projection = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 5))
    row = next(r for r in projection.rows if r.numero_ordre == "EC/18.00101")
    assert "AMBIGUOUS_SNAPSHOT_DATE" in row.anomaly_codes
    assert row.projection_status == "WARNING"


@pytest.mark.asyncio
async def test_projection_insurance_mismatch_tri_state_and_orphans(db_session, test_admin_user):
    pp = _xlsx(["N° Ordre", "Nom", "Assuré"], [["EC/18.00102", "DUPONT", "Non"], ["EC/18.00103", "MARTIN", "Oui"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "mismatch-pp.xlsx", pp, "personnes_physiques", "2026", date(2026, 9, 1))
    headers = ["Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur"]
    insurance = _xlsx(headers, [["EC/18.00102", "Oui", "Oui", date(2027, 12, 31), 2026, "A"], ["UNKNOWN", "?", "?", None, 2026, "" ]])
    out = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "mismatch-ass.xlsx", insurance, "assurances", "2026", date(2026, 9, 2))
    projection = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 3))
    row = next(r for r in projection.rows if r.numero_ordre == "EC/18.00102")
    assert row.insurance_status == "TRUE" and row.assure_source is False and row.insurance_source_mismatch is True
    assert len(projection.orphan_insurance) == 1 and projection.orphan_insurance[0].source_import_id == out.imp.id
    unknown = next(r for r in projection.rows if r.numero_ordre == "EC/18.00103")
    assert unknown.insurance_status == "UNKNOWN" and unknown.insurance_source_mismatch is False


@pytest.mark.asyncio
async def test_projection_keeps_inactive_and_orphan_ca_outside_member_rows(db_session, test_admin_user):
    pp = _xlsx(["N° Ordre", "Nom", "Actif", "Statut"], [["EC/18.00104", "DUPONT", "Non", "indépendant"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "inactive.xlsx", pp, "personnes_physiques", "2026", date(2026, 9, 1))
    ca = _xlsx(["N° d'ordre", "Membre", "Année", "Devise", "CA facturé"], [["UNKNOWN", "X", 2026, "USD", "10"]])
    out = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "orphan-ca-proj.xlsx", ca, "chiffres_affaires", "2026", date(2026, 9, 2))
    projection = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 3))
    row = next(r for r in projection.rows if r.numero_ordre == "EC/18.00104")
    assert row.actif is False
    assert len(projection.orphan_ca) == 1 and projection.orphan_ca[0].source_import_id == out.imp.id


@pytest.mark.asyncio
async def test_projection_reports_identity_without_snapshot(db_session, test_admin_user):
    seed = _xlsx(["N° Ordre", "Nom"], [["EC/18.00105", "DUPONT"]])
    out = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "seed-no-snapshot.xlsx", seed, "personnes_physiques", "2026", date(2026, 9, 1))
    import_ref = (await db_session.execute(select(TableauImport).where(TableauImport.id == out.imp.id))).scalar_one()
    identity = TableauMemberIdentity(organisation_id=test_admin_user.organisation_id, numero_ordre="EC/18.00106", numero_ordre_normalise="EC/18.00106", member_kind="EC", person_kind="PHYSIQUE", first_seen_import_id=import_ref.id, last_seen_import_id=import_ref.id)
    db_session.add(identity)
    await db_session.commit()
    projection = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 2))
    row = next(r for r in projection.rows if r.numero_ordre == "EC/18.00106")
    assert "IDENTITY_WITHOUT_SNAPSHOT" in row.anomaly_codes and row.projection_status == "BLOCKED"


def test_projection_function_has_no_official_or_historical_dependency():
    source = inspect.getsource(tableau_service.build_preparatory_tableau_at_date)
    assert all(token not in source for token in ("ExpertComptable", "experts_comptables", "TableauDossier", "CategoryChangesHistory", "ImportsHistory"))


@pytest.mark.asyncio
async def test_projection_ec_independent_category_and_context(db_session, test_admin_user):
    out = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "proj-ind.xlsx", _xlsx(["N° Ordre", "Nom", "Statut", "NIF"], [["EC/18.00110", "DUPONT", "indépendant", "A1502382C"]]), "personnes_physiques", "2026", date(2026, 9, 1))
    projection = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 2))
    identity_id = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot.identity_id).where(TableauPersonnePhysiqueSnapshot.source_import_id == out.imp.id))).scalar_one()
    row = next(r for r in projection.rows if r.identity_id == identity_id)
    assert (row.member_kind, row.person_kind, row.tableau_category) == ("EC", "PHYSIQUE", "EC Indépendant")
    assert row.nif == "A1502382C" and row.cabinet_attache is None and row.nom_employeur is None


@pytest.mark.asyncio
async def test_projection_ec_salarie_category_and_context(db_session, test_admin_user):
    out = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "proj-sal.xlsx", _xlsx(["N° Ordre", "Nom", "Statut", "Nom de l'employeur"], [["EC/18.00111", "MARTIN", "salarié", "FPI"]]), "personnes_physiques", "2026", date(2026, 9, 1))
    projection = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 2))
    identity_id = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot.identity_id).where(TableauPersonnePhysiqueSnapshot.source_import_id == out.imp.id))).scalar_one()
    row = next(r for r in projection.rows if r.identity_id == identity_id)
    assert (row.member_kind, row.person_kind, row.tableau_category) == ("EC", "PHYSIQUE", "EC Salarié")
    assert row.nom_employeur == "FPI" and row.cabinet_attache is None and row.nif is None


@pytest.mark.asyncio
async def test_projection_sec_category_is_presentation_only(db_session, test_admin_user):
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "proj-sec.xlsx", _xlsx(["N° Ordre", "Société"], [["SEC/18.00112", "SOCIETE"]]), "personnes_morales", "2026", date(2026, 9, 1))
    row = next(r for r in (await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 2))).rows if r.numero_ordre == "SEC/18.00112")
    assert (row.member_kind, row.person_kind, row.tableau_category, row.numero_ordre) == ("SEC", "MORALE", "Société", "SEC/18.00112")


@pytest.mark.asyncio
async def test_projection_detects_ec_with_pm_snapshot(db_session, test_admin_user):
    out = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "conflict-pm.xlsx", _xlsx(["N° Ordre", "Société"], [["SEC/18.00113", "SOCIETE"]]), "personnes_morales", "2026", date(2026, 9, 1))
    identity = TableauMemberIdentity(organisation_id=test_admin_user.organisation_id, numero_ordre="EC/18.00114", numero_ordre_normalise="EC/18.00114", member_kind="EC", person_kind="PHYSIQUE", first_seen_import_id=out.imp.id, last_seen_import_id=out.imp.id)
    db_session.add(identity); await db_session.flush()
    snapshot = (await db_session.execute(select(TableauPersonneMoraleSnapshot).where(TableauPersonneMoraleSnapshot.source_import_id == out.imp.id))).scalar_one()
    snapshot.identity_id = identity.id; snapshot.numero_ordre = identity.numero_ordre; await db_session.commit()
    row = next(r for r in (await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 2))).rows if r.identity_id == identity.id)
    assert "EC_WITH_PM_SNAPSHOT" in row.anomaly_codes and row.projection_status == "BLOCKED" and row.member_kind == "EC"


@pytest.mark.asyncio
async def test_projection_detects_sec_with_pp_snapshot(db_session, test_admin_user):
    out = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "conflict-pp.xlsx", _xlsx(["N° Ordre", "Nom", "Statut"], [["EC/18.00115", "DUPONT", "indépendant"]]), "personnes_physiques", "2026", date(2026, 9, 1))
    identity = TableauMemberIdentity(organisation_id=test_admin_user.organisation_id, numero_ordre="SEC/18.00116", numero_ordre_normalise="SEC/18.00116", member_kind="SEC", person_kind="MORALE", first_seen_import_id=out.imp.id, last_seen_import_id=out.imp.id)
    db_session.add(identity); await db_session.flush()
    snapshot = (await db_session.execute(select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.source_import_id == out.imp.id))).scalar_one()
    snapshot.identity_id = identity.id; snapshot.numero_ordre = identity.numero_ordre; await db_session.commit()
    row = next(r for r in (await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 2))).rows if r.identity_id == identity.id)
    assert "SEC_WITH_PP_SNAPSHOT" in row.anomaly_codes and row.projection_status == "BLOCKED" and row.member_kind == "SEC" and row.tableau_category == "Société"


@pytest.mark.asyncio
async def test_projection_isolation_between_organisations(db_session, test_admin_user):
    from app.models.organisation import Organisation
    other = Organisation(nom="Projection Other", slug="projection-other", is_active=True)
    db_session.add(other); await db_session.flush()
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "iso-a.xlsx", _xlsx(["N° Ordre", "Nom", "Statut"], [["EC/18.00117", "A", "indépendant"]]), "personnes_physiques", "2026", date(2026, 9, 1))
    await import_source_snapshot(db_session, test_admin_user, other.id, "iso-b.xlsx", _xlsx(["N° Ordre", "Nom", "Statut"], [["EC/18.00117", "B", "salarié"]]), "personnes_physiques", "2026", date(2026, 9, 1))
    a = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 2)); b = await build_preparatory_tableau_at_date(db_session, other.id, date(2026, 9, 2))
    assert [r.nom_denomination for r in a.rows if r.numero_ordre == "EC/18.00117"] == ["A"]
    assert [r.nom_denomination for r in b.rows if r.numero_ordre == "EC/18.00117"] == ["B"]
    assert all(r.organisation_id == test_admin_user.organisation_id for r in a.rows) and all(r.organisation_id == other.id for r in b.rows)


@pytest.mark.asyncio
async def test_projection_ca_future_and_multiple_currencies(db_session, test_admin_user):
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ca-proj-pp.xlsx", _xlsx(["N° Ordre", "Nom"], [["EC/18.00118", "DUPONT"]]), "personnes_physiques", "2026", date(2026, 9, 1))
    headers = ["N° d'ordre", "Membre", "Année", "Devise", "CA facturé"]
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ca-cdf.xlsx", _xlsx(headers, [["EC/18.00118", "DUPONT", 2025, "CDF", "10"]]), "chiffres_affaires", "2026", date(2026, 9, 1))
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ca-usd.xlsx", _xlsx(headers, [["EC/18.00118", "DUPONT", 2025, "USD", "20"]]), "chiffres_affaires", "2026", date(2026, 9, 20))
    before = next(r for r in (await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 10))).rows if r.numero_ordre == "EC/18.00118")
    after = next(r for r in (await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 21))).rows if r.numero_ordre == "EC/18.00118")
    assert before.ca_declaration_count == 1 and before.derniere_date_situation_ca == date(2026, 9, 1)
    assert after.ca_declaration_count == 2 and set(after.devises_ca_connues) == {"CDF", "USD"}


@pytest.mark.asyncio
async def test_projection_explicit_insurance_true_and_false(db_session, test_admin_user):
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ins-proj-pp.xlsx", _xlsx(["N° Ordre", "Nom"], [["EC/18.00119", "A"], ["EC/18.00120", "B"]]), "personnes_physiques", "2026", date(2026, 9, 1))
    headers = ["Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur"]
    content = _xlsx(headers, [["EC/18.00119", "Oui", "Oui", date(2027, 1, 1), 2026, "A"], ["EC/18.00120", "Oui", "Oui", date(2026, 1, 1), 2026, "B"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ins-proj.xlsx", content, "assurances", "2026", date(2026, 9, 2))
    projection = await build_preparatory_tableau_at_date(db_session, test_admin_user.organisation_id, date(2026, 9, 10))
    assert next(r for r in projection.rows if r.numero_ordre == "EC/18.00119").insurance_status == "TRUE"
    assert next(r for r in projection.rows if r.numero_ordre == "EC/18.00120").insurance_status == "FALSE"


@pytest.mark.asyncio
async def test_ca_ec_is_rattache_et_decimal_preserved(db_session, test_admin_user):
    pp = _xlsx(["N° Ordre", "Nom"], [["EC/18.00080", "DUPONT"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "pp-ca.xlsx", pp, "personnes_physiques", "2026", date(2026, 9, 1))
    content = _xlsx(["N° d'ordre", "Membre", "Actif", "Année", "Devise", "CA facturé", "CA collecté", "Date mise à jour"], [["EC/18.00080", "DUPONT", "Oui", 2025, "usd", "1 250 000,50", "1 000 000", date(2026, 9, 2)]])
    outcome = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ca.xlsx", content, "chiffres_affaires", "2026", date(2026, 9, 3))
    row = (await db_session.execute(select(TableauCaDeclaration).where(TableauCaDeclaration.source_import_id == outcome.imp.id))).scalar_one()
    assert row.member_identity_id is not None
    assert row.ca_facture == 1250000.50 and row.ca_collecte == 1000000
    assert row.devise_normalisee == "USD" and row.annee == 2025 and row.ca_present is True


@pytest.mark.asyncio
async def test_modele_ca_membres_cp_materializes_the_expected_declaration(db_session, test_admin_user):
    numero = "EC/26.00998"
    await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "pp-ca-modele.xlsx",
        _xlsx(["N° Ordre", "Nom"], [[numero, "MEMBRE TEST"]]),
        "personnes_physiques", "2026", date(2026, 10, 30),
    )
    headers = ["N° d'ordre", "Membre", "Actif", "AN.", "Dev", "CA facturé", "CA collecté", "Date maj"]
    values = [numero, "MEMBRE TEST", "oui", 2024, "USD", 8360, 7000, "18/08/2025"]
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "CA membres CP.xlsx",
        _xlsx(headers, [values]), "chiffres_affaires", "2026", date(2026, 10, 31),
    )

    declaration = (await db_session.execute(
        select(TableauCaDeclaration).where(TableauCaDeclaration.source_import_id == outcome.imp.id)
    )).scalar_one()
    assert declaration.member_identity_id is not None
    assert declaration.actif_normalise is True
    assert declaration.annee == 2024
    assert declaration.devise_normalisee == "USD"
    assert str(declaration.ca_facture) == "8360.00"
    assert str(declaration.ca_collecte) == "7000.00"
    assert declaration.date_mise_a_jour == date(2025, 8, 18)
    assert declaration.ca_present is True


@pytest.mark.asyncio
async def test_ca_sec_orphan_and_invalid_number_do_not_create_identity(db_session, test_admin_user):
    content = _xlsx(["N° d'ordre", "Membre", "Année", "Devise", "CA facturé"], [["SEC/18.00081", "SOCIETE", 2025, "EUR", "10,00"], ["", "ORPHELIN", "bad", "???", "x"]])
    outcome = await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "ca-orphans.xlsx", content, "chiffres_affaires", "2026", date(2026, 9, 4))
    declarations = (await db_session.execute(select(TableauCaDeclaration).where(TableauCaDeclaration.source_import_id == outcome.imp.id).order_by(TableauCaDeclaration.id))).scalars().all()
    assert len(declarations) == 2
    assert declarations[0].member_identity_id is None and declarations[0].numero_ordre_normalise == "SEC/18.00081"
    assert declarations[1].member_identity_id is None and declarations[1].annee is None
    assert (await db_session.execute(select(TableauMemberIdentity).where(TableauMemberIdentity.numero_ordre == "SEC/18.00081"))).scalars().all() == []


@pytest.mark.asyncio
async def test_ca_history_duplicates_and_temporal_projection(db_session, test_admin_user):
    pp = _xlsx(["N° Ordre", "Nom"], [["EC/18.00082", "DUPONT"]])
    await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, "pp-ca2.xlsx", pp, "personnes_physiques", "2026", date(2026, 9, 1))
    headers = ["N° d'ordre", "Membre", "Année", "Devise", "CA facturé", "CA collecté"]
    for name, situation, amount, rows in (("ca-a.xlsx", date(2026, 9, 5), "1000", 1), ("ca-b.xlsx", date(2026, 9, 10), "1200", 2)):
        values = [["EC/18.00082", "DUPONT", 2025, "USD", amount, amount] for _ in range(rows)]
        await import_source_snapshot(db_session, test_admin_user, test_admin_user.organisation_id, name, _xlsx(headers, values), "chiffres_affaires", "2026", situation)
    observations = await get_ca_declarations_at_date(db_session, test_admin_user.organisation_id, "EC/18.00082", date(2026, 9, 10))
    assert len(observations) == 3
    assert await ca_present_at_date(db_session, test_admin_user.organisation_id, "EC/18.00082", date(2026, 9, 10)) is True
    source_rows = (await db_session.execute(select(TableauSourceRow).where(TableauSourceRow.import_id == observations[-1].source_import_id))).scalars().all()
    assert any("Doublon exact" in str(row.normalization_errors) for row in source_rows)


def test_tableau_import_has_no_official_expert_dependency():
    assert not any("experts_comptables" in str(f.target_fullname) for f in TableauImport.__table__.foreign_keys)
    assert not any("experts_comptables" in str(f.target_fullname) for f in TableauSourceRow.__table__.foreign_keys)


def test_multisource_pipeline_has_no_official_or_legacy_import_side_effects():
    source = inspect.getsource(tableau_service.import_source_snapshot)
    assert "ExpertComptable" not in source
    assert "CategoryChangesHistory" not in source
    assert "ImportsHistory" not in source
    assert "TableauDossier(" not in source


def _fills_rouges(sheet, ligne: int) -> set[str]:
    return {cell.fill.fgColor.rgb[-6:] for cell in sheet[ligne] if cell.fill.fill_type == "solid"}


@pytest.mark.asyncio
async def test_aucune_information_manquante_n_ecarte_une_ligne_pp(db_session, test_admin_user):
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "pp-incomplet.xlsx",
        _xlsx(
            ["N°Ordrre", "Nom", "Statut", "N° impôt"],
            [
                ["EC/16.99201", "COMPLET Jean", "indépendant", "A1234567X"],
                ["EC/16.99202", "SANS NIF Marie", "indépendant", None],
                [None, "SANS NUMERO Paul", "indépendant", "A7654321X"],
                ["EC/16.99204", None, "indépendant", "A1111111X"],
            ],
        ),
        "personnes_physiques", "2026", date(2026, 10, 31),
    )
    assert outcome.imp.rejected_rows == 0
    assert outcome.imported == 4
    snapshots = (await db_session.execute(
        select(TableauPersonnePhysiqueSnapshot).where(TableauPersonnePhysiqueSnapshot.source_import_id == outcome.imp.id)
    )).scalars().all()
    assert len(snapshots) == 4
    sans_numero = next(s for s in snapshots if s.nom == "SANS NUMERO Paul")
    assert sans_numero.identity_id is None and sans_numero.numero_ordre == ""
    assert {(e["ligne"], e["champ"]) for e in outcome.errors} == {(3, "nif"), (4, "numero_ordre"), (5, "nom")}


@pytest.mark.asyncio
async def test_une_collision_pp_pm_garde_la_ligne_sans_la_rattacher(db_session, test_admin_user):
    organisation_id = test_admin_user.organisation_id
    await import_source_snapshot(
        db_session, test_admin_user, organisation_id, "pp-collision.xlsx",
        _xlsx(["N° Ordre", "Nom"], [["EC/16.99301", "PERSONNE Jean"]]),
        "personnes_physiques", "2026", date(2026, 10, 30),
    )
    outcome = await import_source_snapshot(
        db_session, test_admin_user, organisation_id, "pm-collision.xlsx",
        _xlsx(["N° Ordre", "Société", "Numéro impôt"], [["EC/16.99301", "SOCIETE SA", "A1X"]]),
        "personnes_morales", "2026", date(2026, 10, 31),
    )
    snapshot = (await db_session.execute(
        select(TableauPersonneMoraleSnapshot).where(TableauPersonneMoraleSnapshot.source_import_id == outcome.imp.id)
    )).scalar_one()
    assert snapshot.identity_id is None
    assert any(e["message"] == "Collision de classification PP/PM" for e in outcome.errors)


@pytest.mark.asyncio
async def test_le_fichier_annote_met_la_ligne_en_rouge_et_se_reimporte(db_session, test_admin_user):
    organisation_id = test_admin_user.organisation_id
    outcome = await import_source_snapshot(
        db_session, test_admin_user, organisation_id, "Membres PP CP.xlsx",
        _xlsx(
            ["N°Ordrre", "Nom", "Date naissance", "Statut", "N° impôt"],
            [
                ["EC/16.99401", "COMPLET Jean", datetime(1980, 6, 15), "indépendant", "A1234567X"],
                ["EC/16.99402", "SANS NIF Marie", "01/01/1975", "indépendant", None],
            ],
        ),
        "personnes_physiques", "2026", date(2026, 10, 31),
    )
    content, nom = await tableau_service.export_import_annote(db_session, organisation_id, outcome.imp.id)
    assert nom == "Membres_PP_CP_annote.xlsx"

    wb = openpyxl.load_workbook(io.BytesIO(content))
    sheet = wb["Sheet"]
    entetes = [cell.value for cell in sheet[1]]
    # L'ordre du fichier d'origine, et le numéro de ligne annoncé par le rapport.
    assert entetes == ["N°Ordrre", "Nom", "Date naissance", "Statut", "N° impôt", "Signalement"]
    assert sheet.cell(2, 3).value == "15/06/1980"
    assert sheet.cell(3, 2).value == "SANS NIF Marie"
    assert sheet.cell(3, 6).value == "Indépendant sans NIF"
    assert _fills_rouges(sheet, 2) == set()
    assert _fills_rouges(sheet, 3) == {"FFC7CE", "FF7C80"}
    assert sheet.cell(3, 5).fill.fgColor.rgb[-6:] == "FF7C80"
    assert [c.value for c in wb["Signalements"][2]] == ["Sheet", 3, "nif", "Indépendant sans NIF"]

    # Le fichier annoté, une fois complété, repasse par l'import.
    sheet.cell(3, 5).value = "A7654321X"
    buffer = io.BytesIO()
    wb.save(buffer)
    reprise = await import_source_snapshot(
        db_session, test_admin_user, organisation_id, nom, buffer.getvalue(),
        "personnes_physiques", "2026", date(2026, 10, 31),
    )
    assert reprise.imported == 2
    assert reprise.errors == []

    # Un second cycle ne cumule pas les colonnes techniques d'annotation.
    content_repris, _ = await tableau_service.export_import_annote(db_session, organisation_id, reprise.imp.id)
    wb_repris = openpyxl.load_workbook(io.BytesIO(content_repris))
    assert [cell.value for cell in wb_repris["Sheet"][1]].count("Signalement") == 1


@pytest.mark.asyncio
async def test_fichier_annote_d_un_autre_conseil_introuvable(db_session, test_admin_user):
    outcome = await import_source_snapshot(
        db_session, test_admin_user, test_admin_user.organisation_id, "pp.xlsx",
        _xlsx(["N° Ordre", "Nom"], [["EC/16.99501", "CLOISON Jean"]]),
        "personnes_physiques", "2026", date(2026, 10, 31),
    )
    with pytest.raises(Exception) as excinfo:
        await tableau_service.export_import_annote(db_session, test_admin_user.organisation_id + 999, outcome.imp.id)
    assert getattr(excinfo.value, "status_code", None) == 404


@pytest.mark.asyncio
async def test_classeur_complet_signale_a_la_bonne_ligne_et_garde_le_membre_sans_nom(db_session, test_admin_user):
    wb = openpyxl.Workbook()
    wb.active.title = "EC EN CABINET"
    cabinet = wb.active
    cabinet.append(["SECTION A1"])
    cabinet.append(["N°", "N° d'ordre", "Nom, post-noms et prénoms", "Cabinet d'attache", "Cotisation"])
    cabinet.append([1, "EC/16.99601", "COMPLET Jean", "CABINET X", "OUI"])
    independants = wb.create_sheet("EC INDEPENDANTS")
    independants.append(["SECTION A2"])
    independants.append(["N°", "N° d'ordre", "Nom, post-noms et prénoms", "NIF", "Cotisation"])
    independants.append([1, "EC/16.99602", "AVEC NIF Marie", "A1X", "OUI"])
    independants.append([2, "EC/16.99603", "SANS NIF Paul", None, "OUI"])
    independants.append([3, "EC/16.99604", None, "A2X", "OUI"])
    independants.append([4, None, None, None, None])
    buffer = io.BytesIO()
    wb.save(buffer)

    outcome = await tableau_service.import_excel(
        db_session, test_admin_user, test_admin_user.organisation_id, "tableau.xlsx", buffer.getvalue(),
        "2026", date_situation=date(2026, 10, 31),
    )
    # La ligne numérotée mais vide n'est pas un membre ; le membre sans nom en est un.
    assert outcome.imported == 4
    assert {(e["feuille"], e["ligne"], e["champ"]) for e in outcome.errors} == {
        ("EC INDEPENDANTS", 4, "nif"),
        ("EC INDEPENDANTS", 5, "nom"),
    }

    content, _nom = await tableau_service.export_import_annote(db_session, test_admin_user.organisation_id, outcome.imp.id)
    annote = openpyxl.load_workbook(io.BytesIO(content))
    assert _fills_rouges(annote["EC EN CABINET"], 3) == set()
    feuille = annote["EC INDEPENDANTS"]
    assert feuille.cell(4, 3).value == "SANS NIF Paul"
    assert "FFC7CE" in _fills_rouges(feuille, 4)
    assert _fills_rouges(feuille, 3) == set()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("source_type", "headers", "complete", "incomplete", "signales"),
    [
        (
            "personnes_morales",
            ["N° Ordre", "Société", "Numéro impôt"],
            ["SEC/18.99701", "COMPLETE SA", "A1X"],
            ["SEC/18.99702", "SANS NIF SA", None],
            {"numero_impot": "Société sans NIF"},
        ),
        (
            "chiffres_affaires",
            ["N° d'ordre", "Membre", "AN.", "Dev", "CA facturé", "CA collecté"],
            ["EC/17.99703", "MEMBRE", 2025, "USD", 100, 90],
            ["EC/17.99704", None, None, None, None, None],
            {
                "membre": "Membre manquant", "annee": "Année manquante", "devise": "Devise manquante",
                "ca_facture": "CA facturé manquant", "ca_collecte": "CA collecté manquant",
            },
        ),
        (
            "assurances",
            ["Membre", "Déclaré", "Souscrit", "Fin couverture", "Année de souscription", "Assureur"],
            ["EC/17.99705 COMPLET", "09/08/2024", "29/05/2024", "21/05/2025", 2024, "ASSUREUR X"],
            ["EC/17.99706 INCOMPLET", None, "29/05/2024", None, None, None],
            {
                "declare": "Déclaré manquant", "fin_couverture": "Fin de couverture manquante",
                "annee_souscription": "Année de souscription manquante", "assureur": "Souscrit sans assureur",
            },
        ),
    ],
)
async def test_chaque_source_garde_signale_et_annote_la_ligne_incomplete(
    db_session, test_admin_user, source_type, headers, complete, incomplete, signales,
):
    organisation_id = test_admin_user.organisation_id
    # Le membre existe : le rapprochement n'ajoute pas son propre signalement.
    for numero in (str(complete[0]).split()[0], str(incomplete[0]).split()[0]):
        if source_type != "personnes_morales":
            await import_source_snapshot(
                db_session, test_admin_user, organisation_id, f"pp-{numero[-5:]}.xlsx",
                _xlsx(["N° Ordre", "Nom"], [[numero, "MEMBRE"]]),
                "personnes_physiques", "2026", date(2026, 10, 1),
            )
    outcome = await import_source_snapshot(
        db_session, test_admin_user, organisation_id, f"{source_type}.xlsx",
        _xlsx(headers, [complete, incomplete]), source_type, "2026", date(2026, 10, 31),
    )
    assert outcome.imported == 2
    assert outcome.imp.rejected_rows == 0
    assert {e["champ"]: e["message"] for e in outcome.errors} == signales
    assert {e["ligne"] for e in outcome.errors} == {3}

    content, nom = await tableau_service.export_import_annote(db_session, organisation_id, outcome.imp.id)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    sheet = wb["Sheet"]
    assert [cell.value for cell in sheet[1]] == [*headers, "Signalement"]
    assert _fills_rouges(sheet, 2) == set()
    assert "FFC7CE" in _fills_rouges(sheet, 3) or "FF7C80" in _fills_rouges(sheet, 3)
    assert sheet.cell(3, 1).value == incomplete[0]
    for message in signales.values():
        assert message in sheet.cell(3, len(headers) + 1).value

    # Le fichier annoté repasse par l'import du même type.
    buffer = io.BytesIO()
    wb.save(buffer)
    reprise = await import_source_snapshot(
        db_session, test_admin_user, organisation_id, nom, buffer.getvalue(), source_type, "2026", date(2026, 11, 1),
    )
    assert reprise.imported == 2


async def _pp(db_session, user, nom_fichier, lignes, situation=date(2026, 10, 31)):
    return await import_source_snapshot(
        db_session, user, user.organisation_id, nom_fichier,
        _xlsx(["N° Ordre", "Nom", "Statut", "N° impôt"], lignes), "personnes_physiques", "2026", situation,
    )


async def _base(db_session, user, **kwargs):
    return await tableau_service.get_base_tableau(db_session, [user.organisation_id], exercice="2026", limit=50, **kwargs)


@pytest.mark.asyncio
async def test_la_base_montre_les_membres_importes_sans_conclusion(db_session, admin_isole):
    await _pp(db_session, admin_isole, "pp-base.xlsx", [
        ["EC/16.99801", "BASE Jean", "indépendant", "A1X"],
        ["EC/16.99802", "BASE Marie", "indépendant", None],
    ])
    await import_source_snapshot(
        db_session, admin_isole, admin_isole.organisation_id, "pm-base.xlsx",
        _xlsx(["N° Ordre", "Société", "Numéro impôt"], [["SEC/18.99803", "BASE SA", "A2X"]]),
        "personnes_morales", "2026", date(2026, 10, 31),
    )
    base = await _base(db_session, admin_isole)

    assert base["dossiers"] == []
    assert base["total_dossiers"] == 0
    assert base["total_membres"] == base["total_a_completer"] == 3
    par_numero = {m["numero_ordre"]: m for m in base["membres_a_completer"]}
    assert par_numero["EC/16.99801"]["categorie"] == "EC Indépendant"
    assert par_numero["EC/16.99801"]["signalements"] == []
    assert par_numero["EC/16.99802"]["signalements"] == ["Indépendant sans NIF"]
    assert par_numero["SEC/18.99803"]["categorie"] == "Société"
    assert "conclusion" not in par_numero["EC/16.99801"]
    # Aucun dossier n'a été créé : l'analyse n'a rien à juger.
    assert (await db_session.execute(select(TableauDossier).where(TableauDossier.organisation_id == admin_isole.organisation_id))).scalars().all() == []


@pytest.mark.asyncio
async def test_le_dossier_du_classeur_prime_sur_la_ligne_a_completer(db_session, admin_isole):
    await _pp(db_session, admin_isole, "pp-prime.xlsx", [
        ["EC/16.99811", "AVEC DOSSIER Jean", "en cabinet", None],
        ["EC/16.99812", "SANS DOSSIER Marie", "en cabinet", None],
    ])
    wb = openpyxl.Workbook()
    wb.active.title = "EC EN CABINET"
    wb.active.append(["N°", "N° d'ordre", "Nom, post-noms et prénoms", "Cabinet d'attache", "Cotisation"])
    wb.active.append([1, "EC/16.99811", "AVEC DOSSIER Jean", "CABINET X", "OUI"])
    buffer = io.BytesIO()
    wb.save(buffer)
    await tableau_service.import_excel(
        db_session, admin_isole, admin_isole.organisation_id, "classeur.xlsx", buffer.getvalue(),
        "2026", date_situation=date(2026, 10, 31),
    )
    base = await _base(db_session, admin_isole)
    assert [d.numero_ordre for d in base["dossiers"]] == ["EC/16.99811"]
    assert [m["numero_ordre"] for m in base["membres_a_completer"]] == ["EC/16.99812"]
    assert base["total_membres"] == 2


@pytest.mark.asyncio
async def test_un_reimport_ne_duplique_pas_les_membres_a_completer(db_session, admin_isole):
    await _pp(db_session, admin_isole, "pp-v1.xlsx", [
        ["EC/16.99821", "ANCIEN NOM", "salarié", None],
        [None, "SANS NUMERO", "salarié", None],
    ], date(2026, 10, 1))
    await _pp(db_session, admin_isole, "pp-v2.xlsx", [
        ["EC/16.99821", "NOUVEAU NOM", "salarié", None],
        [None, "SANS NUMERO", "salarié", None],
    ], date(2026, 10, 31))
    base = await _base(db_session, admin_isole)
    assert sorted(m["nom"] for m in base["membres_a_completer"]) == ["NOUVEAU NOM", "SANS NUMERO"]
    assert base["membres_sans_numero"] == 1


@pytest.mark.asyncio
async def test_pagination_et_filtres_des_membres_a_completer(db_session, admin_isole):
    await _pp(db_session, admin_isole, "pp-pages.xlsx", [
        [f"EC/16.9983{i}", f"PAGE {i}", "en cabinet" if i < 2 else "salarié", None] for i in range(4)
    ])
    page = await tableau_service.get_base_tableau(
        db_session, [admin_isole.organisation_id], exercice="2026", limit=3, offset=3,
    )
    assert [m["nom"] for m in page["membres_a_completer"]] == ["PAGE 3"]
    assert page["total_membres"] == 4
    assert (await _base(db_session, admin_isole, categorie="EC Salarié"))["total_a_completer"] == 2
    assert [m["nom"] for m in (await _base(db_session, admin_isole, recherche="page 1"))["membres_a_completer"]] == ["PAGE 1"]
    assert (await _base(db_session, admin_isole, anomalie_only=True))["total_a_completer"] == 0
    autre = await tableau_service.get_base_tableau(db_session, [admin_isole.organisation_id + 999], exercice="2026", limit=50)
    assert autre["total_a_completer"] == 0


@pytest.mark.asyncio
async def test_le_filtre_societe_retrouve_aussi_les_anciens_dossiers_sec(db_session, admin_isole):
    wb = openpyxl.Workbook()
    wb.active.title = "SOCIETES"
    wb.active.append(["N°", "N° d'ordre", "Dénomination", "NIF", "Cotisation"])
    wb.active.append([1, "SEC/18.99841", "ANCIENNE SEC SA", "A1X", "OUI"])
    wb.active.append([2, "SEC/18.99842", "NOUVELLE SOCIETE SA", "A2X", "OUI"])
    buffer = io.BytesIO()
    wb.save(buffer)
    outcome = await tableau_service.import_excel(
        db_session, admin_isole, admin_isole.organisation_id, "societes.xlsx", buffer.getvalue(),
        "2026", date_situation=date(2026, 10, 31),
    )
    ancien = (await db_session.execute(
        select(TableauDossier).where(TableauDossier.import_id == outcome.imp.id, TableauDossier.numero_ordre == "SEC/18.99841")
    )).scalar_one()
    ancien.categorie = "SEC"
    await db_session.commit()

    base = await _base(db_session, admin_isole, categorie="Société")
    assert sorted(d.numero_ordre for d in base["dossiers"]) == ["SEC/18.99841", "SEC/18.99842"]


async def _membres_filtres(db_session, admin_isole):
    org = admin_isole.organisation_id
    await import_source_snapshot(
        db_session, admin_isole, org, "pp-filtres.xlsx",
        _xlsx(
            ["N° Ordre", "Nom", "Sexe", "Province", "Statut"],
            [
                ["EC/17.99851", "ANCIEN Jean", "M", "KINSHASA", "en cabinet"],
                ["EC/25.99852", "ANCIENNE Marie", "F", "kinshasa ", "salarié"],
                ["EC/26.99853", "NOUVEAU Paul", "Masculin", "HAUT-KATANGA", "indépendant"],
            ],
        ),
        "personnes_physiques", "2026", date(2026, 10, 31),
    )
    await import_source_snapshot(
        db_session, admin_isole, org, "pm-filtres.xlsx",
        _xlsx(["N° Ordre", "Société", "Province", "Numéro impôt"], [["SEC/18.99854", "CABINET SA", "KINSHASA", "A1X"]]),
        "personnes_morales", "2026", date(2026, 10, 31),
    )


@pytest.mark.asyncio
async def test_filtres_personne_province_sexe_et_anciennete(db_session, admin_isole):
    await _membres_filtres(db_session, admin_isole)

    async def noms(**filtres):
        base = await _base(db_session, admin_isole, **filtres)
        return sorted(m["nom"] for m in base["membres_a_completer"])

    assert await noms(personne="morale") == ["CABINET SA"]
    assert await noms(personne="physique") == ["ANCIEN Jean", "ANCIENNE Marie", "NOUVEAU Paul"]
    assert await noms(province="kinshasa") == ["ANCIEN Jean", "ANCIENNE Marie", "CABINET SA"]
    assert await noms(sexe="F") == ["ANCIENNE Marie"]
    assert await noms(sexe="M") == ["ANCIEN Jean", "NOUVEAU Paul"]
    assert await noms(anciennete="Nouveau") == ["NOUVEAU Paul"]
    assert await noms(personne="physique", province="KINSHASA", anciennete="Ancien") == ["ANCIEN Jean", "ANCIENNE Marie"]
    base = await _base(db_session, admin_isole)
    assert base["province_options"] == ["HAUT-KATANGA", "KINSHASA"]


@pytest.mark.asyncio
async def test_l_anciennete_suit_l_exercice_sans_annee_figee(db_session, admin_isole):
    await _membres_filtres(db_session, admin_isole)
    base = await _base(db_session, admin_isole)
    libelles = {m["numero_ordre"]: (m["anciennete_libelle"], m["annee_inscription"]) for m in base["membres_a_completer"]}
    assert libelles["EC/17.99851"] == ("Ancien", 2017)
    assert libelles["EC/25.99852"] == ("Ancien", 2025)
    assert libelles["EC/26.99853"] == ("Nouveau", 2026)

    # L'exercice suivant : les inscrits de 2026 deviennent des anciens, et ceux de 2027 les nouveaux.
    assert tableau_service.verdict_engine.libelle_anciennete("EC/26.99853", 2027) == "Ancien"
    assert tableau_service.verdict_engine.libelle_anciennete("EC/27.00001", 2027) == "Nouveau"
    assert tableau_service.verdict_engine.libelle_anciennete("ANCIEN-FORMAT", 2027) is None


@pytest.mark.asyncio
async def test_le_numero_d_ordre_prime_sur_l_anciennete_du_fichier_et_les_filtres_valent_pour_les_dossiers(db_session, admin_isole):
    org = admin_isole.organisation_id
    await import_source_snapshot(
        db_session, admin_isole, org, "pp-province.xlsx",
        _xlsx(["N° Ordre", "Nom", "Province", "Statut"], [["EC/17.99861", "DOSSIER Jean", "KINSHASA", "en cabinet"]]),
        "personnes_physiques", "2026", date(2026, 10, 30),
    )
    wb = openpyxl.Workbook()
    wb.active.title = "EC EN CABINET"
    wb.active.append(["N°", "N° d'ordre", "Nom, post-noms et prénoms", "Sexe", "Ancienneté", "Cabinet d'attache"])
    wb.active.append([1, "EC/17.99861", "DOSSIER Jean", "M", "Nouveau", "CABINET X"])
    wb.active.append([2, "EC/26.99862", "RECENT Anne", "F", "Ancien", "CABINET X"])
    wb.active.append([3, "ECS-ILLISIBLE", "LIBRE Paul", "M", "Nouveau", "CABINET X"])
    societes = wb.create_sheet("SOCIETES")
    societes.append(["N°", "N° d'ordre", "Dénomination", "NIF"])
    societes.append([1, "SEC/18.99863", "SOCIETE SA", "A1X"])
    buffer = io.BytesIO()
    wb.save(buffer)
    await tableau_service.import_excel(
        db_session, admin_isole, org, "classeur-filtres.xlsx", buffer.getvalue(), "2026", date_situation=date(2026, 10, 31),
    )

    base = await _base(db_session, admin_isole)
    par_numero = {d.numero_ordre: d for d in base["dossiers"]}
    assert par_numero["EC/17.99861"].anciennete == "Ancien"
    assert par_numero["EC/17.99861"].raw_data["anciennete_source"] == "Nouveau"
    assert par_numero["EC/17.99861"].province == "KINSHASA"
    assert par_numero["EC/26.99862"].anciennete == "Nouveau"
    assert par_numero["ECS-ILLISIBLE"].anciennete == "Nouveau"

    async def numeros(**filtres):
        return sorted(d.numero_ordre for d in (await _base(db_session, admin_isole, **filtres))["dossiers"])

    assert await numeros(personne="morale") == ["SEC/18.99863"]
    assert "SEC/18.99863" not in await numeros(personne="physique")
    assert await numeros(sexe="F") == ["EC/26.99862"]
    assert await numeros(province="KINSHASA") == ["EC/17.99861"]
    assert await numeros(anciennete="Nouveau") == ["EC/26.99862", "ECS-ILLISIBLE"]
    assert (await _base(db_session, admin_isole, sexe="F"))["total_membres"] == 1


@pytest.mark.asyncio
async def test_un_numero_posterieur_a_l_exercice_est_a_verifier(db_session, admin_isole):
    moteur = tableau_service.verdict_engine
    assert moteur.libelle_anciennete("SEC/99.99999", 2026) == moteur.ANCIENNETE_A_VERIFIER
    assert moteur.est_nouveau({"numero_ordre": "EC/99.99999"}, 2026) is False
    # Le réglage d'exemption reste, lui, inchangé : 2024 à 2026 pour trois ans.
    assert moteur.est_nouveau({"numero_ordre": "EC/24.00001"}, 2026) is True
    assert moteur.est_nouveau({"numero_ordre": "EC/23.00001"}, 2026) is False

    outcome = await import_source_snapshot(
        db_session, admin_isole, admin_isole.organisation_id, "pm-provisoire.xlsx",
        _xlsx(["N° Ordre", "Société", "Numéro impôt"], [["SEC/99.99999", "PROVISOIRE SA", "A1X"]]),
        "personnes_morales", "2026", date(2026, 10, 31),
    )
    assert [e["message"] for e in outcome.errors] == ["Année du N° d'ordre postérieure à l'exercice"]
    base = await _base(db_session, admin_isole)
    assert base["membres_a_completer"][0]["anciennete_libelle"] == "À vérifier"
    assert (await _base(db_session, admin_isole, anciennete="Nouveau"))["total_a_completer"] == 0
    assert (await _base(db_session, admin_isole, anciennete="Ancien"))["total_a_completer"] == 0


@pytest.mark.asyncio
async def test_le_nhv_de_la_base_reprend_la_colonne_120_for(db_session, admin_isole):
    org = admin_isole.organisation_id
    await import_source_snapshot(
        db_session, admin_isole, org, "pp-nhv.xlsx",
        _xlsx(
            ["N°Ordrre", "Nom", "Statut", "% 120 FOR", "%80 FOR"],
            [
                ["EC/19.99871", "AVEC NHV", "en cabinet", 82, 123],
                ["EC/19.99872", "NHV ZERO", "en cabinet", 0, 0],
                ["EC/19.99873", "SANS NHV", "en cabinet", None, None],
            ],
        ),
        "personnes_physiques", "2026", date(2026, 10, 31),
    )
    await import_source_snapshot(
        db_session, admin_isole, org, "pm-nhv.xlsx",
        _xlsx(["N° Ordre", "Société", "Numéro impôt"], [["SEC/18.99874", "SOCIETE NHV", "A1X"]]),
        "personnes_morales", "2026", date(2026, 10, 31),
    )
    base = await _base(db_session, admin_isole)
    nhv = {m["nom"]: m["nhv"] for m in base["membres_a_completer"]}
    assert nhv == {"AVEC NHV": 82.0, "NHV ZERO": 0.0, "SANS NHV": None, "SOCIETE NHV": None}


@pytest.mark.asyncio
async def test_la_base_montre_le_statut_actif(db_session, admin_isole):
    org = admin_isole.organisation_id
    await import_source_snapshot(
        db_session, admin_isole, org, "pp-actif.xlsx",
        _xlsx(
            ["N° Ordre", "Nom", "Actif", "Statut"],
            [
                ["EC/19.99881", "ACTIF Jean", "oui", "en cabinet"],
                ["EC/19.99882", "INACTIF Marie", "non", "en cabinet"],
                ["EC/19.99883", "DOSSIER Paul", "non", "en cabinet"],
            ],
        ),
        "personnes_physiques", "2026", date(2026, 10, 30),
    )
    await import_source_snapshot(
        db_session, admin_isole, org, "pm-actif.xlsx",
        _xlsx(["N° Ordre", "Société", "Numéro impôt"], [["SEC/18.99884", "SOCIETE SA", "A1X"]]),
        "personnes_morales", "2026", date(2026, 10, 30),
    )
    wb = openpyxl.Workbook()
    wb.active.title = "EC EN CABINET"
    wb.active.append(["N°", "N° d'ordre", "Nom, post-noms et prénoms", "Cabinet d'attache"])
    wb.active.append([1, "EC/19.99883", "DOSSIER Paul", "CABINET X"])
    buffer = io.BytesIO()
    wb.save(buffer)
    await tableau_service.import_excel(
        db_session, admin_isole, org, "classeur-actif.xlsx", buffer.getvalue(), "2026", date_situation=date(2026, 10, 31),
    )
    base = await _base(db_session, admin_isole)
    actifs = {m["nom"]: m["actif"] for m in base["membres_a_completer"]}
    assert actifs == {"ACTIF Jean": True, "INACTIF Marie": False, "SOCIETE SA": None}
    assert [d.actif for d in base["dossiers"]] == [False]
