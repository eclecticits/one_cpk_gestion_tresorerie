"""Rapprochement bancaire : versements, répartition, récapitulatifs, annulation.

Complète `test_comptabilite_exports_banque.py` sur les cas réels qui le
faisaient échouer : une note payée en plusieurs virements, deux mouvements de
même montant le même jour, un relevé avec ses lignes de solde, une décision
prise par erreur, un journal de contrôle aux dates avec fuseau horaire.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

import app.db.session  # noqa: F401 — active le cloisonnement tenant
from app.core.tenant_context import set_current_tenant_id
from app.models.encaissement import Encaissement
from app.models.payment_history import PaymentHistory
from app.modules.comptabilite.services.export_service import creer_controle_bancaire
from app.modules.comptabilite.services.rapprochement_bancaire import (
    decider_rapprochement,
    importer_releve,
    parser_releve,
    releve_out,
)
from tests.test_comptabilite_exports_banque import _org, _setup

LE_20 = datetime(2026, 10, 20, 9, 0, tzinfo=timezone.utc)
LE_22 = datetime(2026, 10, 22, 9, 0, tzinfo=timezone.utc)


@pytest.fixture
async def banque(db_session):
    set_current_tenant_id(None)
    org = _org()
    db_session.add(org)
    await db_session.flush()
    set_current_tenant_id(org.id)
    *_, compte_bancaire = await _setup(db_session, org.id)
    yield db_session, org, compte_bancaire
    set_current_tenant_id(None)


async def _note(db, org, compte_bancaire, *versements: tuple[Decimal, datetime], client="Membre"):
    total = sum(montant for montant, _ in versements)
    note = Encaissement(
        organisation_id=org.id,
        numero_recu=f"ND-{uuid.uuid4().hex[:6]}",
        type_client="personne_physique",
        client_nom=client,
        libelle="Cotisation",
        montant=total,
        montant_total=total,
        montant_paye=total,
        montant_percu=total,
        devise_perception="USD",
        taux_change_applique=Decimal("1"),
        canal="BANQUE",
        compte_bancaire_id=compte_bancaire.id,
        statut_paiement="complet",
        statut_operation="ACTIVE",
        mode_paiement="virement",
        date_encaissement=versements[0][1],
    )
    db.add(note)
    await db.flush()
    for montant, le in versements:
        db.add(
            PaymentHistory(
                organisation_id=org.id,
                encaissement_id=note.id,
                montant=montant,
                devise="USD",
                canal="BANQUE",
                compte_bancaire_id=compte_bancaire.id,
                mode_paiement="virement",
                date_paiement=le,
                statut="ACTIF",
            )
        )
    await db.flush()
    return note


async def _importer(db, org, compte_bancaire, csv: str, nom="releve.csv"):
    releve = await importer_releve(
        db,
        organisation_id=org.id,
        user_id=None,
        compte_bancaire_id=compte_bancaire.id,
        filename=nom,
        content=csv.encode("utf-8"),
    )
    return await releve_out(db, releve)


async def _decider(db, org, ligne, decision, proposition=None):
    return await decider_rapprochement(
        db,
        organisation_id=org.id,
        user_id=None,
        ligne_id=ligne.id,
        decision=decision,
        operation_type=proposition.operation_type if proposition else None,
        operation_id=proposition.operation_id if proposition else None,
    )


# ── Note payée en plusieurs virements ──────────────────────────────────────


async def test_chaque_versement_se_rapproche_de_sa_ligne(banque):
    db, org, compte_bancaire = banque
    note = await _note(db, org, compte_bancaire, (Decimal("40"), LE_20), (Decimal("60"), LE_22))

    out = await _importer(
        db, org, compte_bancaire,
        "Date;Libellé;Débit;Crédit\n"
        "20/10/2026;Virement Membre acompte;;40,00\n"
        "22/10/2026;Virement Membre solde;;60,00\n",
    )
    acompte, solde = out.lignes
    assert acompte.operation_type == solde.operation_type == "VERSEMENT"
    assert acompte.operation_id != solde.operation_id

    await _decider(db, org, acompte, "RAPPROCHER", acompte.propositions[0])
    await db.refresh(note)
    assert note.is_reconciled is False, "la note n'est pointée qu'au dernier virement"

    await _decider(db, org, solde, "RAPPROCHER", solde.propositions[0])
    await db.refresh(note)
    assert note.is_reconciled is True


async def test_annuler_un_rapprochement_le_rend_de_nouveau_decidable(banque):
    db, org, compte_bancaire = banque
    note = await _note(db, org, compte_bancaire, (Decimal("75"), LE_20))
    out = await _importer(
        db, org, compte_bancaire, "Date;Libellé;Débit;Crédit\n20/10/2026;Virement Membre;;75,00\n"
    )
    ligne = out.lignes[0]
    proposition = ligne.propositions[0]

    await _decider(db, org, ligne, "RAPPROCHER", proposition)
    await db.refresh(note)
    assert note.is_reconciled is True

    with pytest.raises(HTTPException) as exc:
        await _decider(db, org, ligne, "IGNORER")
    assert exc.value.status_code == 409

    annulee = await _decider(db, org, ligne, "ANNULER")
    await db.refresh(note)
    assert note.is_reconciled is False
    assert annulee.statut_rapprochement == "INCERTAIN"
    assert annulee.rapproche_par is None

    # L'opération redevient proposable : on peut la confirmer à nouveau.
    refaite = await _decider(db, org, ligne, "RAPPROCHER", proposition)
    assert refaite.statut_rapprochement == "RAPPROCHE"


async def test_annuler_sans_decision_est_refuse(banque):
    db, org, compte_bancaire = banque
    out = await _importer(
        db, org, compte_bancaire, "Date;Libellé;Débit;Crédit\n20/10/2026;Frais de tenue;12,00;\n"
    )
    with pytest.raises(HTTPException) as exc:
        await _decider(db, org, out.lignes[0], "ANNULER")
    assert exc.value.status_code == 409


# ── Répartition des propositions ───────────────────────────────────────────


async def test_deux_mouvements_identiques_recoivent_deux_operations_distinctes(banque):
    """Deux virements de 50 le même jour : chacun sa note, aucune proposée deux fois."""
    db, org, compte_bancaire = banque
    await _note(db, org, compte_bancaire, (Decimal("50"), LE_20), client="Alpha")
    await _note(db, org, compte_bancaire, (Decimal("50"), LE_20), client="Beta")

    out = await _importer(
        db, org, compte_bancaire,
        "Date;Libellé;Débit;Crédit\n"
        "20/10/2026;Virement Alpha;;50,00\n"
        "20/10/2026;Virement Beta;;50,00\n",
    )
    premiere, seconde = out.lignes
    assert premiere.operation_id and seconde.operation_id
    assert premiere.operation_id != seconde.operation_id

    # Les deux confirmations passent, l'une après l'autre.
    await _decider(db, org, premiere, "RAPPROCHER", premiere.propositions[0])
    await _decider(db, org, seconde, "RAPPROCHER", seconde.propositions[0])


# ── Lecture du relevé ──────────────────────────────────────────────────────


def test_lignes_de_solde_et_de_total_sont_ecartees():
    lignes = parser_releve(
        (
            "Date;Libellé;Débit;Crédit;Solde\n"
            "01/10/2026;Solde initial;;;1 000,00\n"
            "15/10/2026;Frais bancaires;15,50;;984,50\n"
            ";Total des mouvements;;;\n"
            "Solde au 31/10/2026;;;;984,50\n"
        ).encode("utf-8"),
        "releve.csv",
    )
    assert [l.libelle for l in lignes] == ["Frais bancaires"]


def test_ligne_sans_montant_qui_n_est_pas_un_recapitulatif_reste_une_erreur():
    with pytest.raises(ValueError, match="soit un débit, soit un crédit"):
        parser_releve(
            "Date;Libellé;Débit;Crédit\n15/10/2026;Virement reçu;;\n".encode("utf-8"),
            "releve.csv",
        )


# ── Journal de contrôle ────────────────────────────────────────────────────


@pytest.mark.parametrize("format_sortie", ["XLSX", "CSV"])
def test_journal_de_controle_accepte_les_dates_avec_fuseau(format_sortie):
    journal = SimpleNamespace(
        lignes=[
            SimpleNamespace(
                date=datetime(2026, 10, 15, 14, 30, tzinfo=timezone.utc),
                reference="REC-1",
                tiers="Membre",
                libelle="Cotisation",
                sortie=Decimal("0"),
                entree=Decimal("125.40"),
                solde=Decimal("125.40"),
            )
        ]
    )
    contenu = creer_controle_bancaire(journal, format_sortie)
    if format_sortie == "CSV":
        assert "2026-10-15;" in contenu.decode("utf-8-sig")
    else:
        ws = load_workbook(io.BytesIO(contenu)).active
        assert ws.cell(2, 1).value.date() == date(2026, 10, 15)
