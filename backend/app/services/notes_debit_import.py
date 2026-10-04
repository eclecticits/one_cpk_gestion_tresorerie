"""Import Excel des notes de débit des experts-comptables et des SEC.

La feuille fait foi. Une ligne par membre, une colonne par libellé :

    N° d'ordre | Nom          | Cotisation 2026 | Pénalité AG | Arriérés
    EC-0412    | KABONGO J.   |       600       |     100     |   450

Chaque ligne devient UNE note de débit — un encaissement non payé, rattaché à
l'expert —, et chaque cellule non vide une ligne de cette note, sous le libellé
de sa colonne. Ajouter une colonne au fichier, c'est ajouter un libellé : rien
n'est à régler dans le code.

Les montants viennent du fichier. Les tarifs réglés (Réglages → Tarifs) ne
servent qu'à deux choses : désigner le poste d'un libellé connu, et signaler un
montant qui s'en écarte. L'écart n'est jamais bloquant — la Comptabilité calcule
la cotisation d'une SEC sur son chiffre d'affaires, ce qu'aucun tarif ne sait
faire —, mais il est consigné, comme tout prix forcé.

Deux temps, sur le même code :
  - `analyser` lit le fichier et dit ce qu'il en ferait, sans rien écrire ;
  - `importer` relit le fichier, refait les mêmes contrôles et crée les notes
    dans une seule transaction : tout ou rien.

Relire plutôt que de garder l'aperçu en mémoire : le serveur reste sans état,
et ce qui est créé est exactement ce qui a été contrôlé.
"""

from __future__ import annotations

import re
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from fastapi import HTTPException, status
from sqlalchemy import extract, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.budget import BudgetExercice, BudgetPoste
from app.models.encaissement import Encaissement, EncaissementArticle
from app.models.encaissement_tarif import normaliser_libelle
from app.models.expert_comptable import ExpertComptable
from app.models.note_debit_import import NoteDebitImport
from app.models.service import Service
from app.models.service_rubrique import ServiceRubrique
from app.services.audit_service import log_action
from app.services.document_sequences import generate_document_number
from app.services.encaissement_tarifs import appliquer_tarifs, exercice_courant, succession, tarifs_resolus
from app.services.mouvements_budgetaires import hors_budget_initial_status
from app.services.service_access import get_user_service_ids

CENT = Decimal("0.01")

#: Lignes parcourues pour trouver l'en-tête : un fichier commence souvent par
#: un titre (« Notes de débit 2026 — Conseil provincial de Kinshasa »).
LIGNES_AVANT_ENTETE = 15

#: En-têtes de numérotation de lignes, pas de montants : « N° » vaut 1, 2, 3…
ENTETES_INDEX = {"n", "no", "n°", "nº", "#", "num", "numero", "n° ligne", "ordre"}

#: Catégories d'import, comme les onglets de l'import national des experts :
#: chacune a son modèle et sa validation. Une SEC déposée dans l'onglet des
#: experts-comptables (ou l'inverse) est une erreur de fichier, pas un détail.
CATEGORIES = {
    "toutes": "Toutes catégories",
    "ec": "Cotisations EC",
    "sec": "Cotisations SEC",
    "penalites": "Pénalités",
}

#: Mots qui désignent une colonne d'information même quand elle ne contient
#: que des nombres (un téléphone, une année, un nombre d'experts).
MOTS_INFORMATION = {
    "telephone", "tel", "phone", "gsm", "annee", "exercice", "nif", "code",
    "nombre", "nb", "effectif", "age", "id",
}


# ---------------------------------------------------------------------------
# Lecture de la feuille
# ---------------------------------------------------------------------------


def _sans_accents(valeur: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", valeur) if not unicodedata.combining(c))


def normaliser_entete(valeur: Any) -> str:
    """« N° d’Ordre » → « n° d'ordre » : la clé de reconnaissance d'un en-tête."""
    texte = str(valeur or "").replace("’", "'").replace("\xa0", " ")
    return " ".join(_sans_accents(texte).lower().split())


def libelle_colonne(valeur: Any) -> str:
    """Le libellé tel qu'il s'imprimera sur la note : l'en-tête, espaces resserrés."""
    return " ".join(str(valeur or "").replace("\xa0", " ").split())[:255]


def _est_entete_numero(entete: str) -> bool:
    return (
        ("ordre" in entete and entete not in ENTETES_INDEX)
        or "matricule" in entete
        or "onec" in entete.split()
    )


def _est_entete_nom(entete: str) -> bool:
    return (
        entete.startswith("nom")
        or "denomination" in entete
        or "raison sociale" in entete
        or entete in {"membre", "expert", "expert-comptable", "expert comptable", "societe"}
    )


def _est_entete_arrieres(entete: str) -> bool:
    return "arriere" in entete


def _est_entete_total(entete: str) -> bool:
    mots = set(re.split(r"[^a-z0-9]+", entete))
    return bool(mots & {"total", "solde", "cumul"})


def lire_montant(valeur: Any) -> Decimal | None:
    """Un montant de cellule, quelle que soit sa forme ; `None` si la cellule est vide.

    Accepte « 1 500 », « 1.500,00 », « 1,500.00 », « $600 », « 600 USD ».
    Lève `ValueError` pour ce qui n'est pas un nombre : une cellule « payé » ou
    « exonéré » n'est pas un zéro, et l'importer comme tel effacerait une dette.
    """
    if valeur is None:
        return None
    if isinstance(valeur, bool):
        raise ValueError(str(valeur))
    if isinstance(valeur, (int, float, Decimal)):
        return Decimal(str(valeur)).quantize(CENT)
    texte = str(valeur).strip()
    if texte in {"", "-", "—", "–"}:
        return None
    texte = re.sub(r"(?i)usd|us\$|\$|fc|cdf", "", texte)
    texte = texte.replace("\xa0", "").replace(" ", "").replace("'", "")
    if "," in texte and "." in texte:
        if texte.rfind(",") > texte.rfind("."):
            texte = texte.replace(".", "").replace(",", ".")
        else:
            texte = texte.replace(",", "")
    elif "," in texte:
        texte = texte.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", texte) else texte.replace(",", ".")
    try:
        return Decimal(texte).quantize(CENT)
    except InvalidOperation as exc:
        raise ValueError(str(valeur)) from exc


def cle_numero_ordre(valeur: Any) -> str:
    """Clé de rapprochement d'un n° d'ordre : « ec-0412 », « EC 412 » → « EC412 »."""
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    cle = re.sub(r"[^A-Z0-9]", "", _sans_accents(str(valeur)).upper())
    return re.sub(r"(?<!\d)0+(?=\d)", "", cle)


def cle_nom(valeur: Any) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", _sans_accents(str(valeur or "")).lower()).split())


@dataclass
class Colonne:
    index: int
    libelle: str
    arrieres: bool = False


@dataclass
class Feuille:
    ligne_entete: int
    col_numero: int | None
    col_nom: int | None
    col_total: int | None
    colonnes: list[Colonne]
    ignorees: list[dict[str, str]]
    #: (n° de ligne Excel, valeurs)
    lignes: list[tuple[int, list[Any]]]


def lire_feuille(contenu: bytes) -> Feuille:
    """Repère l'en-tête et classe les colonnes : identité, montants, le reste."""
    try:
        from openpyxl import load_workbook
    except Exception:  # pragma: no cover - dépendance de l'image
        raise HTTPException(status_code=500, detail="openpyxl n'est pas installé")
    try:
        classeur = load_workbook(filename=BytesIO(contenu), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Fichier illisible : déposez un classeur Excel (.xlsx).",
        )
    feuille = classeur.active
    toutes = [list(ligne) for ligne in feuille.iter_rows(values_only=True)]
    classeur.close()

    ligne_entete = None
    for position, ligne in enumerate(toutes[:LIGNES_AVANT_ENTETE]):
        entetes = [normaliser_entete(v) for v in ligne]
        if any(_est_entete_numero(e) for e in entetes if e) or (
            any(_est_entete_nom(e) for e in entetes if e) and sum(1 for e in entetes if e) >= 2
        ):
            ligne_entete = position
            break
    if ligne_entete is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Aucune ligne d'en-tête reconnue : la feuille doit avoir une colonne « N° d'ordre » ou « Nom ».",
        )

    brutes = toutes[ligne_entete]
    # N° de ligne Excel (1-indexé) compté AVANT d'écarter les lignes vides :
    # c'est celui que l'utilisateur retrouvera dans son fichier.
    donnees = [
        (ligne_entete + 2 + decalage, ligne)
        for decalage, ligne in enumerate(toutes[ligne_entete + 1:])
        if any(v not in (None, "") for v in ligne)
    ]

    col_numero = col_nom = col_total = None
    colonnes: list[Colonne] = []
    ignorees: list[dict[str, str]] = []
    for index, brut in enumerate(brutes):
        entete = normaliser_entete(brut)
        if not entete:
            continue
        libelle = libelle_colonne(brut)
        if col_numero is None and _est_entete_numero(entete):
            col_numero = index
            continue
        if col_nom is None and _est_entete_nom(entete):
            col_nom = index
            continue
        if entete in ENTETES_INDEX:
            ignorees.append({"libelle": libelle, "raison": "Numérotation des lignes"})
            continue
        if _est_entete_total(entete):
            if col_total is None:
                col_total = index
            ignorees.append({"libelle": libelle, "raison": "Total : recalculé, et comparé à la somme des colonnes"})
            continue

        valeurs = [ligne[index] for _, ligne in donnees if index < len(ligne) and ligne[index] not in (None, "")]
        arrieres = _est_entete_arrieres(entete)
        if not valeurs and not arrieres:
            ignorees.append({"libelle": libelle, "raison": "Colonne vide"})
            continue
        if not arrieres and set(re.split(r"[^a-z0-9]+", entete)) & MOTS_INFORMATION:
            ignorees.append({"libelle": libelle, "raison": "Colonne d'information"})
            continue
        numeriques = 0
        for valeur in valeurs:
            try:
                lire_montant(valeur)
                numeriques += 1
            except ValueError:
                pass
        # Une colonne d'information (province, catégorie…) n'a aucun nombre ; une
        # colonne de montants peut en avoir un ou deux mal saisis, que le contrôle
        # signalera ligne par ligne plutôt que de taire la colonne entière.
        if not arrieres and numeriques * 2 < len(valeurs):
            ignorees.append({"libelle": libelle, "raison": "Colonne d'information"})
            continue
        colonnes.append(Colonne(index=index, libelle=libelle, arrieres=arrieres))

    if col_numero is None and col_nom is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La feuille doit avoir une colonne « N° d'ordre » (ou, à défaut, « Nom »).",
        )
    if not colonnes:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Aucune colonne de montants : chaque libellé de la note doit avoir sa colonne.",
        )
    return Feuille(
        ligne_entete=ligne_entete + 1,
        col_numero=col_numero,
        col_nom=col_nom,
        col_total=col_total,
        colonnes=colonnes,
        ignorees=ignorees,
        lignes=donnees,
    )


# ---------------------------------------------------------------------------
# Contexte : experts, postes, services, tarifs, notes déjà émises
# ---------------------------------------------------------------------------


async def services_disponibles(db: AsyncSession, user: Any, tenant_id: int) -> list[Service]:
    """Les services sous lesquels cet utilisateur peut émettre une note."""
    requete = select(Service).where(Service.organisation_id == tenant_id, Service.is_active.is_(True))
    if (getattr(user, "role", "") or "").lower() not in {"admin", "super_admin"}:
        ids = await get_user_service_ids(db, user)
        if not ids:
            return []
        requete = requete.where(Service.id.in_(ids))
    return list((await db.execute(requete.order_by(Service.libelle))).scalars().all())


async def resoudre_service(db: AsyncSession, user: Any, tenant_id: int, service_id: int | None) -> int | None:
    """Même règle que la saisie d'un encaissement : un agent émet sous l'un de ses services."""
    role = (getattr(user, "role", "") or "").lower()
    if role in {"admin", "super_admin"}:
        if service_id is None:
            return None
        trouve = (
            await db.execute(
                select(Service.id).where(
                    Service.id == service_id, Service.organisation_id == tenant_id, Service.is_active.is_(True)
                )
            )
        ).scalar_one_or_none()
        if trouve is None:
            raise HTTPException(status_code=400, detail="service_id invalide")
        return service_id
    ids = await get_user_service_ids(db, user)
    if not ids:
        raise HTTPException(status_code=403, detail="Aucun service assigné")
    if service_id is None:
        if len(ids) == 1:
            return ids[0]
        raise HTTPException(status_code=400, detail="Choisissez le service sous lequel émettre les notes.")
    if service_id not in ids:
        raise HTTPException(status_code=403, detail="Accès interdit à ce service")
    return service_id


async def postes_disponibles(
    db: AsyncSession, tenant_id: int, service_id: int | None
) -> tuple[BudgetExercice | None, list[BudgetPoste]]:
    """Les postes de recette de l'exercice courant, bornés aux rubriques du service."""
    exercice = await exercice_courant(db, tenant_id)
    if exercice is None:
        return None, []
    requete = select(BudgetPoste).where(
        BudgetPoste.organisation_id == tenant_id,
        BudgetPoste.exercice_id == exercice.id,
        BudgetPoste.is_deleted.is_(False),
        BudgetPoste.active.is_(True),
        func.upper(BudgetPoste.type) == "RECETTE",
    )
    if service_id is not None:
        requete = requete.where(
            BudgetPoste.id.in_(
                select(ServiceRubrique.budget_poste_id).where(ServiceRubrique.service_id == service_id)
            )
        )
    return exercice, list((await db.execute(requete.order_by(BudgetPoste.code))).scalars().all())


async def _codes_arrieres(db: AsyncSession, tenant_id: int, exercice: BudgetExercice) -> set[str]:
    """Les codes que l'exercice précédent désigne comme ses postes d'arriérés."""
    lignes = (
        await db.execute(
            select(BudgetPoste.code_poste_arrieres)
            .join(BudgetExercice, BudgetExercice.id == BudgetPoste.exercice_id)
            .where(
                BudgetPoste.organisation_id == tenant_id,
                BudgetExercice.annee == exercice.annee - 1,
                BudgetPoste.is_deleted.is_(False),
                BudgetPoste.code_poste_arrieres.is_not(None),
            )
        )
    ).scalars().all()
    return {code.strip().upper() for code in lignes if code and code.strip()}


async def _experts(db: AsyncSession) -> tuple[dict[str, ExpertComptable], dict[str, list[ExpertComptable]]]:
    """Le référentiel, indexé par n° d'ordre et par nom."""
    experts = (await db.execute(select(ExpertComptable))).scalars().all()
    par_numero: dict[str, ExpertComptable] = {}
    par_nom: dict[str, list[ExpertComptable]] = {}
    for expert in experts:
        cle = cle_numero_ordre(expert.numero_ordre)
        if cle:
            par_numero.setdefault(cle, expert)
        par_nom.setdefault(cle_nom(expert.nom_denomination), []).append(expert)
    return par_numero, par_nom


async def _deja_emises(
    db: AsyncSession, tenant_id: int, expert_ids: set[uuid.UUID], annee: int
) -> dict[tuple[uuid.UUID, str], str]:
    """(expert, libellé) → n° de la note qui le porte déjà cette année."""
    if not expert_ids:
        return {}
    lignes = (
        await db.execute(
            select(Encaissement.expert_comptable_id, EncaissementArticle.libelle, Encaissement.numero_recu)
            .join(EncaissementArticle, EncaissementArticle.encaissement_id == Encaissement.id)
            .where(
                Encaissement.organisation_id == tenant_id,
                Encaissement.expert_comptable_id.in_(expert_ids),
                Encaissement.is_deleted.is_(False),
                Encaissement.est_proforma.is_(False),
                Encaissement.statut_operation != "ANNULEE",
                extract("year", Encaissement.date_encaissement) == annee,
            )
        )
    ).all()
    return {(ligne[0], normaliser_libelle(ligne[1])): ligne[2] or "—" for ligne in lignes}


# ---------------------------------------------------------------------------
# Analyse
# ---------------------------------------------------------------------------


@dataclass
class LigneAnalysee:
    ligne: int
    numero_ordre: str
    nom: str
    expert: ExpertComptable | None = None
    montants: dict[int, Decimal] = field(default_factory=dict)
    erreurs: list[str] = field(default_factory=list)
    avertissements: list[str] = field(default_factory=list)
    doublon: bool = False

    @property
    def total(self) -> Decimal:
        return sum(self.montants.values(), Decimal("0.00"))

    @property
    def statut(self) -> str:
        if self.erreurs:
            return "erreur"
        if self.avertissements or self.doublon:
            return "avertissement"
        return "ok"


@dataclass
class Analyse:
    feuille: Feuille
    exercice: BudgetExercice | None
    annee: int
    lignes: list[LigneAnalysee]
    colonnes: list[dict[str, Any]]
    postes: list[BudgetPoste]
    services: list[Service]
    service_id: int | None
    #: index de colonne → (tarif, poste du tarif)
    tarifs: dict[int, tuple[Any, BudgetPoste | None]]
    categorie: str = "toutes"


def _montant_texte(valeur: Decimal) -> str:
    return f"{valeur:,.2f}".replace(",", " ")


async def analyser(
    db: AsyncSession,
    *,
    tenant_id: int,
    user: Any,
    contenu: bytes,
    service_id: int | None,
    categorie: str = "toutes",
) -> Analyse:
    if categorie not in CATEGORIES:
        raise HTTPException(status_code=400, detail="Catégorie d'import inconnue")
    feuille = lire_feuille(contenu)
    services = await services_disponibles(db, user, tenant_id)
    # L'aperçu ne bloque pas sur le service : un agent de plusieurs services le
    # choisira à l'écran, et c'est l'import qui l'exigera (`resoudre_service`).
    if service_id is not None:
        service_retenu = await resoudre_service(db, user, tenant_id, service_id)
    elif (getattr(user, "role", "") or "").lower() not in {"admin", "super_admin"} and len(services) == 1:
        service_retenu = services[0].id
    else:
        service_retenu = None
    exercice, postes = await postes_disponibles(db, tenant_id, service_retenu)
    annee = exercice.annee if exercice is not None else datetime.now().year
    postes_par_code = {(p.code or "").strip().upper(): p for p in postes}

    aujourdhui = datetime.now().date()
    par_libelle = {
        tarif.libelle_normalise: (tarif, poste)
        for tarif, poste in await tarifs_resolus(db, tenant_id, actifs_seulement=True, a_la_date=aujourdhui)
    }
    codes_arrieres = await _codes_arrieres(db, tenant_id, exercice) if exercice is not None else set()
    postes_arrieres = [p for code, p in postes_par_code.items() if code in codes_arrieres]

    tarifs: dict[int, tuple[Any, BudgetPoste | None]] = {}
    colonnes: list[dict[str, Any]] = []
    for colonne in feuille.colonnes:
        cle = normaliser_libelle(colonne.libelle)
        erreur = None
        trouve = par_libelle.get(cle)
        if trouve is not None:
            tarifs[colonne.index] = trouve
        else:
            remplacant = await succession(db, tenant_id, cle)
            if remplacant is not None:
                erreur = (
                    f"« {colonne.libelle} » est devenu « {remplacant.libelle} » : "
                    f"renommez la colonne du fichier."
                )
        tarif, poste_tarif = trouve if trouve is not None else (None, None)
        suggere = poste_tarif
        if suggere is None and colonne.arrieres:
            if len(postes_arrieres) == 1:
                suggere = postes_arrieres[0]
            else:
                suggere = next((p for p in postes if "arri" in normaliser_entete(p.libelle)), None)
        if categorie == "penalites" and not colonne.arrieres and "penalit" not in normaliser_entete(colonne.libelle):
            avertissement = f"« {colonne.libelle} » ne ressemble pas à une pénalité : vérifiez l'onglet choisi."
        else:
            avertissement = None
        colonnes.append(
            {
                "avertissement": avertissement,
                "cle": str(colonne.index),
                "libelle": colonne.libelle,
                "arrieres": colonne.arrieres,
                "tarif": (
                    {
                        "id": tarif.id,
                        "libelle": tarif.libelle,
                        "montant": str(tarif.montant) if tarif.montant is not None else None,
                        "poste_code": tarif.budget_poste_code,
                    }
                    if tarif is not None
                    else None
                ),
                # Un tarif qui fixe le poste l'impose : le choisir ailleurs serait
                # un forçage, que l'import n'a pas à faire en silence.
                "poste_impose": poste_tarif is not None,
                "poste_suggere_id": suggere.id if suggere is not None else None,
                "erreur": erreur,
                "total": "0.00",
                "nb_lignes": 0,
            }
        )

    par_numero, par_nom = await _experts(db)
    lignes: list[LigneAnalysee] = []
    numeros_vus: dict[str, int] = {}
    for numero_ligne, valeurs in feuille.lignes:
        def cellule(index: int | None) -> Any:
            return valeurs[index] if index is not None and index < len(valeurs) else None

        brut_numero = cellule(feuille.col_numero)
        numero = "" if brut_numero is None else str(int(brut_numero) if isinstance(brut_numero, float) and brut_numero.is_integer() else brut_numero).strip()
        nom = libelle_colonne(cellule(feuille.col_nom))
        ligne = LigneAnalysee(ligne=numero_ligne, numero_ordre=numero, nom=nom)

        cle = cle_numero_ordre(numero)
        if cle:
            ligne.expert = par_numero.get(cle)
            if ligne.expert is None:
                ligne.erreurs.append(f"N° d'ordre « {numero} » introuvable dans la liste des experts")
            elif cle in numeros_vus:
                ligne.erreurs.append(f"Membre déjà présent à la ligne {numeros_vus[cle]} du fichier")
            else:
                numeros_vus[cle] = numero_ligne
        elif nom:
            candidats = par_nom.get(cle_nom(nom), [])
            if len(candidats) == 1:
                ligne.expert = candidats[0]
                ligne.avertissements.append("Rapproché par le nom, faute de n° d'ordre")
            elif candidats:
                ligne.erreurs.append(f"Plusieurs experts s'appellent « {nom} » : indiquez le n° d'ordre")
            else:
                ligne.erreurs.append(f"« {nom} » introuvable dans la liste des experts")
        else:
            ligne.erreurs.append("Ni n° d'ordre ni nom")

        for colonne in feuille.colonnes:
            try:
                montant = lire_montant(cellule(colonne.index))
            except ValueError:
                ligne.erreurs.append(f"« {colonne.libelle} » : « {cellule(colonne.index)} » n'est pas un montant")
                continue
            if montant is None or montant == 0:
                continue
            if montant < 0:
                ligne.erreurs.append(f"« {colonne.libelle} » : montant négatif")
                continue
            ligne.montants[colonne.index] = montant
            tarif = tarifs.get(colonne.index, (None, None))[0]
            if tarif is not None and tarif.montant is not None:
                prix = Decimal(str(tarif.montant))
                if (montant % prix) != 0:
                    ligne.avertissements.append(
                        f"« {colonne.libelle} » : {_montant_texte(montant)} au lieu du tarif "
                        f"{_montant_texte(prix)} — le montant du fichier est retenu"
                    )

        if not ligne.montants and not ligne.erreurs:
            ligne.erreurs.append("Aucun montant sur la ligne")

        if feuille.col_total is not None and ligne.montants:
            try:
                total_fichier = lire_montant(cellule(feuille.col_total))
            except ValueError:
                total_fichier = None
            if total_fichier is not None and total_fichier != ligne.total:
                ligne.avertissements.append(
                    f"Total du fichier {_montant_texte(total_fichier)} ≠ somme des colonnes "
                    f"{_montant_texte(ligne.total)} — la somme est retenue"
                )

        if ligne.expert is not None and categorie in {"ec", "sec"}:
            est_sec = ligne.expert.type_ec == "SEC"
            if categorie == "ec" and est_sec:
                ligne.erreurs.append(
                    f"{ligne.expert.nom_denomination} est une SEC : importez-la dans l'onglet « Cotisations SEC »"
                )
            elif categorie == "sec" and not est_sec:
                ligne.erreurs.append(
                    f"{ligne.expert.nom_denomination} n'est pas une SEC : importez-le dans l'onglet « Cotisations EC »"
                )

        if ligne.expert is not None and not ligne.expert.active:
            ligne.avertissements.append("Expert inactif dans la liste")
        lignes.append(ligne)

    deja = await _deja_emises(db, tenant_id, {l.expert.id for l in lignes if l.expert is not None}, annee)
    libelles = {c.index: c.libelle for c in feuille.colonnes}
    for ligne in lignes:
        if ligne.expert is None or ligne.erreurs:
            continue
        for index in ligne.montants:
            numero = deja.get((ligne.expert.id, normaliser_libelle(libelles[index])))
            if numero:
                ligne.doublon = True
                ligne.avertissements.append(f"« {libelles[index]} » déjà émis en {annee} (note {numero})")

    par_cle = {c["cle"]: c for c in colonnes}
    for ligne in lignes:
        if ligne.erreurs:
            continue
        for index, montant in ligne.montants.items():
            info = par_cle[str(index)]
            info["total"] = str(Decimal(info["total"]) + montant)
            info["nb_lignes"] += 1

    return Analyse(
        feuille=feuille,
        exercice=exercice,
        annee=annee,
        lignes=lignes,
        colonnes=colonnes,
        postes=postes,
        services=services,
        service_id=service_retenu,
        tarifs=tarifs,
        categorie=categorie,
    )


def analyse_en_reponse(analyse: Analyse) -> dict[str, Any]:
    valides = [l for l in analyse.lignes if not l.erreurs]
    arrieres = {c.index for c in analyse.feuille.colonnes if c.arrieres}
    return {
        "categorie": analyse.categorie,
        "ligne_entete": analyse.feuille.ligne_entete,
        "exercice": analyse.annee,
        "exercice_ouvert": analyse.exercice is not None,
        "service_id": analyse.service_id,
        "services": [{"id": s.id, "code": s.code, "libelle": s.libelle} for s in analyse.services],
        "postes": [{"id": p.id, "code": p.code, "libelle": p.libelle} for p in analyse.postes],
        "colonnes": analyse.colonnes,
        "colonnes_ignorees": analyse.feuille.ignorees,
        "colonne_numero": analyse.feuille.col_numero is not None,
        "lignes": [
            {
                "ligne": l.ligne,
                "numero_ordre": l.numero_ordre,
                "nom": l.nom,
                "expert": (
                    {
                        "id": str(l.expert.id),
                        "numero_ordre": l.expert.numero_ordre,
                        "nom": l.expert.nom_denomination,
                        "type_ec": l.expert.type_ec,
                    }
                    if l.expert is not None
                    else None
                ),
                "montants": {str(k): str(v) for k, v in l.montants.items()},
                "total": str(l.total),
                "statut": l.statut,
                "doublon": l.doublon,
                "erreurs": l.erreurs,
                "avertissements": l.avertissements,
            }
            for l in analyse.lignes
        ],
        "resume": {
            "nb_lignes": len(analyse.lignes),
            "nb_ok": sum(1 for l in analyse.lignes if l.statut == "ok"),
            "nb_avertissements": sum(1 for l in analyse.lignes if l.statut == "avertissement"),
            "nb_erreurs": sum(1 for l in analyse.lignes if l.statut == "erreur"),
            "nb_doublons": sum(1 for l in valides if l.doublon),
            "total": str(sum((l.total for l in valides), Decimal("0.00"))),
            "total_arrieres": str(
                sum((m for l in valides for i, m in l.montants.items() if i in arrieres), Decimal("0.00"))
            ),
        },
    }


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------


async def importer(
    db: AsyncSession,
    *,
    tenant_id: int,
    user: Any,
    fichier: str,
    contenu: bytes,
    service_id: int | None,
    postes: dict[str, int],
    importer_doublons: bool = False,
    categorie: str = "toutes",
) -> dict[str, Any]:
    """Crée les notes du fichier, toutes ou aucune."""
    # Importé ici : l'endpoint des encaissements importe déjà ce module-ci par
    # le routeur, et ces deux fonctions sont celles de la saisie — les recopier
    # serait laisser l'import diverger du formulaire au premier correctif.
    from app.api.v1.endpoints.encaissements import _add_encaissement_articles, _valider_postes_articles
    from app.core.horodatage import resoudre_date_operation

    user_id = getattr(user, "id", None)
    analyse = await analyser(
        db, tenant_id=tenant_id, user=user, contenu=contenu, service_id=service_id, categorie=categorie
    )
    service_retenu = await resoudre_service(db, user, tenant_id, service_id)
    if analyse.exercice is None:
        raise HTTPException(status_code=400, detail="Aucun exercice budgétaire : ouvrez l'exercice avant d'importer.")

    erreurs_colonnes = [c["erreur"] for c in analyse.colonnes if c["erreur"]]
    if erreurs_colonnes:
        raise HTTPException(status_code=400, detail=" ".join(erreurs_colonnes))

    postes_valides = {p.id for p in analyse.postes}
    poste_par_colonne: dict[int, int] = {}
    manquants: list[str] = []
    for info in analyse.colonnes:
        index = int(info["cle"])
        if not info["nb_lignes"]:
            continue
        choisi = postes.get(info["cle"]) or info["poste_suggere_id"]
        if info["poste_impose"]:
            choisi = info["poste_suggere_id"]
        if choisi is None:
            manquants.append(info["libelle"])
            continue
        if int(choisi) not in postes_valides:
            raise HTTPException(
                status_code=400,
                detail=f"« {info['libelle']} » : poste non autorisé pour ce service ou hors de l'exercice {analyse.annee}.",
            )
        poste_par_colonne[index] = int(choisi)
    if manquants:
        raise HTTPException(
            status_code=400,
            detail="Choisissez le poste budgétaire de : " + ", ".join(f"« {m} »" for m in manquants),
        )

    a_creer = [l for l in analyse.lignes if not l.erreurs and (importer_doublons or not l.doublon)]
    if not a_creer:
        raise HTTPException(status_code=400, detail="Aucune note à créer : toutes les lignes sont en erreur ou déjà émises.")

    await _valider_postes_articles(
        db,
        tenant_id=tenant_id,
        service_id=service_retenu,
        articles=[{"budget_poste_id": pid} for pid in set(poste_par_colonne.values())],
        impact_budgetaire=True,
    )
    postes_par_id = {p.id: p for p in analyse.postes}
    libelles = {c.index: c.libelle for c in analyse.feuille.colonnes}
    arrieres = {c.index for c in analyse.feuille.colonnes if c.arrieres}
    date_note = resoudre_date_operation(None, user=user, champ="date_encaissement")

    total = Decimal("0.00")
    total_arrieres = Decimal("0.00")
    enregistrement = NoteDebitImport(
        organisation_id=tenant_id,
        fichier=(fichier or "notes_de_debit.xlsx")[:300],
        service_id=service_retenu,
        created_by=user_id,
        colonnes=[
            {
                "libelle": libelles[index],
                "arrieres": index in arrieres,
                "poste_id": poste_id,
                "poste_code": postes_par_id[poste_id].code,
            }
            for index, poste_id in poste_par_colonne.items()
        ],
    )
    db.add(enregistrement)
    await db.flush()

    notes: list[dict[str, Any]] = []
    for ligne in a_creer:
        expert = ligne.expert
        assert expert is not None
        articles: list[dict[str, Any]] = []
        for ordre, (index, montant) in enumerate(ligne.montants.items()):
            quantite, prix = Decimal("1.00"), montant
            tarif = analyse.tarifs.get(index, (None, None))[0]
            # Trois pénalités à 100, c'est 3 × 100 et non une ligne de 300 qui
            # s'écarterait du tarif : la quantité se déduit quand elle tombe juste.
            if tarif is not None and tarif.montant is not None:
                tarife = Decimal(str(tarif.montant))
                if tarife > 0 and montant % tarife == 0:
                    quantite, prix = (montant / tarife).quantize(CENT), tarife
            articles.append(
                {
                    "libelle": libelles[index],
                    "description": None,
                    "quantite": quantite,
                    "prix_unitaire": prix,
                    "montant": montant,
                    "budget_poste_id": poste_par_colonne[index],
                    "sort_order": ordre,
                }
            )
        ecarts = await appliquer_tarifs(
            db, tenant_id, articles, peut_forcer=True, date_encaissement=date_note.date()
        )
        montant_note = ligne.total
        poste_entete = postes_par_id[articles[0]["budget_poste_id"]]
        note = Encaissement(
            numero_recu=await generate_document_number(db, doc_type="ND", tenant_id=tenant_id, service_id=None),
            numero_proforma=None,
            est_proforma=False,
            organisation_id=tenant_id,
            type_client="sec" if expert.type_ec == "SEC" else "expert_comptable",
            expert_comptable_id=expert.id,
            client_nom=None,
            client_id=None,
            libelle=", ".join(a["libelle"] for a in articles)[:255],
            description=f"Import « {enregistrement.fichier} », ligne {ligne.ligne}",
            montant=montant_note,
            montant_total=montant_note,
            montant_paye=Decimal("0.00"),
            montant_percu=Decimal("0.00"),
            devise_perception="USD",
            taux_change_applique=Decimal("1.00"),
            budget_poste_id=poste_entete.id,
            budget_poste_code=poste_entete.code,
            budget_poste_libelle=poste_entete.libelle,
            service_id=service_retenu,
            statut_paiement="non_paye",
            mode_paiement="cash",
            nature_mouvement="BUDGETAIRE",
            impact_budgetaire=True,
            hors_budget_status=hors_budget_initial_status("BUDGETAIRE"),
            canal="CAISSE",
            date_encaissement=date_note,
            date_paiement=None,
            created_by=user_id,
            note_debit_import_id=enregistrement.id,
        )
        db.add(note)
        await db.flush()
        _add_encaissement_articles(db, note, tenant_id, articles)
        if ecarts:
            await log_action(
                db,
                user_id=user_id,
                action="ENCAISSEMENT_TARIF_FORCE",
                target_table="encaissements",
                target_id=str(note.id),
                new_value={"ecarts": ecarts, "import": str(enregistrement.id)},
            )
        total += montant_note
        total_arrieres += sum((m for i, m in ligne.montants.items() if i in arrieres), Decimal("0.00"))
        notes.append(
            {
                "id": str(note.id),
                "numero_recu": note.numero_recu,
                "numero_ordre": expert.numero_ordre,
                "nom": expert.nom_denomination,
                "montant_total": str(montant_note),
            }
        )

    enregistrement.nb_notes = len(notes)
    enregistrement.nb_lignes_ecartees = len(analyse.lignes) - len(notes)
    enregistrement.montant_total = total
    enregistrement.montant_arrieres = total_arrieres
    await log_action(
        db,
        user_id=user_id,
        action="NOTES_DEBIT_IMPORT",
        target_table="notes_debit_imports",
        target_id=str(enregistrement.id),
        new_value={
            "categorie": categorie,
            "fichier": enregistrement.fichier,
            "nb_notes": len(notes),
            "montant_total": str(total),
            "montant_arrieres": str(total_arrieres),
        },
    )
    await db.commit()
    return {
        "categorie": categorie,
        "import_id": str(enregistrement.id),
        "fichier": enregistrement.fichier,
        "nb_notes": len(notes),
        "nb_lignes_ecartees": enregistrement.nb_lignes_ecartees,
        "montant_total": str(total),
        "montant_arrieres": str(total_arrieres),
        "notes": notes,
    }
