"""Journaux de trésorerie joints aux rapports planifiés (hebdomadaire, mensuel).

Un fichier par journal : la caisse en USD, la caisse en CDF, puis chaque compte
bancaire actif, sur la période du rapport. En PDF pour le rapport hebdomadaire,
en Excel pour le rapport mensuel — chacun au modèle de l'écran Rapports.

Deux règles :

* **Mêmes chiffres que l'écran.** Les montants viennent du calcul de l'écran
  « Journal de trésorerie » (`GET /reports/journal-tresorerie`).
* **Même document que l'écran.** La mise en page reproduit celle du PDF tiré
  depuis Rapports (`frontend/src/utils/pdfJournalGenerator.ts` et
  `pdfTenantIdentity.ts`) : en-tête de l'organisation, cartouches de soldes,
  tableau avec report à nouveau et totaux, cadres de signature, pied de page.
  Le navigateur dessine l'un, le serveur l'autre ; une évolution de l'un doit
  être reportée sur l'autre.

Un journal sans mouvement et à solde nul n'est pas joint : il n'apprendrait
rien au lecteur et alourdirait le courriel.
"""

from __future__ import annotations

import io
import logging
import os
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

import anyio
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import BaseDocTemplate, Flowable, Frame, PageTemplate, Paragraph, Table, TableStyle
from reportlab.platypus.flowables import KeepTogether, Spacer
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.models.banque import Banque
from app.models.compte_bancaire import CompteBancaire
from app.models.print_settings import PrintSettings

logger = logging.getLogger("onec_cpk_api.journaux_annexes")

# ── Palette et mesures de pdfJournalGenerator.ts / pdfTenantIdentity.ts ─────

def _rgb(r: int, g: int, b: int) -> colors.Color:
    return colors.Color(r / 255, g / 255, b / 255)


VERT = _rgb(6, 95, 70)
FOND_ENTETE = colors.HexColor("#f0fdf4")
ENCRE = _rgb(15, 23, 42)
GRIS_TEXTE = _rgb(90, 100, 115)
GRIS_PIED = _rgb(110, 120, 135)
TRAIT = _rgb(226, 232, 240)
TRAIT_CADRE = _rgb(203, 213, 225)
LIGNE_PAIRE = _rgb(248, 250, 252)
FOND_REPORT = _rgb(241, 245, 249)
ENTREE = _rgb(21, 128, 61)
SORTIE = _rgb(185, 28, 28)
TIRET = _rgb(148, 163, 184)
FILET_JOUR = _rgb(203, 213, 225)

TITRE_ONEC = "ORDRE NATIONAL DES EXPERTS-COMPTABLES"
NOM_PAR_DEFAUT = "ONEC-RDC"
MARGE = 12 * mm
LARGEUR, HAUTEUR = A4
#: Cadre du logo dans le bandeau, en mm.
LOGO_LARGEUR_MAX = 34
LOGO_HAUTEUR_MAX = 21
LOGO_PAR_DEFAUT = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "imge_onec.png")


@dataclass(frozen=True)
class Identite:
    nom: str
    contact: str
    logo: bytes | None


@dataclass(frozen=True)
class JournalAnnexe:
    nom_compte: str
    devise: str
    lignes: list
    solde_initial: Decimal
    total_entrees: Decimal
    total_sorties: Decimal
    solde_final: Decimal


def _montant(value) -> str:
    """`formatAmount` du frontend : « 1 234,56 »."""
    try:
        return f"{float(value or 0):,.2f}".replace(",", " ").replace(".", ",")
    except (TypeError, ValueError):
        return "0,00"


def _libelle_complet(ligne) -> str:
    """`libelleJournalComplet` du frontend : précisions et référence dans une seule parenthèse."""
    reference = (getattr(ligne, "reference", None) or "").strip()
    precision = getattr(ligne, "precision", None)
    if precision:
        base = (getattr(ligne, "libelle_base", None) or "").strip()
        details = " · ".join(p for p in (precision, f"réf. {reference}" if reference else "") if p)
        return f"{base} ({details})" if base else details
    libelle = (getattr(ligne, "libelle", None) or "").strip()
    return f"{libelle} ({reference})".strip() if reference else libelle


def _echapper(texte: str) -> str:
    return texte.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ── Données ──────────────────────────────────────────────────────────────────


async def _identite(db: AsyncSession, tenant_id: int, organisation: str) -> Identite:
    ps = (
        await db.execute(select(PrintSettings).where(PrintSettings.organisation_id == tenant_id).limit(1))
    ).scalar_one_or_none()
    nom = (getattr(ps, "organization_name", "") or "").strip() or organisation or NOM_PAR_DEFAUT
    contact = "  •  ".join(
        v for v in ((getattr(ps, "address", "") or ""), (getattr(ps, "phone", "") or ""), (getattr(ps, "email", "") or "")) if v.strip()
    )
    logo: bytes | None = None
    if ps is None or getattr(ps, "show_header_logo", True) is not False:
        chemin = LOGO_PAR_DEFAUT
        logo_url = (getattr(ps, "logo_url", "") or "").strip()
        if logo_url:
            rel = logo_url.replace("/uploads/", "", 1).lstrip("/") if logo_url.startswith("/uploads/") else logo_url.lstrip("/")
            racine = os.path.abspath(settings.upload_dir) if settings.upload_dir else ""
            candidat = os.path.abspath(os.path.join(racine, rel)) if racine else ""
            if candidat and candidat.startswith(racine) and os.path.exists(candidat):
                chemin = candidat
        try:
            with open(chemin, "rb") as f:
                logo = f.read()
        except OSError:
            logo = None
    return Identite(nom=nom, contact=contact, logo=logo)


async def _journaux(db: AsyncSession, tenant_id: int, debut: str, fin: str) -> list[JournalAnnexe]:
    # Import différé : le module des rapports est un endpoint, qui importe
    # lui-même des services ; l'importer au chargement créerait un cycle.
    from app.api.v1.endpoints.reports import journal_tresorerie

    # Mêmes noms de compte que l'écran Rapports.
    sources: list[tuple[str, str, str, int | None]] = [
        ("Caisse USD", "CAISSE", "USD", None),
        ("Caisse CDF", "CAISSE", "CDF", None),
    ]
    comptes = (
        await db.execute(
            select(CompteBancaire, Banque)
            .outerjoin(Banque, CompteBancaire.banque_id == Banque.id)
            .where(
                CompteBancaire.organisation_id == tenant_id,
                CompteBancaire.is_active.is_(True),
                CompteBancaire.account_type == "BANK",
            )
            .order_by(CompteBancaire.devise, CompteBancaire.intitule)
        )
    ).all()
    for compte, banque in comptes:
        sources.append((f"{banque.nom if banque else 'Banque'} - {compte.intitule}", "BANQUE", compte.devise, compte.id))

    journaux: list[JournalAnnexe] = []
    for nom_compte, canal, devise, compte_id in sources:
        try:
            reponse = await journal_tresorerie(
                canal=canal,
                devise=devise,
                compte_bancaire_id=compte_id,
                date_debut=debut,
                date_fin=fin,
                user=None,
                db=db,
                tenant_id=tenant_id,
            )
        except Exception:
            # Un journal qui ne se calcule pas ne doit pas priver le rapport des
            # autres : il manque, et le journal du serveur dit pourquoi.
            logger.exception("Journal non calculé pour l'annexe : %s (organisation %s)", nom_compte, tenant_id)
            continue
        if not reponse.lignes and not reponse.solde_initial and not reponse.solde_final:
            continue
        journaux.append(
            JournalAnnexe(
                nom_compte=nom_compte,
                devise=devise,
                lignes=list(reponse.lignes),
                solde_initial=reponse.solde_initial,
                total_entrees=reponse.total_entrees,
                total_sorties=reponse.total_sorties,
                solde_final=reponse.solde_final,
            )
        )
    return journaux


# ── Tracé ────────────────────────────────────────────────────────────────────


def _y(y_mm: float) -> float:
    """Ordonnée jsPDF (depuis le haut, en mm) vers reportlab (depuis le bas, en points)."""
    return HAUTEUR - y_mm * mm


def _entete(c: rl_canvas.Canvas, identite: Identite, titre: str, sous_titre: str) -> None:
    """`drawTenantHeader` : bandeau de l'organisation, filet vert, titre, période."""
    bande = 26
    c.setFillColor(FOND_ENTETE)
    c.roundRect(MARGE, _y(8 + bande), LARGEUR - 2 * MARGE, bande * mm, 2.5 * mm, stroke=0, fill=1)

    gauche = MARGE + 5 * mm
    if identite.logo:
        try:
            image = ImageReader(io.BytesIO(identite.logo))
            largeur_px, hauteur_px = image.getSize()
            # Proportions gardées : le logo tient dans 34 × 21 mm, centré dans
            # la hauteur du bandeau, au lieu d'être écrasé dans un carré.
            echelle = min(LOGO_LARGEUR_MAX / largeur_px, LOGO_HAUTEUR_MAX / hauteur_px)
            w, h = largeur_px * echelle, hauteur_px * echelle
            y_haut = 8 + (bande - h) / 2
            c.drawImage(image, MARGE + 3 * mm, _y(y_haut + h), w * mm, h * mm, mask="auto")
            gauche = MARGE + (3 + w + 4) * mm
        except Exception:
            # Un logo illisible ne doit jamais empêcher l'émission du document.
            pass

    c.setFillColor(VERT)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(gauche, _y(16), TITRE_ONEC)
    c.setFillColor(ENCRE)
    c.setFont("Times-BoldItalic", 11)
    c.drawString(gauche, _y(22.5), identite.nom)
    if identite.contact:
        c.setFillColor(GRIS_TEXTE)
        c.setFont("Helvetica", 7.5)
        c.drawString(gauche, _y(28), identite.contact)

    c.setStrokeColor(VERT)
    c.setLineWidth(1.2 * mm)
    c.line(MARGE, _y(8 + bande + 1.5), LARGEUR - MARGE, _y(8 + bande + 1.5))

    c.setFillColor(ENCRE)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(MARGE, _y(44), titre.upper())
    c.setFillColor(GRIS_TEXTE)
    c.setFont("Helvetica", 9)
    c.drawString(MARGE, _y(50), sous_titre)


def _cartouches(c: rl_canvas.Canvas, journal: JournalAnnexe, y_haut: float) -> None:
    """`drawSummary` : solde d'ouverture, flux, solde de clôture."""
    ecart = 4
    largeur = (LARGEUR / mm - 2 * 12 - ecart * 3) / 4
    hauteur = 17
    cartes = [
        ("Solde initial", journal.solde_initial, _rgb(241, 245, 249)),
        ("Total entrées", journal.total_entrees, _rgb(236, 253, 245)),
        ("Total sorties", journal.total_sorties, _rgb(254, 242, 242)),
        ("Solde final", journal.solde_final, VERT),
    ]
    for i, (libelle, valeur, fond) in enumerate(cartes):
        x = 12 + i * (largeur + ecart)
        final = i == len(cartes) - 1
        c.setFillColor(fond)
        c.roundRect(x * mm, _y(y_haut + hauteur), largeur * mm, hauteur * mm, 2 * mm, stroke=0, fill=1)
        c.setFillColor(_rgb(220, 252, 231) if final else GRIS_TEXTE)
        c.setFont("Helvetica", 7)
        c.drawString((x + 3) * mm, _y(y_haut + 6), libelle.upper())
        c.setFillColor(colors.white if final else ENCRE)
        c.setFont("Helvetica-Bold", 10)
        c.drawString((x + 3) * mm, _y(y_haut + 13), f"{_montant(valeur)} {journal.devise}")


class _Signatures(Flowable):
    """`drawSignatures` : trois cadres, gardés ensemble sur une même page."""

    def __init__(self) -> None:
        super().__init__()
        self.width = LARGEUR - 2 * MARGE
        self.height = 26 * mm

    def draw(self) -> None:
        c = self.canv
        ecart = 5 * mm
        largeur = (self.width - 2 * ecart) / 3
        for i, titre in enumerate(("Établi par (Caissier)", "Vérifié par (Comptabilité)", "Approuvé par (Direction)")):
            x = i * (largeur + ecart)
            c.setStrokeColor(TRAIT_CADRE)
            c.setLineWidth(0.3 * mm)
            c.roundRect(x, 0, largeur, self.height, 2 * mm, stroke=1, fill=0)
            c.setFillColor(GRIS_TEXTE)
            c.setFont("Helvetica-Bold", 7.5)
            c.drawString(x + 3 * mm, self.height - 5.5 * mm, titre)
            c.setStrokeColor(TRAIT)
            c.line(x + 3 * mm, 7 * mm, x + largeur - 3 * mm, 7 * mm)


def _canvas_numerote(identite: Identite, emis: str):
    """Canvas qui connaît le nombre total de pages pour le pied « Page n / N »."""

    class CanvasNumerote(rl_canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages: list[dict] = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for etat in self._pages:
                self.__dict__.update(etat)
                self._pied(total)
                super().showPage()
            super().save()

        def _pied(self, total: int) -> None:
            """`drawTenantFooter` : l'organisation sur chaque page."""
            self.setStrokeColor(TRAIT)
            self.setLineWidth(0.3 * mm)
            self.line(MARGE, 14 * mm, LARGEUR - MARGE, 14 * mm)
            self.setFillColor(GRIS_PIED)
            self.setFont("Helvetica", 7.5)
            self.drawString(MARGE, 9.5 * mm, identite.nom)
            self.drawCentredString(LARGEUR / 2, 9.5 * mm, emis)
            self.drawRightString(LARGEUR - MARGE, 9.5 * mm, f"Page {self._pageNumber} / {total}")

    return CanvasNumerote


def _pdf(journal: JournalAnnexe, identite: Identite, debut: datetime, fin_incluse: datetime, tz: ZoneInfo) -> bytes:
    titre = f"Journal de trésorerie — {journal.nom_compte}"
    sous_titre = f"Période du {debut:%d/%m/%Y} au {fin_incluse:%d/%m/%Y}  •  Devise : {journal.devise}"
    emis = f"Émis le {datetime.now(tz):%d/%m/%Y à %H:%M} — envoi automatique"

    # Première page : le tableau commence sous l'en-tête (56 mm) et les
    # cartouches (17 mm + 6) ; les suivantes à 20 mm du haut, comme autoTable.
    haut_premiere = 56 + 17 + 6
    tampon = io.BytesIO()
    doc = BaseDocTemplate(
        tampon,
        pagesize=A4,
        leftMargin=MARGE,
        rightMargin=MARGE,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=titre,
        author=identite.nom,
    )
    largeur_utile = LARGEUR - 2 * MARGE
    premiere = Frame(MARGE, 20 * mm, largeur_utile, HAUTEUR - haut_premiere * mm - 20 * mm, id="premiere", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    suivantes = Frame(MARGE, 20 * mm, largeur_utile, HAUTEUR - 40 * mm, id="suivantes", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)

    def sur_premiere(c, _doc):
        _entete(c, identite, titre, sous_titre)
        _cartouches(c, journal, 56)

    doc.addPageTemplates(
        [
            PageTemplate(id="Premiere", frames=[premiere], onPage=sur_premiere, autoNextPageTemplate="Suivantes"),
            PageTemplate(id="Suivantes", frames=[suivantes]),
        ]
    )

    cellule = ParagraphStyle("cellule", fontName="Helvetica", fontSize=8.5, leading=10, textColor=ENCRE)
    devise = journal.devise
    donnees: list[list] = [["Date", "Libellé / Référence", f"Entrée ({devise})", f"Sortie ({devise})", f"Solde ({devise})"]]
    donnees.append(["", Paragraph("<b>REPORT À NOUVEAU (solde initial)</b>", cellule), "", "", _montant(journal.solde_initial)])
    # Mises en forme ligne à ligne : couleur des flux, filet de changement de jour.
    styles_lignes: list[tuple] = []
    jour_precedent = None
    for ligne in journal.lignes:
        date = ligne.date
        if date is not None and date.tzinfo is not None:
            date = date.astimezone(tz)
        entree = Decimal(ligne.entree or 0)
        sortie = Decimal(ligne.sortie or 0)
        rang = len(donnees)
        styles_lignes.append(("TEXTCOLOR", (2, rang), (2, rang), ENTREE if entree > 0 else TIRET))
        styles_lignes.append(("TEXTCOLOR", (3, rang), (3, rang), SORTIE if sortie > 0 else TIRET))
        jour = date.date() if date else None
        if jour_precedent is not None and jour != jour_precedent:
            styles_lignes.append(("LINEABOVE", (0, rang), (-1, rang), 0.45 * mm, FILET_JOUR))
        jour_precedent = jour
        donnees.append(
            [
                date.strftime("%d/%m/%Y") if date else "",
                Paragraph(_echapper(_libelle_complet(ligne)), cellule),
                _montant(entree) if entree > 0 else "—",
                _montant(sortie) if sortie > 0 else "—",
                _montant(ligne.solde),
            ]
        )
    blanc = ParagraphStyle("total", parent=cellule, textColor=colors.white)
    donnees.append(
        [
            "",
            Paragraph("<b>TOTAUX DE LA PÉRIODE</b>", blanc),
            _montant(journal.total_entrees),
            _montant(journal.total_sorties),
            _montant(journal.solde_final),
        ]
    )
    derniere = len(donnees) - 1

    table = Table(
        donnees,
        colWidths=[21 * mm, largeur_utile - (21 + 27 + 27 + 30) * mm, 27 * mm, 27 * mm, 30 * mm],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("TEXTCOLOR", (0, 1), (-1, -1), ENCRE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 1), (-1, -1), 1.4 * mm),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 1.4 * mm),
                ("LEFTPADDING", (0, 0), (-1, -1), 2 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2 * mm),
                # En-tête
                ("BACKGROUND", (0, 0), (-1, 0), VERT),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("TOPPADDING", (0, 0), (-1, 0), 2.5 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 2.5 * mm),
                # Corps
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGNE_PAIRE]),
                ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
                ("FONTNAME", (4, 1), (4, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.15 * mm, TRAIT),
                # Report à nouveau
                ("BACKGROUND", (0, 1), (-1, 1), FOND_REPORT),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                # Totaux
            ]
            + styles_lignes
            + [
                ("BACKGROUND", (0, derniere), (-1, derniere), VERT),
                ("TEXTCOLOR", (0, derniere), (-1, derniere), colors.white),
                ("FONTNAME", (0, derniere), (-1, derniere), "Helvetica-Bold"),
            ]
        )
    )

    doc.build(
        [table, KeepTogether([Spacer(1, 10 * mm), _Signatures()])],
        canvasmaker=_canvas_numerote(identite, emis),
    )
    return tampon.getvalue()


def _nom_fichier(organisation: str, nom_compte: str, fin_incluse: datetime, extension: str = "pdf") -> str:
    """Même nom que le fichier tiré de l'écran : Journal_<org>_<compte>_<jj-mm-aaaa>.<ext>."""
    org = re.sub(r"[^\w\-]+", "_", organisation or "organisation")
    compte = re.sub(r"[^\w\-]+", "_", nom_compte or "caisse")
    return f"Journal_{org}_{compte}_{fin_incluse:%d-%m-%Y}.{extension}"


# ── Excel : le modèle de l'export « Journal » de l'écran Rapports ───────────
# Reproduit `createReportListSheet` (frontend/src/utils/reportExcelStyles.ts)
# tel que `exportJournalToExcel` (Rapports.tsx) l'appelle : bandeau, cinq
# cartouches, tableau N° / Date / Libellé / Entrée / Sortie / Solde, ligne
# TOTAL MOUVEMENTS, en-tête figé, filtre, impression paysage.

XL_VERT = "065F46"
XL_VERT_FONCE = "064E3B"
XL_VERT_CLAIR = "D1FAE5"
XL_AMBRE_CLAIR = "FEF3C7"
XL_ROUGE_CLAIR = "FEE2E2"
XL_ARDOISE = "334155"
XL_ARDOISE_CLAIR = "F8FAFC"
XL_ROUGE = "DC2626"
XL_AMBRE = "B45309"
XL_BORDURE = "D1D5DB"
XL_FORMAT_MONTANT = "#,##0.00;[Red]-#,##0.00"
XL_FORMAT_DATE = "dd/mm/yyyy"

# Tons des cartouches : (fond, couleur de la valeur), comme `toneFills` / `toneValueColors`.
XL_TONS = {
    "default": (XL_ARDOISE_CLAIR, XL_ARDOISE),
    "positive": (XL_VERT_CLAIR, XL_VERT_FONCE),
    "negative": (XL_ROUGE_CLAIR, XL_ROUGE),
    "warning": (XL_AMBRE_CLAIR, XL_AMBRE),
}


def _xlsx(journal: JournalAnnexe, identite: Identite, debut: datetime, fin_incluse: datetime, tz: ZoneInfo) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    fin_trait = Side(style="thin", color=XL_BORDURE)
    bordure = Border(left=fin_trait, right=fin_trait, top=fin_trait, bottom=fin_trait)

    def remplir(couleur: str) -> PatternFill:
        return PatternFill("solid", fgColor=couleur)

    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    nb_col = 6
    derniere_lettre = get_column_letter(nb_col)
    periode = f"du {debut:%d/%m/%Y} au {fin_incluse:%d/%m/%Y}"
    genere = datetime.now(tz).strftime("%d/%m/%Y à %H:%M")
    canal = "CAISSE" if journal.nom_compte.startswith("Caisse") else "BANQUE"

    # Bandeau : titre, organisation, sous-titre.
    bandeau = [
        (f"JOURNAL DE TRÉSORERIE — {journal.nom_compte}", Font(bold=True, color="FFFFFF", size=16), remplir(XL_VERT), 28),
        (identite.nom, Font(bold=True, color=XL_VERT_FONCE, size=11), None, 20),
        (
            f"{periode} | Canal : {canal} | Devise : {journal.devise} | Généré le {genere} — envoi automatique",
            Font(italic=True, color=XL_ARDOISE, size=10),
            None,
            24,
        ),
    ]
    for rang, (texte, police, fond, hauteur) in enumerate(bandeau, start=1):
        ws.merge_cells(f"A{rang}:{derniere_lettre}{rang}")
        cellule = ws.cell(row=rang, column=1, value=texte)
        for col in range(1, nb_col + 1):
            c = ws.cell(row=rang, column=col)
            c.font = police
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=rang == 3)
            if fond:
                c.fill = fond
        ws.row_dimensions[rang].height = hauteur

    # Cartouches (lignes 4 et 5), réparties sur les six colonnes comme `cardBounds`.
    total_entrees = sum((Decimal(l.entree or 0) for l in journal.lignes), Decimal("0"))
    total_sorties = sum((Decimal(l.sortie or 0) for l in journal.lignes), Decimal("0"))
    cartes = [
        ("Solde initial", float(journal.solde_initial or 0), "default", "money"),
        ("Total entrées", float(total_entrees), "positive", "money"),
        ("Total sorties", float(total_sorties), "warning", "money"),
        ("Solde final", float(journal.solde_final or 0), "negative" if (journal.solde_final or 0) < 0 else "positive", "money"),
        ("Mouvements", len(journal.lignes), "default", "integer"),
    ]
    for index, (libelle, valeur, ton, fmt) in enumerate(cartes):
        debut_col = (index * nb_col) // len(cartes) + 1
        fin_col = ((index + 1) * nb_col) // len(cartes)
        fond, couleur = XL_TONS[ton]
        for rang, contenu, police in (
            (4, libelle, Font(bold=True, color=XL_ARDOISE, size=9)),
            (5, valeur, Font(bold=True, color=couleur, size=12)),
        ):
            ws.cell(row=rang, column=debut_col, value=contenu)
            for col in range(debut_col, fin_col + 1):
                c = ws.cell(row=rang, column=col)
                c.font = police
                c.fill = remplir(fond)
                c.border = bordure
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if rang == 5:
                    c.number_format = XL_FORMAT_MONTANT if fmt == "money" else "#,##0"
            if fin_col > debut_col:
                ws.merge_cells(start_row=rang, start_column=debut_col, end_row=rang, end_column=fin_col)
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 27

    # En-tête du tableau (ligne 7).
    ligne_entete = 7
    for col, titre in enumerate(["N°", "Date", "Libellé", "Entrée", "Sortie", "Solde"], start=1):
        c = ws.cell(row=ligne_entete, column=col, value=titre)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = remplir(XL_VERT)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = bordure
    ws.row_dimensions[ligne_entete].height = 25

    police_corps = Font(color=XL_ARDOISE, size=10)
    zebre = remplir(XL_ARDOISE_CLAIR)
    for index, ligne in enumerate(journal.lignes):
        rang = ligne_entete + 1 + index
        date = ligne.date
        if date is not None and date.tzinfo is not None:
            date = date.astimezone(tz)
        valeurs = [
            index + 1,
            date.date() if date else None,
            _libelle_complet(ligne),
            float(ligne.entree or 0),
            float(ligne.sortie or 0),
            float(ligne.solde or 0),
        ]
        for col, valeur in enumerate(valeurs, start=1):
            c = ws.cell(row=rang, column=col, value=valeur)
            c.font = police_corps
            c.border = bordure
            if index % 2 == 1:
                c.fill = zebre
            if col == 1:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col == 2:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.number_format = XL_FORMAT_DATE
            elif col == 3:
                c.alignment = Alignment(vertical="center", wrap_text=True)
            else:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = XL_FORMAT_MONTANT

    # Ligne TOTAL MOUVEMENTS : « TOTAL » s'étale sur N° et Date.
    derniere_donnee = ligne_entete + len(journal.lignes)
    rang_total = derniere_donnee + 1
    for col, valeur in enumerate(
        ["TOTAL MOUVEMENTS", "", "", float(total_entrees), float(total_sorties), float(journal.solde_final or 0)],
        start=1,
    ):
        c = ws.cell(row=rang_total, column=col, value=valeur if valeur != "" else None)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = remplir(XL_VERT)
        c.border = bordure
        c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="right" if col >= 4 else None)
        if col >= 4:
            c.number_format = XL_FORMAT_MONTANT
    ws.merge_cells(start_row=rang_total, start_column=1, end_row=rang_total, end_column=2)
    ws.row_dimensions[rang_total].height = 24

    for col, largeur in enumerate([6, 16, 58, 18, 18, 18], start=1):
        ws.column_dimensions[get_column_letter(col)].width = largeur

    # Mise en page : en-tête figé, filtre, quadrillage masqué, onglet vert,
    # impression paysage ajustée à la largeur, en-tête répété à chaque page.
    ws.freeze_panes = f"A{ligne_entete + 1}"
    ws.auto_filter.ref = f"A{ligne_entete}:{derniere_lettre}{max(derniere_donnee, ligne_entete)}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = XL_VERT
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.35
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.page_margins.header = ws.page_margins.footer = 0.2
    ws.print_title_rows = f"{ligne_entete}:{ligne_entete}"

    wb.properties.title = f"Journal de trésorerie — {journal.nom_compte}"
    wb.properties.subject = periode
    wb.properties.creator = identite.nom

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


async def construire_annexes(
    db: AsyncSession,
    *,
    tenant_id: int,
    organisation: str,
    start: datetime,
    end: datetime,
    tz: ZoneInfo,
    format: str = "pdf",
) -> list[tuple[str, bytes]]:
    """(nom de fichier, contenu) pour chaque journal de la période [start, end[.

    `format` : "pdf" (le PDF du journal de l'écran) ou "xlsx" (son export Excel).
    """
    fin_incluse = end - timedelta(days=1)
    identite = await _identite(db, tenant_id, organisation)
    journaux = await _journaux(db, tenant_id, start.date().isoformat(), fin_incluse.date().isoformat())
    annexes: list[tuple[str, bytes]] = []
    for journal in journaux:
        # Le tracé est du CPU pur : hors de la boucle d'événements.
        fabrique = _xlsx if format == "xlsx" else _pdf
        contenu = await anyio.to_thread.run_sync(fabrique, journal, identite, start, fin_incluse, tz)
        annexes.append((_nom_fichier(identite.nom, journal.nom_compte, fin_incluse, format), contenu))
    return annexes
