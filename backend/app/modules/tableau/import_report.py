"""Fichier annoté d'un import : les lignes du classeur, celles à compléter en rouge.

Chaque ligne garde son numéro d'origine, pour que « Ligne 6 » du rapport d'import
désigne bien la ligne 6 du fichier. Une fois complété, le fichier se réimporte : la
colonne « Signalement » n'est reconnue par aucun parseur.
"""
from __future__ import annotations

import io
import re
from typing import Any, Iterable

from .source_adapters import _header_map

COLONNE_SIGNALEMENT = "Signalement"
FEUILLE_SIGNALEMENTS = "Signalements"

# Les champs signalés qui ne portent pas le nom du champ de leur colonne.
_ALIAS_CHAMPS = {
    "nom": {"membre"},
    "societe": {"membre"},
    "nif": {"nif", "numero_impot"},
    "numero_impot": {"numero_impot", "nif"},
    "cabinet": {"cabinet_attache"},
    "employeur": {"nom_employeur"},
}

_DATE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:T00:00:00)?$")


def nom_fichier_annote(file_name: str | None) -> str:
    base = re.sub(r"\.xlsx?$", "", file_name or "import", flags=re.IGNORECASE)
    base = re.sub(r"[^A-Za-z0-9_-]+", "_", base).strip("_") or "import"
    return f"{base}_annote.xlsx"


def _valeur(value: Any) -> Any:
    """Les dates ont été conservées en ISO : on les rend lisibles, et réimportables."""
    if isinstance(value, str):
        match = _DATE_ISO.match(value)
        if match:
            return f"{match.group(3)}/{match.group(2)}/{match.group(1)}"
    return value


def _message(erreur: dict) -> str:
    return str(erreur.get("message") or "Information à vérifier")


def _est_colonne_signalement(colonne: Any) -> bool:
    return str(colonne or "").strip().casefold() == COLONNE_SIGNALEMENT.casefold()


def _colonnes_signalees(colonnes: list[str], erreurs: Iterable[dict]) -> set[int]:
    """Index (base 0) des colonnes visées : c'est la cellule à compléter."""
    champs_colonnes = _header_map(colonnes)
    index: set[int] = set()
    for erreur in erreurs:
        champ = str(erreur.get("champ") or "")
        attendus = _ALIAS_CHAMPS.get(champ, {champ})
        index.update(i for i, colonne in enumerate(colonnes) if champs_colonnes.get(colonne) in attendus)
    return index


def build_import_report(source_rows: Iterable[Any], colonnes_source: dict[str, list[str]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    entete_fill = PatternFill("solid", fgColor="1F4E78")
    entete_font = Font(bold=True, color="FFFFFF")
    ligne_fill = PatternFill("solid", fgColor="FFC7CE")
    cellule_fill = PatternFill("solid", fgColor="FF7C80")
    rouge_font = Font(color="9C0006")

    par_feuille: dict[str, list[Any]] = {}
    for row in source_rows:
        par_feuille.setdefault(row.feuille or "Feuille1", []).append(row)

    wb = Workbook()
    synthese = wb.active
    synthese.title = FEUILLE_SIGNALEMENTS
    entetes_synthese = ["Feuille", "Ligne", "Colonne", "Message"]
    synthese.append(entetes_synthese)
    for cell in synthese[1]:
        cell.fill, cell.font = entete_fill, entete_font

    for feuille, lignes in par_feuille.items():
        lignes.sort(key=lambda item: item.line_number)
        # Un fichier annoté peut être corrigé, réimporté puis annoté à nouveau.
        # Remplacer notre colonne technique évite d'en ajouter une copie à chaque cycle.
        colonnes = [
            colonne
            for colonne in (colonnes_source.get(feuille) or [])
            if not _est_colonne_signalement(colonne)
        ]
        for ligne in lignes:
            # Import antérieur à l'enregistrement de l'ordre des colonnes, ou ligne
            # portant une colonne que la première n'avait pas.
            colonnes.extend(
                cle
                for cle in ligne.raw_data
                if cle not in colonnes and not _est_colonne_signalement(cle)
            )

        titre = feuille[:31]
        ws = wb.create_sheet(f"{titre[:28]} (2)" if titre.casefold() == FEUILLE_SIGNALEMENTS.casefold() else titre)
        ligne_entete = max(lignes[0].line_number - 1, 1)
        colonne_signalement = len(colonnes) + 1
        for index, colonne in enumerate([*colonnes, COLONNE_SIGNALEMENT], start=1):
            # « __col_3 » désigne une colonne sans intitulé dans le fichier d'origine.
            cell = ws.cell(ligne_entete, index, "" if colonne.startswith("__col_") else colonne)
            cell.fill, cell.font = entete_fill, entete_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        for ligne in lignes:
            erreurs = ligne.normalization_errors or []
            # Une ligne ne peut pas s'écrire sur l'en-tête, même dans un import ancien.
            numero = max(ligne.line_number, ligne_entete + 1)
            for index, colonne in enumerate(colonnes, start=1):
                ws.cell(numero, index, _valeur(ligne.raw_data.get(colonne)))
            if not erreurs:
                continue
            messages = list(dict.fromkeys(_message(erreur) for erreur in erreurs))
            ws.cell(numero, colonne_signalement, " ; ".join(messages))
            a_completer = _colonnes_signalees(colonnes, erreurs)
            for index in range(1, colonne_signalement + 1):
                cell = ws.cell(numero, index)
                cell.fill = cellule_fill if index - 1 in a_completer else ligne_fill
                cell.font = rouge_font
            for erreur in erreurs:
                synthese.append([ws.title, ligne.line_number, erreur.get("champ"), _message(erreur)])
                for cell in synthese[synthese.max_row]:
                    cell.fill, cell.font = ligne_fill, rouge_font

        for index, colonne in enumerate(colonnes, start=1):
            valeurs = (ligne.raw_data.get(colonne) for ligne in lignes)
            largeur = max([len(str(colonne)), *(len(str(v)) for v in valeurs if v is not None)])
            ws.column_dimensions[get_column_letter(index)].width = min(largeur + 2, 40)
        ws.column_dimensions[get_column_letter(colonne_signalement)].width = 50
        ws.freeze_panes = ws.cell(ligne_entete + 1, 1)
        ws.auto_filter.ref = (
            f"A{ligne_entete}:{get_column_letter(colonne_signalement)}{ws.max_row}"
        )

    if synthese.max_row == 1:
        synthese.append(["", "", "", "Aucune ligne à compléter dans cet import."])
    for lettre, largeur in zip("ABCD", (24, 8, 20, 70)):
        synthese.column_dimensions[lettre].width = largeur
    synthese.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
