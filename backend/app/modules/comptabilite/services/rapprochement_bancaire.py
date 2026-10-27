"""Import déterministe de relevés et propositions de rapprochement.

Une proposition n'est jamais validée automatiquement. Même un score de 100
reste au statut PROPOSE jusqu'à la décision explicite d'un utilisateur.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from uuid import UUID

from fastapi import HTTPException, status
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.compte_bancaire import CompteBancaire
from app.models.encaissement import Encaissement
from app.models.payment_history import PaymentHistory
from app.models.sortie_fonds import SortieFonds
from app.models.transfert_interne import TransfertInterne
from app.modules.comptabilite.models import (
    ComptaReleveBancaireImport,
    ComptaReleveBancaireLigne,
)
from app.modules.comptabilite.schemas.rapprochement import (
    PropositionRapprochementOut,
    ReleveBancaireImportOut,
    ReleveBancaireLigneOut,
)
from app.services.encaissement_flux import PAYMENT_STATUS_ACTIVE, flux_encaissements


CENT = Decimal("0.01")
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_LINES = 10_000


def _q(value: Decimal | int | float | str | None) -> Decimal:
    return Decimal(value or 0).quantize(CENT, rounding=ROUND_HALF_UP)


def _normaliser(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


COLONNES = {
    "date_operation": {
        "date", "date operation", "date d operation", "date transaction", "date comptable",
    },
    "date_valeur": {"date valeur", "date de valeur", "valeur date"},
    "reference": {"reference", "ref", "numero", "no operation", "piece"},
    "tiers": {"tiers", "beneficiaire", "payeur", "contrepartie", "donneur ordre"},
    "libelle": {"libelle", "description", "motif", "details", "detail", "narration"},
    # Convention du relevé bancaire : débit = sortie du compte, crédit = entrée.
    "debit": {"debit", "montant debit", "sortie", "retrait", "withdrawal"},
    "credit": {"credit", "montant credit", "entree", "depot", "deposit"},
    "solde": {"solde", "balance", "solde apres operation"},
}


# Lignes de récapitulatif que les banques glissent dans leurs relevés : sans
# montant, elles ne sont pas des mouvements. On les écarte au lieu de rejeter
# tout le fichier ; une ligne sans montant qui n'en est pas une reste une erreur.
_RECAPITULATIFS = (
    "solde", "ancien solde", "nouveau solde", "solde initial", "solde final",
    "solde d ouverture", "solde de cloture", "total", "totaux", "report",
    "a reporter", "opening balance", "closing balance", "balance",
)


def _est_recapitulatif(textes: tuple[object, ...], debit: object, credit: object) -> bool:
    """Vrai pour une ligne sans débit ni crédit dont le libellé, ou la date
    quand la banque y écrit « Solde au 01/10 », annonce un récapitulatif."""
    if debit not in (None, "") or credit not in (None, ""):
        return False
    for valeur in textes:
        texte = _normaliser(valeur)
        if any(texte == motif or texte.startswith(motif + " ") for motif in _RECAPITULATIFS):
            return True
    return False


@dataclass(slots=True)
class LigneParse:
    numero_ligne: int
    date_operation: date
    date_valeur: date | None
    reference: str | None
    tiers: str | None
    libelle: str
    debit: Decimal
    credit: Decimal
    solde: Decimal | None
    empreinte: str


@dataclass(slots=True)
class Candidat:
    operation_type: str
    operation_id: str
    date_operation: date
    reference: str | None
    libelle: str
    tiers: str | None
    debit: Decimal
    credit: Decimal
    score: Decimal = Decimal("0")

    def as_dict(self) -> dict[str, Any]:
        return {
            "operation_type": self.operation_type,
            "operation_id": self.operation_id,
            "date_operation": self.date_operation.isoformat(),
            "reference": self.reference,
            "libelle": self.libelle,
            "tiers": self.tiers,
            "debit": f"{self.debit:.2f}",
            "credit": f"{self.credit:.2f}",
            "score": f"{self.score:.2f}",
        }


def _mapping_entetes(values: list[object]) -> dict[str, int]:
    normalisees = [_normaliser(value) for value in values]
    mapping: dict[str, int] = {}
    for cible, synonymes in COLONNES.items():
        for index, nom in enumerate(normalisees):
            if nom in synonymes:
                mapping[cible] = index
                break
    requis = {"date_operation", "libelle", "debit", "credit"}
    manquantes = sorted(requis - set(mapping))
    if manquantes:
        noms = ", ".join(manquantes)
        raise ValueError(
            f"Colonnes obligatoires absentes : {noms}. Le relevé doit distinguer Débit et Crédit."
        )
    return mapping


def _date(value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            parsed = from_excel(value)
            return parsed.date() if isinstance(parsed, datetime) else parsed
        except (TypeError, ValueError, OverflowError):
            return None
    raw = str(value).strip()
    if not raw:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(raw).date()
    except ValueError:
        return None


def _montant(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, Decimal):
        return _q(abs(value))
    if isinstance(value, (int, float)):
        return _q(abs(Decimal(str(value))))
    raw = str(value).strip()
    if not raw or raw.lower() in {"nan", "none", "-"}:
        return Decimal("0")
    negatif = raw.startswith("(") and raw.endswith(")")
    raw = raw.strip("()").replace("\u00a0", "").replace(" ", "").replace("'", "")
    raw = re.sub(r"[^0-9,.-]", "", raw)
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        parsed = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"Montant illisible : {value}") from exc
    if negatif:
        parsed = -parsed
    return _q(abs(parsed))


def _solde(value: object) -> Decimal:
    """Lit un solde signé (un découvert ne doit jamais devenir positif)."""
    if isinstance(value, (int, float, Decimal)):
        return _q(Decimal(str(value)))
    raw = str(value or "").strip()
    negatif = raw.startswith("(") and raw.endswith(")")
    raw = raw.strip("()").replace("\u00a0", "").replace(" ", "").replace("'", "")
    raw = re.sub(r"[^0-9,.-]", "", raw)
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        parsed = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"Solde illisible : {value}") from exc
    return _q(-abs(parsed) if negatif else parsed)


def _lire_csv(content: bytes) -> list[list[object]]:
    text: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Encodage CSV non reconnu.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _lire_xlsx(content: bytes) -> list[list[object]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Classeur Excel illisible ou protégé.") from exc
    worksheet = workbook.active
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def parser_releve(content: bytes, filename: str) -> list[LigneParse]:
    if not content:
        raise ValueError("Le fichier est vide.")
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("Le fichier dépasse la limite de 10 Mo.")
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        rows = _lire_csv(content)
    elif suffix == ".xlsx":
        rows = _lire_xlsx(content)
    else:
        raise ValueError("Format non pris en charge : utilisez CSV ou XLSX.")
    if not rows:
        raise ValueError("Le relevé ne contient aucune ligne.")

    header_index: int | None = None
    mapping: dict[str, int] | None = None
    derniere_erreur = "En-tête introuvable."
    for index, row in enumerate(rows[:20]):
        try:
            mapping = _mapping_entetes(row)
            header_index = index
            break
        except ValueError as exc:
            derniere_erreur = str(exc)
    if mapping is None or header_index is None:
        raise ValueError(derniere_erreur)

    parsed: list[LigneParse] = []
    erreurs: list[str] = []
    for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        def value(cle: str):
            idx = mapping.get(cle)
            return row[idx] if idx is not None and idx < len(row) else None

        if not any(cell not in (None, "") for cell in row):
            continue
        if _est_recapitulatif(
            (value("libelle"), value("date_operation")), value("debit"), value("credit")
        ):
            continue
        try:
            date_operation = _date(value("date_operation"))
            if date_operation is None:
                raise ValueError("date d'opération invalide")
            libelle = " ".join(str(value("libelle") or "").strip().split())
            if not libelle:
                raise ValueError("libellé absent")
            debit = _montant(value("debit"))
            credit = _montant(value("credit"))
            if (debit > 0 and credit > 0) or (debit == 0 and credit == 0):
                raise ValueError("une ligne doit porter soit un débit, soit un crédit")
            reference = " ".join(str(value("reference") or "").strip().split()) or None
            tiers = " ".join(str(value("tiers") or "").strip().split()) or None
            solde_raw = value("solde")
            solde = _solde(solde_raw) if solde_raw not in (None, "") else None
            date_valeur = _date(value("date_valeur"))
            empreinte_payload = "|".join(
                (
                    date_operation.isoformat(),
                    date_valeur.isoformat() if date_valeur else "",
                    _normaliser(reference),
                    _normaliser(libelle),
                    f"{debit:.2f}",
                    f"{credit:.2f}",
                )
            )
            parsed.append(
                LigneParse(
                    numero_ligne=row_index,
                    date_operation=date_operation,
                    date_valeur=date_valeur,
                    reference=reference,
                    tiers=tiers,
                    libelle=libelle,
                    debit=debit,
                    credit=credit,
                    solde=solde,
                    empreinte=hashlib.sha256(empreinte_payload.encode("utf-8")).hexdigest(),
                )
            )
        except ValueError as exc:
            erreurs.append(f"Ligne {row_index}: {exc}")
    if erreurs:
        raise ValueError("Le relevé contient des lignes invalides : " + " | ".join(erreurs[:20]))
    if not parsed:
        raise ValueError("Aucun mouvement bancaire exploitable n'a été trouvé.")
    if len(parsed) > MAX_LINES:
        raise ValueError(f"Le relevé dépasse la limite de {MAX_LINES} mouvements.")
    return parsed


def scorer(ligne: LigneParse, candidat: Candidat) -> Decimal:
    if ligne.debit != candidat.debit or ligne.credit != candidat.credit:
        return Decimal("0")
    ecart_jours = abs((ligne.date_operation - candidat.date_operation).days)
    if ecart_jours > 7:
        return Decimal("0")
    score = Decimal("65")
    score += max(Decimal("0"), Decimal("20") - Decimal(ecart_jours * 3))
    ref_ligne = _normaliser(ligne.reference)
    ref_candidat = _normaliser(candidat.reference)
    if ref_ligne and ref_candidat:
        if ref_ligne == ref_candidat:
            score += Decimal("10")
        elif ref_ligne in ref_candidat or ref_candidat in ref_ligne:
            score += Decimal("6")
    texte_ligne = _normaliser(f"{ligne.tiers or ''} {ligne.libelle}")
    texte_candidat = _normaliser(f"{candidat.tiers or ''} {candidat.libelle}")
    if texte_ligne and texte_candidat:
        score += Decimal(str(round(SequenceMatcher(None, texte_ligne, texte_candidat).ratio() * 5, 2)))
    return min(Decimal("100"), score.quantize(Decimal("0.01")))


async def _operations_rapprochees(
    db: AsyncSession, *, organisation_id: int, compte_bancaire_id: int
) -> dict[str, set[str]]:
    """Opérations sans drapeau de pointage déjà confirmées sur ce compte."""
    rows = (
        await db.execute(
            select(ComptaReleveBancaireLigne.operation_type, ComptaReleveBancaireLigne.operation_id).where(
                ComptaReleveBancaireLigne.organisation_id == organisation_id,
                ComptaReleveBancaireLigne.compte_bancaire_id == compte_bancaire_id,
                ComptaReleveBancaireLigne.statut_rapprochement == "RAPPROCHE",
                ComptaReleveBancaireLigne.operation_type.in_(("TRANSFERT", "VERSEMENT")),
            )
        )
    ).all()
    rapproches: dict[str, set[str]] = {"TRANSFERT": set(), "VERSEMENT": set()}
    for operation_type, operation_id in rows:
        rapproches[operation_type].add(operation_id)
    return rapproches


async def _candidats(
    db: AsyncSession,
    *,
    organisation_id: int,
    compte_bancaire_id: int,
    devise: str,
    date_min: date,
    date_max: date,
) -> list[Candidat]:
    debut = datetime.combine(date_min - timedelta(days=7), datetime.min.time(), tzinfo=timezone.utc)
    fin = datetime.combine(date_max + timedelta(days=8), datetime.min.time(), tzinfo=timezone.utc)
    result: list[Candidat] = []
    # Encaissements et sorties portent leur propre drapeau de pointage. Les
    # transferts internes et les versements n'en ont pas : la ligne de relevé
    # validée est donc la source de vérité pour empêcher de rapprocher deux fois
    # le même transfert, ou le même versement, sur un même compte bancaire.
    deja_rapproches = await _operations_rapprochees(
        db, organisation_id=organisation_id, compte_bancaire_id=compte_bancaire_id
    )

    flux = flux_encaissements(organisation_id)
    enc_rows = (
        await db.execute(
            select(
                flux.c.encaissement_id,
                flux.c.versement_id,
                flux.c.date_flux,
                flux.c.reference,
                flux.c.montant,
                Encaissement.libelle,
                Encaissement.client_nom,
            )
            .join(Encaissement, Encaissement.id == flux.c.encaissement_id)
            .where(
                flux.c.canal == "BANQUE",
                flux.c.compte_bancaire_id == compte_bancaire_id,
                flux.c.devise == devise,
                flux.c.date_flux >= debut,
                flux.c.date_flux < fin,
                Encaissement.is_reconciled.is_(False),
            )
        )
    ).all()
    for row in enc_rows:
        # Une note payée en plusieurs virements donne une ligne par versement :
        # chacun se rapproche de sa propre ligne de relevé. Seule une note sans
        # versement détaillé se rapproche en bloc.
        if row.versement_id is not None:
            operation_type, operation_id = "VERSEMENT", str(row.versement_id)
            if operation_id in deja_rapproches["VERSEMENT"]:
                continue
        else:
            operation_type, operation_id = "ENCAISSEMENT", str(row.encaissement_id)
        result.append(
            Candidat(
                operation_type=operation_type,
                operation_id=operation_id,
                date_operation=row.date_flux.date(),
                reference=row.reference,
                libelle=row.libelle,
                tiers=row.client_nom,
                debit=Decimal("0"),
                credit=_q(row.montant),
            )
        )

    paiement_ts = func.coalesce(SortieFonds.date_paiement, SortieFonds.created_at)
    sortie_rows = (
        await db.execute(
            select(
                SortieFonds.id,
                paiement_ts.label("date_operation"),
                SortieFonds.reference_numero,
                SortieFonds.reference,
                SortieFonds.motif,
                SortieFonds.beneficiaire,
                SortieFonds.montant_paye,
                SortieFonds.type_sortie,
                SortieFonds.canal,
            ).where(
                SortieFonds.organisation_id == organisation_id,
                SortieFonds.compte_bancaire_id == compte_bancaire_id,
                SortieFonds.devise == devise,
                paiement_ts >= debut,
                paiement_ts < fin,
                or_(SortieFonds.statut.is_(None), SortieFonds.statut == "VALIDE"),
                SortieFonds.is_reconciled.is_(False),
            )
        )
    ).all()
    for row in sortie_rows:
        incoming = row.type_sortie == "versement_banque"
        if not incoming and row.canal != "BANQUE":
            continue
        result.append(
            Candidat(
                operation_type="SORTIE",
                operation_id=str(row.id),
                date_operation=row.date_operation.date(),
                reference=row.reference_numero or row.reference,
                libelle=row.motif,
                tiers=row.beneficiaire,
                debit=Decimal("0") if incoming else _q(row.montant_paye),
                credit=_q(row.montant_paye) if incoming else Decimal("0"),
            )
        )

    transfer_rows = (
        await db.execute(
            select(TransfertInterne).where(
                TransfertInterne.organisation_id == organisation_id,
                TransfertInterne.devise == devise,
                TransfertInterne.date_transfert >= debut,
                TransfertInterne.date_transfert < fin,
                or_(
                    TransfertInterne.source_id == compte_bancaire_id,
                    TransfertInterne.destination_id == compte_bancaire_id,
                ),
            )
        )
    ).scalars().all()
    for transfer in transfer_rows:
        if str(transfer.id) in deja_rapproches["TRANSFERT"]:
            continue
        incoming = transfer.destination_type == "BANQUE" and transfer.destination_id == compte_bancaire_id
        outgoing = transfer.source_type == "BANQUE" and transfer.source_id == compte_bancaire_id
        if not incoming and not outgoing:
            continue
        result.append(
            Candidat(
                operation_type="TRANSFERT",
                operation_id=str(transfer.id),
                date_operation=transfer.date_transfert.date(),
                reference=transfer.reference,
                libelle="Transfert interne",
                tiers=None,
                debit=_q(transfer.montant) if outgoing else Decimal("0"),
                credit=_q(transfer.montant) if incoming else Decimal("0"),
            )
        )
    return result


async def importer_releve(
    db: AsyncSession,
    *,
    organisation_id: int,
    user_id: UUID | None,
    compte_bancaire_id: int,
    filename: str,
    content: bytes,
) -> ComptaReleveBancaireImport:
    compte = (
        await db.execute(
            select(CompteBancaire).where(
                CompteBancaire.id == compte_bancaire_id,
                CompteBancaire.organisation_id == organisation_id,
                CompteBancaire.account_type == "BANK",
                CompteBancaire.is_active.is_(True),
            )
        )
    ).scalar_one_or_none()
    if compte is None:
        raise HTTPException(status_code=404, detail="Compte bancaire introuvable ou inactif.")
    try:
        lignes_parsees = parser_releve(content, filename)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    empreinte_fichier = hashlib.sha256(content).hexdigest()
    duplicate_file = (
        await db.execute(
            select(ComptaReleveBancaireImport.id).where(
                ComptaReleveBancaireImport.organisation_id == organisation_id,
                ComptaReleveBancaireImport.compte_bancaire_id == compte_bancaire_id,
                ComptaReleveBancaireImport.empreinte_fichier == empreinte_fichier,
            )
        )
    ).scalar_one_or_none()
    if duplicate_file:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Ce fichier a déjà été importé ({duplicate_file}).",
        )

    candidats = await _candidats(
        db,
        organisation_id=organisation_id,
        compte_bancaire_id=compte_bancaire_id,
        devise=compte.devise,
        date_min=min(l.date_operation for l in lignes_parsees),
        date_max=max(l.date_operation for l in lignes_parsees),
    )
    empreintes = {ligne.empreinte for ligne in lignes_parsees}
    existantes = (
        await db.execute(
            select(ComptaReleveBancaireLigne).where(
                ComptaReleveBancaireLigne.organisation_id == organisation_id,
                ComptaReleveBancaireLigne.compte_bancaire_id == compte_bancaire_id,
                ComptaReleveBancaireLigne.empreinte_ligne.in_(empreintes),
            )
        )
    ).scalars().all()
    doublons_par_empreinte = {ligne.empreinte_ligne: ligne for ligne in existantes}

    releve = ComptaReleveBancaireImport(
        organisation_id=organisation_id,
        compte_bancaire_id=compte_bancaire_id,
        devise=compte.devise,
        nom_fichier=(filename or "releve")[:255],
        format=Path(filename or "").suffix.lstrip(".").upper(),
        empreinte_fichier=empreinte_fichier,
        statut="IMPORTE",
        importe_par=user_id,
    )
    db.add(releve)
    await db.flush()

    # Doublons d'abord : une ligne déjà vue (dans un relevé précédent ou plus
    # haut dans celui-ci) n'entre pas dans la répartition.
    doublon_de: dict[int, ComptaReleveBancaireLigne | int] = {}
    premiere_par_empreinte: dict[str, int] = {}
    for index, parsed in enumerate(lignes_parsees):
        existante = doublons_par_empreinte.get(parsed.empreinte)
        if existante is not None:
            doublon_de[index] = existante
        elif parsed.empreinte in premiere_par_empreinte:
            doublon_de[index] = premiere_par_empreinte[parsed.empreinte]
        else:
            premiere_par_empreinte[parsed.empreinte] = index

    scores: dict[int, list[Candidat]] = {}
    for index, parsed in enumerate(lignes_parsees):
        if index in doublon_de:
            continue
        scores[index] = sorted(
            (
                replace(candidat, score=score)
                for candidat in candidats
                if (score := scorer(parsed, candidat)) >= Decimal("65")
            ),
            key=lambda c: (c.score, c.date_operation),
            reverse=True,
        )

    # Répartition gloutonne, meilleurs scores d'abord : une opération n'est
    # proposée qu'à une seule ligne. Sans elle, deux retraits identiques du
    # même jour se verraient proposer la même sortie, et la seconde
    # confirmation échouerait.
    attribution: dict[int, tuple[str, str]] = {}
    prises: set[tuple[str, str]] = set()
    paires = sorted(
        (
            (candidat.score, index, (candidat.operation_type, candidat.operation_id))
            for index, liste in scores.items()
            for candidat in liste
        ),
        key=lambda paire: (paire[0], -paire[1]),
        reverse=True,
    )
    for _, index, cle in paires:
        if index in attribution or cle in prises:
            continue
        attribution[index] = cle
        prises.add(cle)

    lignes_creees: dict[int, ComptaReleveBancaireLigne] = {}
    doublons = 0
    propositions = 0
    for index, parsed in enumerate(lignes_parsees):
        doublon = doublon_de.get(index)
        if isinstance(doublon, int):
            doublon = lignes_creees[doublon]
        scored: list[Candidat] = []
        if doublon is None and index in attribution:
            cle = attribution[index]
            liste = scores[index]
            # L'opération attribuée en tête, puis les autres encore libres.
            scored = [c for c in liste if (c.operation_type, c.operation_id) == cle] + [
                c for c in liste if (c.operation_type, c.operation_id) not in prises
            ]
            scored = scored[:5]
        top = scored[0] if scored else None
        ambigu = bool(top and len(scored) > 1 and top.score - scored[1].score <= Decimal("5"))
        statut_ligne = "NON_RAPPROCHE"
        if doublon is not None:
            doublons += 1
        elif top is not None:
            statut_ligne = "INCERTAIN" if ambigu or top.score < Decimal("90") else "PROPOSE"
            propositions += 1
        ligne = ComptaReleveBancaireLigne(
            organisation_id=organisation_id,
            releve_import_id=releve.id,
            compte_bancaire_id=compte_bancaire_id,
            numero_ligne=parsed.numero_ligne,
            date_operation=parsed.date_operation,
            date_valeur=parsed.date_valeur,
            reference=parsed.reference,
            tiers=parsed.tiers,
            libelle=parsed.libelle,
            debit=parsed.debit,
            credit=parsed.credit,
            solde=parsed.solde,
            devise=compte.devise,
            empreinte_ligne=parsed.empreinte,
            est_doublon=doublon is not None,
            doublon_de_id=doublon.id if doublon else None,
            statut_rapprochement=statut_ligne,
            operation_type=top.operation_type if top else None,
            operation_id=top.operation_id if top else None,
            score=top.score if top else None,
            propositions=[candidate.as_dict() for candidate in scored],
        )
        db.add(ligne)
        await db.flush()
        lignes_creees[index] = ligne

    releve.nombre_lignes = len(lignes_parsees)
    releve.nombre_doublons = doublons
    releve.nombre_propositions = propositions
    await db.flush()
    return releve


def _ligne_out(ligne: ComptaReleveBancaireLigne) -> ReleveBancaireLigneOut:
    propositions = [PropositionRapprochementOut(**item) for item in (ligne.propositions or [])]
    return ReleveBancaireLigneOut(
        id=ligne.id,
        releve_import_id=ligne.releve_import_id,
        compte_bancaire_id=ligne.compte_bancaire_id,
        numero_ligne=ligne.numero_ligne,
        date_operation=ligne.date_operation,
        date_valeur=ligne.date_valeur,
        reference=ligne.reference,
        tiers=ligne.tiers,
        libelle=ligne.libelle,
        debit=ligne.debit,
        credit=ligne.credit,
        solde=ligne.solde,
        devise=ligne.devise,
        est_doublon=ligne.est_doublon,
        doublon_de_id=ligne.doublon_de_id,
        statut_rapprochement=ligne.statut_rapprochement,
        operation_type=ligne.operation_type,
        operation_id=ligne.operation_id,
        score=ligne.score,
        propositions=propositions,
        rapproche_par=ligne.rapproche_par,
        rapproche_le=ligne.rapproche_le,
    )


async def releve_out(db: AsyncSession, releve: ComptaReleveBancaireImport) -> ReleveBancaireImportOut:
    lignes = list(
        (
            await db.execute(
                select(ComptaReleveBancaireLigne)
                .where(ComptaReleveBancaireLigne.releve_import_id == releve.id)
                .order_by(ComptaReleveBancaireLigne.numero_ligne)
            )
        ).scalars().all()
    )
    return ReleveBancaireImportOut(
        id=releve.id,
        compte_bancaire_id=releve.compte_bancaire_id,
        devise=releve.devise,
        nom_fichier=releve.nom_fichier,
        format=releve.format,
        statut=releve.statut,
        nombre_lignes=releve.nombre_lignes,
        nombre_doublons=releve.nombre_doublons,
        nombre_propositions=releve.nombre_propositions,
        nombre_rapprochees=releve.nombre_rapprochees,
        importe_par=releve.importe_par,
        created_at=releve.created_at,
        lignes=[_ligne_out(ligne) for ligne in lignes],
    )


async def _recompter_releve(db: AsyncSession, releve: ComptaReleveBancaireImport) -> None:
    await db.flush()
    rows = (
        await db.execute(
            select(ComptaReleveBancaireLigne.statut_rapprochement, func.count())
            .where(
                ComptaReleveBancaireLigne.releve_import_id == releve.id,
                ComptaReleveBancaireLigne.est_doublon.is_(False),
            )
            .group_by(ComptaReleveBancaireLigne.statut_rapprochement)
        )
    ).all()
    par_statut = dict(rows)
    releve.nombre_rapprochees = par_statut.get("RAPPROCHE", 0)
    decidees = par_statut.get("RAPPROCHE", 0) + par_statut.get("IGNORE", 0)
    a_decider = releve.nombre_lignes - releve.nombre_doublons
    if decidees == 0:
        releve.statut = "IMPORTE"
    else:
        releve.statut = "RAPPROCHE" if decidees >= a_decider else "PARTIEL"


async def _encaissement_du_versement(db: AsyncSession, versement_id: str) -> Encaissement | None:
    versement = await db.get(PaymentHistory, UUID(versement_id))
    return await db.get(Encaissement, versement.encaissement_id) if versement else None


async def _versements_bancaires_tous_rapproches(
    db: AsyncSession, *, organisation_id: int, encaissement_id: UUID
) -> bool:
    """Une note payée en plusieurs virements n'est pointée qu'au dernier."""
    versements = (
        await db.execute(
            select(PaymentHistory.id).where(
                PaymentHistory.organisation_id == organisation_id,
                PaymentHistory.encaissement_id == encaissement_id,
                PaymentHistory.statut == PAYMENT_STATUS_ACTIVE,
                PaymentHistory.canal == "BANQUE",
            )
        )
    ).scalars().all()
    if not versements:
        return False
    rapproches = set(
        (
            await db.execute(
                select(ComptaReleveBancaireLigne.operation_id).where(
                    ComptaReleveBancaireLigne.organisation_id == organisation_id,
                    ComptaReleveBancaireLigne.operation_type == "VERSEMENT",
                    ComptaReleveBancaireLigne.statut_rapprochement == "RAPPROCHE",
                    ComptaReleveBancaireLigne.operation_id.in_([str(v) for v in versements]),
                )
            )
        ).scalars().all()
    )
    return all(str(v) in rapproches for v in versements)


def _pointer(source, *, user_id: UUID | None, reference: str | None) -> None:
    source.is_reconciled = True
    source.reconciled_at = datetime.now(timezone.utc)
    source.reconciled_by_id = user_id
    source.bank_statement_ref = reference


def _depointer(source) -> None:
    source.is_reconciled = False
    source.reconciled_at = None
    source.reconciled_by_id = None
    source.bank_statement_ref = None


async def _source_pointee(db: AsyncSession, operation_type: str | None, operation_id: str | None):
    """L'encaissement ou la sortie dont le drapeau de pointage suit cette ligne."""
    if not operation_id:
        return None
    if operation_type == "ENCAISSEMENT":
        return await db.get(Encaissement, UUID(operation_id))
    if operation_type == "SORTIE":
        return await db.get(SortieFonds, UUID(operation_id))
    if operation_type == "VERSEMENT":
        return await _encaissement_du_versement(db, operation_id)
    return None


async def decider_rapprochement(
    db: AsyncSession,
    *,
    organisation_id: int,
    user_id: UUID | None,
    ligne_id: UUID,
    decision: str,
    operation_type: str | None,
    operation_id: str | None,
) -> ComptaReleveBancaireLigne:
    ligne = (
        await db.execute(
            select(ComptaReleveBancaireLigne)
            .where(
                ComptaReleveBancaireLigne.id == ligne_id,
                ComptaReleveBancaireLigne.organisation_id == organisation_id,
            )
            .with_for_update()
        )
    ).scalar_one_or_none()
    if ligne is None:
        raise HTTPException(status_code=404, detail="Ligne de relevé introuvable.")
    if ligne.est_doublon:
        raise HTTPException(status_code=409, detail="Un doublon de relevé ne peut pas être rapproché.")

    if decision == "ANNULER":
        await _annuler_decision(db, ligne)
    elif ligne.statut_rapprochement in {"RAPPROCHE", "IGNORE"}:
        raise HTTPException(status_code=409, detail="Cette ligne a déjà fait l'objet d'une décision.")
    elif decision == "IGNORER":
        ligne.statut_rapprochement = "IGNORE"
        ligne.operation_type = None
        ligne.operation_id = None
        ligne.score = None
        ligne.rapproche_par = user_id
        ligne.rapproche_le = datetime.now(timezone.utc)
    else:
        candidats = await _candidats(
            db,
            organisation_id=organisation_id,
            compte_bancaire_id=ligne.compte_bancaire_id,
            devise=ligne.devise,
            date_min=ligne.date_operation,
            date_max=ligne.date_operation,
        )
        candidat = next(
            (
                item for item in candidats
                if item.operation_type == operation_type and item.operation_id == operation_id
            ),
            None,
        )
        if candidat is None or ligne.debit != candidat.debit or ligne.credit != candidat.credit:
            raise HTTPException(
                status_code=409,
                detail="L'opération ne correspond plus à la ligne du relevé ou a déjà été rapprochée.",
            )
        ligne.statut_rapprochement = "RAPPROCHE"
        ligne.operation_type = candidat.operation_type
        ligne.operation_id = candidat.operation_id
        ligne.score = scorer(
            LigneParse(
                numero_ligne=ligne.numero_ligne,
                date_operation=ligne.date_operation,
                date_valeur=ligne.date_valeur,
                reference=ligne.reference,
                tiers=ligne.tiers,
                libelle=ligne.libelle,
                debit=ligne.debit,
                credit=ligne.credit,
                solde=ligne.solde,
                empreinte=ligne.empreinte_ligne,
            ),
            candidat,
        )
        ligne.rapproche_par = user_id
        ligne.rapproche_le = datetime.now(timezone.utc)
        await db.flush()

        source = await _source_pointee(db, candidat.operation_type, candidat.operation_id)
        if source is not None and (
            candidat.operation_type != "VERSEMENT"
            or await _versements_bancaires_tous_rapproches(
                db, organisation_id=organisation_id, encaissement_id=source.id
            )
        ):
            _pointer(source, user_id=user_id, reference=ligne.reference)

    releve = await db.get(ComptaReleveBancaireImport, ligne.releve_import_id)
    if releve:
        await _recompter_releve(db, releve)
    await db.flush()
    return ligne


async def _annuler_decision(db: AsyncSession, ligne: ComptaReleveBancaireLigne) -> None:
    """Remet une ligne à décider : une erreur de pointage doit pouvoir se défaire.

    Le drapeau de pointage de l'opération est retiré avec elle (il n'a pu être
    posé que par cette ligne : une opération déjà pointée n'est jamais
    proposée). La ligne retrouve sa meilleure proposition d'origine.
    """
    if ligne.statut_rapprochement not in {"RAPPROCHE", "IGNORE"}:
        raise HTTPException(status_code=409, detail="Aucune décision à annuler sur cette ligne.")
    if ligne.statut_rapprochement == "RAPPROCHE":
        source = await _source_pointee(db, ligne.operation_type, ligne.operation_id)
        if source is not None and source.is_reconciled:
            _depointer(source)

    propositions = ligne.propositions or []
    top = propositions[0] if propositions else None
    ligne.statut_rapprochement = "INCERTAIN" if top else "NON_RAPPROCHE"
    ligne.operation_type = top["operation_type"] if top else None
    ligne.operation_id = top["operation_id"] if top else None
    ligne.score = Decimal(top["score"]) if top else None
    ligne.rapproche_par = None
    ligne.rapproche_le = None
