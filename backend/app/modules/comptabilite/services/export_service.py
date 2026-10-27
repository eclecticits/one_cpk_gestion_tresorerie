"""Exports comptables stricts et traçables.

Le fichier logiciel est volontairement construit côté serveur : les contrôles,
les types Excel et l'historique ne peuvent ainsi pas être contournés par une
version ancienne du navigateur.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from uuid import UUID

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.compte_bancaire import CompteBancaire
from app.modules.comptabilite.models import (
    ComptaCompte,
    ComptaEcriture,
    ComptaExportEcriture,
    ComptaExportHistorique,
    ComptaJournal,
    ComptaLigneEcriture,
    ComptaMappingCompteBancaire,
)
from app.modules.comptabilite.schemas.exports import (
    ExportAnomalieOut,
    ExportComptableApercuOut,
    ExportComptableFiltres,
    ExportLigneApercuOut,
)
from app.schemas.reports import ReportJournalResponse


CENT = Decimal("0.01")
STATUTS_EXPORTABLES = ("VALIDEE", "CLOTUREE")
ENTETE_LOGICIEL = ("Date>", "Pièce", "Compte", "Libellé", "Débit", "Crédit")
ENTETE_CONTROLE = (
    "Date d'opération",
    "Date de valeur",
    "Référence",
    "Bénéficiaire ou payeur",
    "Libellé",
    "Débit",
    "Crédit",
    "Solde",
)


def _q(value: Decimal | int | float | None) -> Decimal:
    return Decimal(value or 0).quantize(CENT, rounding=ROUND_HALF_UP)


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


@dataclass(slots=True)
class SelectionComptable:
    apercu: ExportComptableApercuOut
    lignes_export: list[ExportLigneApercuOut]
    ecritures: list[ComptaEcriture]
    comptes: dict[int, ComptaCompte]
    empreinte: str
    premier_export_id: UUID | None


async def _resoudre_compte_bancaire(
    db: AsyncSession,
    *,
    organisation_id: int,
    compte_bancaire_id: int,
    devise: str,
) -> tuple[CompteBancaire, int]:
    compte = (
        await db.execute(
            select(CompteBancaire).where(
                CompteBancaire.id == compte_bancaire_id,
                CompteBancaire.organisation_id == organisation_id,
                CompteBancaire.account_type == "BANK",
                CompteBancaire.is_active.is_(True),
            )
        )
    ).scalar_one_or_none()
    if compte is None:
        raise HTTPException(status_code=400, detail="Compte bancaire introuvable ou inactif.")
    if compte.devise != devise:
        raise HTTPException(
            status_code=400,
            detail=f"Le compte bancaire est tenu en {compte.devise}, pas en {devise}.",
        )
    mapping = (
        await db.execute(
            select(ComptaMappingCompteBancaire).where(
                ComptaMappingCompteBancaire.organisation_id == organisation_id,
                ComptaMappingCompteBancaire.compte_bancaire_id == compte_bancaire_id,
            )
        )
    ).scalar_one_or_none()
    if mapping is None:
        raise HTTPException(
            status_code=400,
            detail="Ce compte bancaire n'a pas de compte comptable associé. Configurez son imputation avant l'export.",
        )
    return compte, mapping.compte_id


async def preparer_export_comptable(
    db: AsyncSession,
    *,
    organisation_id: int,
    filtres: ExportComptableFiltres,
) -> SelectionComptable:
    """Charge et contrôle toutes les écritures correspondant aux filtres."""

    compte_comptable_banque_id: int | None = None
    if filtres.compte_bancaire_id is not None:
        _, compte_comptable_banque_id = await _resoudre_compte_bancaire(
            db,
            organisation_id=organisation_id,
            compte_bancaire_id=filtres.compte_bancaire_id,
            devise=filtres.devise,
        )

    stmt = (
        select(ComptaEcriture)
        .join(ComptaJournal, ComptaJournal.id == ComptaEcriture.journal_id)
        .options(selectinload(ComptaEcriture.lignes))
        .where(
            ComptaEcriture.organisation_id == organisation_id,
            ComptaEcriture.statut.in_(STATUTS_EXPORTABLES),
            ComptaEcriture.date_ecriture >= filtres.date_debut,
            ComptaEcriture.date_ecriture <= filtres.date_fin,
            ComptaEcriture.devise == filtres.devise,
        )
    )
    if filtres.journal == "CAISSE":
        stmt = stmt.where(ComptaJournal.type_journal == "CA")
    elif filtres.journal == "BANQUE":
        stmt = stmt.where(ComptaJournal.type_journal == "BQ")
    else:
        stmt = stmt.where(ComptaJournal.type_journal.in_(("CA", "BQ")))
    if compte_comptable_banque_id is not None:
        stmt = stmt.where(
            ComptaEcriture.lignes.any(ComptaLigneEcriture.compte_id == compte_comptable_banque_id)
        )
    stmt = stmt.order_by(ComptaEcriture.date_ecriture, ComptaEcriture.numero, ComptaEcriture.id)
    ecritures = list((await db.execute(stmt)).scalars().unique().all())

    compte_ids = {ligne.compte_id for ecriture in ecritures for ligne in (ecriture.lignes or [])}
    comptes: dict[int, ComptaCompte] = {}
    if compte_ids:
        comptes = {
            compte.id: compte
            for compte in (
                await db.execute(
                    select(ComptaCompte).where(
                        ComptaCompte.organisation_id == organisation_id,
                        ComptaCompte.id.in_(compte_ids),
                    )
                )
            ).scalars().all()
        }

    deja_exportees: dict[UUID, UUID] = {}
    if ecritures:
        hist_rows = (
            await db.execute(
                select(ComptaExportEcriture.ecriture_id, ComptaExportEcriture.export_id)
                .join(ComptaExportHistorique, ComptaExportHistorique.id == ComptaExportEcriture.export_id)
                .where(
                    ComptaExportEcriture.organisation_id == organisation_id,
                    ComptaExportHistorique.type_export == "LOGICIEL_COMPTABLE",
                    ComptaExportEcriture.ecriture_id.in_([e.id for e in ecritures]),
                )
                .order_by(ComptaExportHistorique.created_at)
            )
        ).all()
        deja_exportees = {ecriture_id: export_id for ecriture_id, export_id in hist_rows}

    anomalies: list[ExportAnomalieOut] = []
    apercu_lignes: list[ExportLigneApercuOut] = []
    total_debit = Decimal("0")
    total_credit = Decimal("0")
    origines_vues: dict[tuple[str, str, str], UUID] = {}

    if not ecritures:
        anomalies.append(
            ExportAnomalieOut(
                code="AUCUNE_ECRITURE",
                message="Aucune écriture validée ou clôturée ne correspond à cette sélection.",
            )
        )

    for ecriture in ecritures:
        numero = ecriture.numero or None
        piece = (ecriture.reference_piece or numero or "").strip()
        if not numero:
            anomalies.append(
                ExportAnomalieOut(
                    code="NUMERO_MANQUANT",
                    message="L'écriture validée ne possède pas de numéro comptable.",
                    ecriture_id=ecriture.id,
                )
            )
        if not piece:
            anomalies.append(
                ExportAnomalieOut(
                    code="PIECE_MANQUANTE",
                    message="La référence de pièce est absente.",
                    ecriture_id=ecriture.id,
                    numero=numero,
                )
            )
        if not ecriture.libelle.strip():
            anomalies.append(
                ExportAnomalieOut(
                    code="LIBELLE_MANQUANT",
                    message="Le libellé de l'écriture est vide.",
                    ecriture_id=ecriture.id,
                    numero=numero,
                )
            )
        if ecriture.id in deja_exportees:
            anomalies.append(
                ExportAnomalieOut(
                    code="DEJA_EXPORTEE",
                    message=f"Cette écriture figure déjà dans l'export {deja_exportees[ecriture.id]}.",
                    bloquante=not filtres.autoriser_reexport,
                    ecriture_id=ecriture.id,
                    numero=numero,
                )
            )
        if ecriture.module_origine and ecriture.type_origine and ecriture.objet_origine_id:
            cle_origine = (ecriture.module_origine, ecriture.type_origine, ecriture.objet_origine_id)
            precedente = origines_vues.get(cle_origine)
            if precedente and precedente != ecriture.id:
                anomalies.append(
                    ExportAnomalieOut(
                        code="ORIGINE_DUPLIQUEE",
                        message="Deux écritures actives portent la même origine métier.",
                        ecriture_id=ecriture.id,
                        numero=numero,
                    )
                )
            origines_vues[cle_origine] = ecriture.id

        lignes = sorted(ecriture.lignes or [], key=lambda ligne: ligne.ordre)
        if len(lignes) < 2:
            anomalies.append(
                ExportAnomalieOut(
                    code="ECRITURE_INCOMPLETE",
                    message="Une écriture doit comporter au moins deux lignes.",
                    ecriture_id=ecriture.id,
                    numero=numero,
                )
            )

        debit_ecriture = Decimal("0")
        credit_ecriture = Decimal("0")
        for ligne in lignes:
            debit = _q(ligne.debit)
            credit = _q(ligne.credit)
            debit_ecriture += debit
            credit_ecriture += credit
            compte = comptes.get(ligne.compte_id)
            compte_numero = (compte.numero if compte else "").strip()
            if compte is None:
                anomalies.append(
                    ExportAnomalieOut(
                        code="COMPTE_INTROUVABLE",
                        message=f"Le compte #{ligne.compte_id} est introuvable dans l'organisation.",
                        ecriture_id=ecriture.id,
                        numero=numero,
                    )
                )
            elif not compte.actif:
                anomalies.append(
                    ExportAnomalieOut(
                        code="COMPTE_INACTIF",
                        message=f"Le compte {compte.numero} est inactif.",
                        ecriture_id=ecriture.id,
                        numero=numero,
                    )
                )
            if not compte_numero.isdigit():
                anomalies.append(
                    ExportAnomalieOut(
                        code="COMPTE_NON_NUMERIQUE",
                        message=f"Le compte « {compte_numero or ligne.compte_id} » n'est pas numérique et ne respecte pas le modèle cible.",
                        ecriture_id=ecriture.id,
                        numero=numero,
                    )
                )
            elif compte_numero.startswith("0") or len(compte_numero) > 15:
                anomalies.append(
                    ExportAnomalieOut(
                        code="COMPTE_NON_REPRESENTABLE",
                        message=f"Le compte {compte_numero} ne peut pas être écrit comme nombre Excel sans perte.",
                        ecriture_id=ecriture.id,
                        numero=numero,
                    )
                )
            if ligne.devise != filtres.devise:
                anomalies.append(
                    ExportAnomalieOut(
                        code="DEVISE_INCOHERENTE",
                        message=f"La ligne est en {ligne.devise}, la sélection en {filtres.devise}.",
                        ecriture_id=ecriture.id,
                        numero=numero,
                    )
                )
            if debit < 0 or credit < 0 or (debit > 0 and credit > 0) or (debit == 0 and credit == 0):
                anomalies.append(
                    ExportAnomalieOut(
                        code="MONTANT_INVALIDE",
                        message="Une ligne doit porter un débit ou un crédit positif, jamais les deux.",
                        ecriture_id=ecriture.id,
                        numero=numero,
                    )
                )
            apercu_lignes.append(
                ExportLigneApercuOut(
                    ecriture_id=ecriture.id,
                    date=ecriture.date_ecriture,
                    piece=piece,
                    compte=compte_numero,
                    libelle=(ligne.libelle or ecriture.libelle).strip(),
                    debit=debit,
                    credit=credit,
                )
            )

        debit_ecriture = _q(debit_ecriture)
        credit_ecriture = _q(credit_ecriture)
        total_debit += debit_ecriture
        total_credit += credit_ecriture
        if debit_ecriture != credit_ecriture:
            anomalies.append(
                ExportAnomalieOut(
                    code="ECRITURE_DESEQUILIBREE",
                    message=f"Débit {debit_ecriture} ≠ crédit {credit_ecriture}.",
                    ecriture_id=ecriture.id,
                    numero=numero,
                )
            )

    total_debit = _q(total_debit)
    total_credit = _q(total_credit)
    ecart = _q(total_debit - total_credit)
    if ecritures and ecart != 0:
        anomalies.append(
            ExportAnomalieOut(
                code="SELECTION_DESEQUILIBREE",
                message=f"La sélection est déséquilibrée de {ecart} {filtres.devise}.",
            )
        )

    empreinte_payload = {
        "organisation_id": organisation_id,
        "date_debut": filtres.date_debut.isoformat(),
        "date_fin": filtres.date_fin.isoformat(),
        "journal": filtres.journal,
        "compte_bancaire_id": filtres.compte_bancaire_id,
        "devise": filtres.devise,
        "ecritures": [(str(e.id), e.version) for e in ecritures],
    }
    empreinte = _sha256(json.dumps(empreinte_payload, sort_keys=True).encode("utf-8"))
    exact_export = (
        await db.execute(
            select(ComptaExportHistorique.id)
            .where(
                ComptaExportHistorique.organisation_id == organisation_id,
                ComptaExportHistorique.type_export == "LOGICIEL_COMPTABLE",
                ComptaExportHistorique.empreinte_selection == empreinte,
            )
            .order_by(ComptaExportHistorique.created_at)
            .limit(1)
        )
    ).scalar_one_or_none()

    apercu = ExportComptableApercuOut(
        date_debut=filtres.date_debut,
        date_fin=filtres.date_fin,
        journal=filtres.journal,
        compte_bancaire_id=filtres.compte_bancaire_id,
        devise=filtres.devise,
        nombre_ecritures=len(ecritures),
        nombre_lignes=len(apercu_lignes),
        total_debit=total_debit,
        total_credit=total_credit,
        ecart=ecart,
        equilibre=ecart == 0,
        exportable=bool(ecritures) and not any(a.bloquante for a in anomalies),
        nombre_deja_exportees=len(deja_exportees),
        lignes=apercu_lignes[:500],
        anomalies=anomalies,
    )
    return SelectionComptable(
        apercu=apercu,
        lignes_export=apercu_lignes,
        ecritures=ecritures,
        comptes=comptes,
        empreinte=empreinte,
        premier_export_id=exact_export,
    )


def creer_xlsx_logiciel(lignes: list[ExportLigneApercuOut]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    ws.freeze_panes = "A2"
    remplissage = PatternFill("solid", fgColor="1F4E78")
    for col, valeur in enumerate(ENTETE_LOGICIEL, start=1):
        cellule = ws.cell(1, col, valeur)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = remplissage
        cellule.alignment = Alignment(horizontal="center")
    for row, ligne in enumerate(lignes, start=2):
        ws.cell(row, 1, datetime.combine(ligne.date, datetime.min.time())).number_format = "mm-dd-yy"
        ws.cell(row, 2, ligne.piece)
        ws.cell(row, 3, int(ligne.compte)).number_format = "0"
        ws.cell(row, 4, ligne.libelle)
        if ligne.debit:
            ws.cell(row, 5, float(ligne.debit)).number_format = "0.00"
        if ligne.credit:
            ws.cell(row, 6, float(ligne.credit)).number_format = "0.00"
    for index, largeur in enumerate((13, 24, 18, 72, 16, 16), start=1):
        ws.column_dimensions[get_column_letter(index)].width = largeur
    wb.properties.title = "Import comptable ONEC Smart"
    wb.properties.creator = "ONEC Smart"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _jour(value: datetime | date) -> date:
    """Le jour d'un mouvement, tel que le journal l'affiche.

    Les dates de trésorerie sont stockées avec fuseau horaire, qu'Excel refuse
    (openpyxl lève une TypeError à l'enregistrement) : on n'en garde que le jour.
    """
    return value.date() if isinstance(value, datetime) else value


def creer_controle_bancaire(journal: ReportJournalResponse, format_sortie: str) -> bytes:
    rows = [
        (
            _jour(ligne.date),
            None,
            ligne.reference or "",
            ligne.tiers or "",
            ligne.libelle or "",
            _q(ligne.sortie),
            _q(ligne.entree),
            _q(ligne.solde),
        )
        for ligne in journal.lignes
    ]
    if format_sortie == "CSV":
        text_buffer = io.StringIO(newline="")
        writer = csv.writer(text_buffer, delimiter=";")
        writer.writerow(ENTETE_CONTROLE)
        for row in rows:
            writer.writerow(
                [
                    row[0].isoformat(),
                    "",
                    row[2], row[3], row[4],
                    f"{row[5]:.2f}" if row[5] else "",
                    f"{row[6]:.2f}" if row[6] else "",
                    f"{row[7]:.2f}",
                ]
            )
        return ("\ufeff" + text_buffer.getvalue()).encode("utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Contrôle bancaire"
    ws.freeze_panes = "A2"
    for col, valeur in enumerate(ENTETE_CONTROLE, start=1):
        cellule = ws.cell(1, col, valeur)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F4E78")
        cellule.alignment = Alignment(horizontal="center")
    for row_index, row in enumerate(rows, start=2):
        ws.cell(row_index, 1, row[0]).number_format = "dd/mm/yyyy"
        ws.cell(row_index, 2, row[1]).number_format = "dd/mm/yyyy"
        ws.cell(row_index, 3, row[2])
        ws.cell(row_index, 4, row[3])
        ws.cell(row_index, 5, row[4])
        for col_index, value in enumerate(row[5:], start=6):
            ws.cell(row_index, col_index, float(value) if value else None).number_format = "#,##0.00"
    for index, largeur in enumerate((16, 16, 24, 32, 64, 16, 16, 18), start=1):
        ws.column_dimensions[get_column_letter(index)].width = largeur
    wb.properties.title = "Journal de contrôle bancaire ONEC Smart"
    wb.properties.creator = "ONEC Smart"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def historiser_export(
    db: AsyncSession,
    *,
    organisation_id: int,
    user_id: UUID | None,
    type_export: str,
    filtres: ExportComptableFiltres,
    nom_fichier: str,
    contenu: bytes,
    empreinte_selection: str,
    nombre_ecritures: int,
    nombre_lignes: int,
    total_debit: Decimal,
    total_credit: Decimal,
    ecriture_ids: list[UUID] | None = None,
    reexport_de_id: UUID | None = None,
) -> ComptaExportHistorique:
    historique = ComptaExportHistorique(
        organisation_id=organisation_id,
        type_export=type_export,
        format=filtres.format,
        journal=filtres.journal,
        compte_bancaire_id=filtres.compte_bancaire_id,
        devise=filtres.devise,
        date_debut=filtres.date_debut,
        date_fin=filtres.date_fin,
        nombre_ecritures=nombre_ecritures,
        nombre_lignes=nombre_lignes,
        total_debit=_q(total_debit),
        total_credit=_q(total_credit),
        ecart=_q(total_debit - total_credit),
        nom_fichier=nom_fichier,
        empreinte_fichier=_sha256(contenu),
        empreinte_selection=empreinte_selection,
        genere_par=user_id,
        reexport_de_id=reexport_de_id,
    )
    db.add(historique)
    await db.flush()
    for ecriture_id in ecriture_ids or []:
        db.add(
            ComptaExportEcriture(
                organisation_id=organisation_id,
                export_id=historique.id,
                ecriture_id=ecriture_id,
            )
        )
    await db.flush()
    return historique


def exiger_exportable(selection: SelectionComptable, filtres: ExportComptableFiltres) -> None:
    if filtres.format != "XLSX":
        raise HTTPException(
            status_code=400,
            detail="Le format logiciel comptable est XLSX : CSV ne conserve pas les types exigés par le modèle.",
        )
    if selection.premier_export_id and not filtres.autoriser_reexport:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"Cette sélection a déjà été générée dans l'export {selection.premier_export_id}. "
                "Activez explicitement la réexportation si le comptable confirme qu'elle est nécessaire."
            ),
        )
    if not selection.apercu.exportable:
        messages = [a.message for a in selection.apercu.anomalies if a.bloquante]
        raise HTTPException(
            status_code=400,
            detail={"message": "Export bloqué par les contrôles comptables.", "anomalies": messages[:30]},
        )


def empreinte_controle(journal: ReportJournalResponse, filtres: ExportComptableFiltres) -> str:
    payload = {
        "date_debut": filtres.date_debut.isoformat(),
        "date_fin": filtres.date_fin.isoformat(),
        "compte_bancaire_id": filtres.compte_bancaire_id,
        "devise": filtres.devise,
        "lignes": [
            (
                ligne.date.isoformat(), ligne.reference, str(ligne.entree),
                str(ligne.sortie), str(ligne.solde), ligne.transaction_id,
            )
            for ligne in journal.lignes
        ],
    }
    return _sha256(json.dumps(payload, sort_keys=True).encode("utf-8"))
